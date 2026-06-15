#!/usr/bin/env python3
"""Collect read-only host evidence through doops and write an inventory JSON."""

from __future__ import annotations

import argparse
import base64
import csv
import datetime as dt
import gzip
import ipaddress
import json
import os
import re
import shlex
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path
from typing import Any


START_MARKER = "__SERVER_HEALTH_JSON_START__"
END_MARKER = "__SERVER_HEALTH_JSON_END__"
PAYLOAD_PREFIX = "SERVER_HEALTH_B64:"

MANAGED_ENVIRONMENTS = {
    "zheyin": {
        "target": "zheyin",
        "display_name": "Zhejiang Conservatory of Music",
        "aliases": {"zheyin", "zy", "zjcm", "浙音", "浙江音乐学院", "zhejiang-conservatory"},
    },
    "hdu": {
        "target": "hdu",
        "display_name": "Hangzhou Dianzi University",
        "aliases": {"hdu", "hangdian", "杭电", "杭州电子科技大学", "hangzhou-dianzi"},
    },
}

REMOTE_PYTHON = r"""
import json
import os
import platform
import re
import shutil
import subprocess
import time

AUTH_SECURITY_ENABLED = "__AUTH_SECURITY_ENABLED__" == "1"
AUTH_LOG_WINDOW_DAYS = int("__AUTH_LOG_WINDOW_DAYS__")


def run(command, timeout=4):
    try:
        result = subprocess.run(
            command,
            shell=isinstance(command, str),
            text=True,
            capture_output=True,
            timeout=timeout,
            check=False,
        )
        return result.stdout.strip(), result.stderr.strip(), result.returncode
    except Exception as exc:
        return "", exc.__class__.__name__, 127


def cpu_percent():
    def read_stat():
        with open("/proc/stat", "r", encoding="utf-8", errors="ignore") as handle:
            for line in handle:
                if line.startswith("cpu "):
                    values = [float(item) for item in line.split()[1:]]
                    idle = values[3] + (values[4] if len(values) > 4 else 0)
                    return sum(values), idle
        return None

    first = read_stat()
    time.sleep(1)
    second = read_stat()
    if not first or not second:
        return None
    total_delta = second[0] - first[0]
    idle_delta = second[1] - first[1]
    if total_delta <= 0:
        return None
    return round((1 - idle_delta / total_delta) * 100, 1)


def memory_percent():
    values = {}
    try:
        with open("/proc/meminfo", "r", encoding="utf-8", errors="ignore") as handle:
            for line in handle:
                key, raw = line.split(":", 1)
                values[key] = float(raw.strip().split()[0])
    except Exception:
        return None
    total = values.get("MemTotal")
    available = values.get("MemAvailable")
    if not total or available is None:
        return None
    return round((total - available) * 100 / total, 1)


def load_per_core():
    try:
        load1 = os.getloadavg()[0]
        cores = os.cpu_count() or 1
        return round(load1 / cores, 2)
    except Exception:
        return None


def uptime_days():
    try:
        with open("/proc/uptime", "r", encoding="utf-8", errors="ignore") as handle:
            seconds = float(handle.read().split()[0])
        return round(seconds / 86400, 2)
    except Exception:
        return None


def disks():
    stdout, _, rc = run(["df", "-P", "-x", "tmpfs", "-x", "devtmpfs"], timeout=6)
    if rc != 0:
        return []
    rows = []
    for line in stdout.splitlines()[1:]:
        parts = line.split()
        if len(parts) < 6:
            continue
        mount = parts[5]
        used = parts[4].rstrip("%")
        try:
            used_percent = float(used)
        except ValueError:
            continue
        if mount.startswith(("/proc", "/sys", "/run")):
            continue
        rows.append({"mount": mount, "used_percent": used_percent})
    return rows[:20]


def process_running(pattern):
    if shutil.which("pgrep"):
        _, _, rc = run(["pgrep", "-f", pattern], timeout=3)
        return rc == 0
    stdout, _, rc = run("ps -ef | grep -F '{}' | grep -v grep".format(pattern), timeout=4)
    return rc == 0 and bool(stdout)


def service_status(name):
    if shutil.which("systemctl"):
        stdout, _, rc = run(["systemctl", "is-active", name], timeout=3)
        state = stdout.strip().lower()
        if rc == 0 and state == "active":
            return "running"
        if state == "failed":
            return "failed"
    if process_running(name):
        return "running"
    return None


def services():
    observed = []
    seen = set()
    for name in (
        "doops-agent",
        "kubelet",
        "containerd",
        "docker",
        "nginx",
        "mysql",
        "mariadb",
        "postgresql",
        "redis-server",
        "redis",
        "sshd",
    ):
        state = service_status(name)
        if state:
            observed.append({"name": name, "status": state, "expected": "running"})
            seen.add(name)
    if shutil.which("systemctl"):
        stdout, _, rc = run(["systemctl", "--failed", "--no-legend", "--plain", "--type=service"], timeout=5)
        if rc == 0:
            for line in stdout.splitlines():
                parts = line.split()
                if not parts:
                    continue
                name = parts[0]
                if name in seen:
                    continue
                observed.append({"name": name, "status": "failed", "expected": "running"})
                seen.add(name)
                if len(observed) >= 30:
                    break
    return observed


def processes():
    rows_by_pid = {}

    def parse(stdout):
        for line in stdout.splitlines()[1:]:
            parts = line.split(None, 5)
            if len(parts) < 6:
                continue
            pid, user, state, cpu, memory, name = parts
            try:
                pid_value = int(pid)
                cpu_value = round(float(cpu), 1)
                memory_value = round(float(memory), 1)
            except ValueError:
                continue
            rows_by_pid[pid_value] = {
                "pid": pid_value,
                "user": user[:32],
                "state": state[:16],
                "cpu_percent": cpu_value,
                "memory_percent": memory_value,
                "name": name[:80],
            }

    commands = [
        ["ps", "-eo", "pid,user,stat,pcpu,pmem,comm", "--sort=-pcpu"],
        ["ps", "-eo", "pid,user,stat,pcpu,pmem,comm", "--sort=-pmem"],
    ]
    for command in commands:
        stdout, _, rc = run(command, timeout=5)
        if rc == 0:
            parse(stdout)
    return sorted(
        rows_by_pid.values(),
        key=lambda item: (float(item.get("cpu_percent") or 0), float(item.get("memory_percent") or 0)),
        reverse=True,
    )[:12]


def mask_identifier(value):
    text = str(value or "").strip()
    if not text:
        return ""
    return (text[:3] if len(text) > 4 else text[:1]) + "***"


def top_counts(counter, key_name, mask=False):
    rows = []
    for key, count in sorted(counter.items(), key=lambda item: (-item[1], item[0]))[:5]:
        rows.append({key_name: mask_identifier(key) if mask else key, "count": count})
    return rows


def parse_auth_log(text):
    failed = 0
    success = 0
    sources = {}
    users = {}
    failed_pairs = {}
    successful_pairs = set()
    for raw_line in str(text or "").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        failed_match = re.search(r"Failed password for (?:invalid user )?(?P<user>\S+) from (?P<source>[0-9a-fA-F:.]+)", line)
        accepted_match = re.search(r"Accepted (?:password|publickey) for (?P<user>\S+) from (?P<source>[0-9a-fA-F:.]+)", line)
        pam_match = re.search(r"authentication failure;.*rhost=(?P<source>\S+).*user=(?P<user>\S+)", line)
        if failed_match or pam_match:
            failed += 1
            match = failed_match or pam_match
            user = match.group("user")
            source = match.group("source")
            failed_pairs[(source, user)] = failed_pairs.get((source, user), 0) + 1
        elif accepted_match:
            success += 1
            user = accepted_match.group("user")
            source = accepted_match.group("source")
            successful_pairs.add((source, user))
        else:
            continue
        sources[source] = sources.get(source, 0) + 1
        users[user] = users.get(user, 0) + 1
    patterns = []
    for source, user in successful_pairs:
        count = failed_pairs.get((source, user), 0)
        if count:
            patterns.append({"type": "failed-then-success", "source": source, "user": mask_identifier(user), "failed_count": count})
    return {
        "failed_login_count": failed,
        "successful_login_count": success,
        "unique_source_count": len(sources),
        "top_sources": top_counts(sources, "source"),
        "top_users": top_counts(users, "user", mask=True),
        "patterns": patterns[:20],
    }


def read_auth_log_text():
    chunks = []
    if shutil.which("journalctl"):
        stdout, _, rc = run(["journalctl", "--since", "{} days ago".format(AUTH_LOG_WINDOW_DAYS), "-u", "sshd", "--no-pager"], timeout=8)
        if rc == 0 and stdout:
            chunks.append(stdout)
    for path in ("/var/log/auth.log", "/var/log/secure"):
        if os.path.exists(path) and os.access(path, os.R_OK):
            try:
                with open(path, "r", encoding="utf-8", errors="ignore") as handle:
                    chunks.append("\n".join(handle.readlines()[-5000:]))
            except Exception:
                pass
    return "\n".join(chunks)


def sshd_config_value(key):
    if shutil.which("sshd"):
        stdout, _, rc = run(["sshd", "-T"], timeout=4)
        if rc == 0:
            for line in stdout.splitlines():
                parts = line.strip().split(None, 1)
                if len(parts) == 2 and parts[0].lower() == key.lower():
                    return parts[1].strip().lower()
    for path in ("/etc/ssh/sshd_config",):
        if os.path.exists(path) and os.access(path, os.R_OK):
            try:
                with open(path, "r", encoding="utf-8", errors="ignore") as handle:
                    for line in handle:
                        line = line.strip()
                        if not line or line.startswith("#"):
                            continue
                        parts = line.split(None, 1)
                        if len(parts) == 2 and parts[0].lower() == key.lower():
                            return parts[1].strip().lower()
            except Exception:
                pass
    return None


def collect_auth_security():
    if not AUTH_SECURITY_ENABLED:
        return {"status": "not_enabled", "window_days": AUTH_LOG_WINDOW_DAYS, "method": "doops-self"}
    text = read_auth_log_text()
    status = "collected" if text else "unavailable"
    log_evidence = parse_auth_log(text)
    password_login = sshd_config_value("passwordauthentication")
    permit_root = sshd_config_value("permitrootlogin")
    max_auth = sshd_config_value("maxauthtries")
    fail2ban = service_status("fail2ban")
    return {
        "status": status,
        "window_days": AUTH_LOG_WINDOW_DAYS,
        "method": "doops-self",
        "log_evidence": log_evidence,
        "protection_config": {
            "password_login_enabled": None if password_login is None else password_login in ("yes", "true"),
            "root_login_enabled": None if permit_root is None else permit_root not in ("no", "prohibit-password", "forced-commands-only"),
            "max_auth_tries": int(max_auth) if str(max_auth or "").isdigit() else None,
            "account_lockout_enabled": any(os.path.exists(path) for path in ("/etc/security/faillock.conf", "/etc/pam.d/system-auth", "/etc/pam.d/password-auth")),
            "fail2ban_enabled": fail2ban == "running",
        },
        "password_strength": {"status": "not_enabled", "weak_count": 0, "medium_count": 0, "strong_count": 0, "reuse_group_count": 0, "findings": []},
    }


def main():
    metrics = {
        "cpu_percent": cpu_percent(),
        "memory_percent": memory_percent(),
        "load_per_core": load_per_core(),
        "uptime_days": uptime_days(),
        "disks": disks(),
    }
    metrics = {key: value for key, value in metrics.items() if value is not None}
    payload = {
        "hostname": platform.node() or os.environ.get("HOSTNAME") or "unknown",
        "platform": platform.platform(),
        "os_type": platform.system().lower() or "unknown",
        "metrics": metrics,
        "services": services(),
        "processes": processes(),
        "auth_security": collect_auth_security(),
        "collected_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
    }
    print(json.dumps(payload, ensure_ascii=False, sort_keys=True))


try:
    main()
except Exception as exc:
    print(json.dumps({"error": exc.__class__.__name__, "metrics": {}, "services": [], "processes": []}, ensure_ascii=False))
"""

REMOTE_PROBE_PYTHON = r"""
import base64
import ipaddress
import json
import platform
import shutil
import socket
import subprocess
import tempfile
import time


SERVERS = json.loads(base64.b64decode("__SERVERS_B64__").decode("utf-8"))
PROBE_TARGET = "__PROBE_TARGET__"
TIMEOUT = float("__PROBE_TIMEOUT__")


def valid_ip(address):
    try:
        ipaddress.ip_address(str(address))
        return True
    except Exception:
        return False


def ping(address):
    if not valid_ip(address):
        return {"status": "skipped", "latency_ms": None, "error": "invalid-address"}
    if not shutil.which("ping"):
        return {"status": "not-available", "latency_ms": None, "error": "ping-unavailable"}
    started = time.time()
    try:
        result = subprocess.run(
            ["ping", "-c", "1", "-W", str(max(1, int(TIMEOUT))), str(address)],
            text=True,
            capture_output=True,
            timeout=TIMEOUT + 1.5,
            check=False,
        )
    except Exception as exc:
        return {"status": "error", "latency_ms": None, "error": exc.__class__.__name__}
    elapsed = int((time.time() - started) * 1000)
    return {"status": "alive" if result.returncode == 0 else "unreachable", "latency_ms": elapsed if result.returncode == 0 else None}


def probe_port(address, port):
    if not valid_ip(address):
        return {"port": port, "status": "not-tested", "latency_ms": None, "error": "invalid-address"}
    started = time.time()
    try:
        with socket.create_connection((str(address), int(port)), timeout=TIMEOUT):
            return {"port": int(port), "status": "open", "latency_ms": int((time.time() - started) * 1000)}
    except socket.timeout:
        return {"port": int(port), "status": "timeout", "latency_ms": None}
    except OSError as exc:
        return {"port": int(port), "status": "closed", "latency_ms": None, "error": exc.__class__.__name__}


def main():
    observations = []
    for item in SERVERS:
        address = str(item.get("address") or "")
        ports = [int(port) for port in item.get("ports", []) if str(port).isdigit()]
        ping_result = ping(address)
        port_results = [probe_port(address, port) for port in ports]
        reachable = ping_result.get("status") == "alive" or any(result.get("status") == "open" for result in port_results)
        observations.append(
            {
                "index": item.get("index"),
                "address": address,
                "reachable": bool(reachable),
                "ping": ping_result,
                "ports": port_results,
            }
        )
    payload = {
        "probe_target": PROBE_TARGET,
        "probe_hostname": platform.node() or "unknown",
        "checked_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "servers": observations,
    }
    encoded = base64.b64encode(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")).decode("ascii")
    print("SERVER_HEALTH_B64:" + encoded)


try:
    main()
except Exception as exc:
    payload = {"probe_target": PROBE_TARGET, "error": exc.__class__.__name__, "servers": []}
    encoded = base64.b64encode(json.dumps(payload, ensure_ascii=False).encode("utf-8")).decode("ascii")
    print("SERVER_HEALTH_B64:" + encoded)
"""

REMOTE_HOST_METRICS_PYTHON = r"""
import base64
import http.client
import ipaddress
import json
import os
import shutil
import socket
import subprocess
import tempfile
import time


SERVERS = json.loads(base64.b64decode("__SERVERS_B64__").decode("utf-8"))
HOST_METRICS_TARGET = "__HOST_METRICS_TARGET__"
HOST_METRICS_PROFILE = "__HOST_METRICS_PROFILE__"
METRICS_TIMEOUT = int(float("__METRICS_TIMEOUT__"))
AUTH_SECURITY_ENABLED = "__AUTH_SECURITY_ENABLED__" == "1"
AUTH_LOG_WINDOW_DAYS = int("__AUTH_LOG_WINDOW_DAYS__")
SSH_USER = os.environ.get("SERVER_HEALTH_SSH_USER", "monitor")
SSH_KEY = os.environ.get("SERVER_HEALTH_SSH_KEY", "/opt/server-health-monitor/secrets/zheyin_monitor_ed25519")
WINRM_SECRET = os.environ.get("SERVER_HEALTH_WINRM_SECRET", "/opt/server-health-monitor/secrets/winrm_zheyin_monitor.json")

LINUX_REMOTE_PYTHON = r'''
import json
import os
import shutil
import subprocess
import time

def run(command, timeout=4):
    try:
        result = subprocess.run(command, text=True, capture_output=True, timeout=timeout, check=False)
        return result.stdout.strip(), result.stderr.strip(), result.returncode
    except Exception as exc:
        return "", exc.__class__.__name__, 127

def cpu_percent():
    def read_stat():
        with open("/proc/stat", "r", encoding="utf-8", errors="ignore") as handle:
            for line in handle:
                if line.startswith("cpu "):
                    values = [float(item) for item in line.split()[1:]]
                    idle = values[3] + (values[4] if len(values) > 4 else 0)
                    return sum(values), idle
        return None
    first = read_stat()
    time.sleep(1)
    second = read_stat()
    if not first or not second or second[0] <= first[0]:
        return None
    return round((1 - ((second[1] - first[1]) / (second[0] - first[0]))) * 100, 1)

def memory_percent():
    values = {}
    with open("/proc/meminfo", "r", encoding="utf-8", errors="ignore") as handle:
        for line in handle:
            key, raw = line.split(":", 1)
            values[key] = float(raw.strip().split()[0])
    total = values.get("MemTotal")
    available = values.get("MemAvailable")
    if not total or available is None:
        return None
    return round((total - available) * 100 / total, 1)

def uptime_days():
    with open("/proc/uptime", "r", encoding="utf-8", errors="ignore") as handle:
        return round(float(handle.read().split()[0]) / 86400, 2)

def disks():
    stdout, _, rc = run(["df", "-P", "-x", "tmpfs", "-x", "devtmpfs"], timeout=6)
    if rc != 0:
        return []
    rows = []
    for line in stdout.splitlines()[1:]:
        parts = line.split()
        if len(parts) < 6:
            continue
        try:
            used_percent = float(parts[4].rstrip("%"))
        except ValueError:
            continue
        mount = parts[5]
        if not mount.startswith(("/proc", "/sys", "/run")):
            rows.append({"mount": mount, "used_percent": used_percent})
    return rows[:20]

def service_status(name):
    if shutil.which("systemctl"):
        stdout, _, rc = run(["systemctl", "is-active", name], timeout=3)
        if rc == 0 and stdout.strip() == "active":
            return "running"
        if stdout.strip() == "failed":
            return "failed"
    stdout, _, rc = run(["pgrep", "-f", name], timeout=3) if shutil.which("pgrep") else ("", "", 1)
    return "running" if rc == 0 and stdout else None

def services():
    observed = []
    for name in ("sshd", "nginx", "mysql", "mariadb", "postgresql", "redis-server", "redis", "docker", "containerd", "kubelet"):
        state = service_status(name)
        if state:
            observed.append({"name": name, "status": state, "expected": "running"})
    return observed

def processes():
    rows_by_pid = {}
    def parse(stdout):
        for line in stdout.splitlines()[1:]:
            parts = line.split(None, 5)
            if len(parts) < 6:
                continue
            pid, user, state, cpu, memory, name = parts
            try:
                pid_value = int(pid)
                cpu_value = round(float(cpu), 1)
                memory_value = round(float(memory), 1)
            except ValueError:
                continue
            rows_by_pid[pid_value] = {
                "pid": pid_value,
                "user": user[:32],
                "state": state[:16],
                "cpu_percent": cpu_value,
                "memory_percent": memory_value,
                "name": name[:80],
            }
    for command in (
        ["ps", "-eo", "pid,user,stat,pcpu,pmem,comm", "--sort=-pcpu"],
        ["ps", "-eo", "pid,user,stat,pcpu,pmem,comm", "--sort=-pmem"],
    ):
        stdout, _, rc = run(command, timeout=5)
        if rc == 0:
            parse(stdout)
    return sorted(
        rows_by_pid.values(),
        key=lambda item: (float(item.get("cpu_percent") or 0), float(item.get("memory_percent") or 0)),
        reverse=True,
    )[:12]

metrics = {
    "cpu_percent": cpu_percent(),
    "memory_percent": memory_percent(),
    "load_per_core": round(os.getloadavg()[0] / (os.cpu_count() or 1), 2),
    "uptime_days": uptime_days(),
    "disks": disks(),
}
print(json.dumps({"metrics": {k: v for k, v in metrics.items() if v is not None}, "services": services(), "processes": processes()}, ensure_ascii=False))
'''

LINUX_REMOTE_SH_FALLBACK = r'''
read_cpu() {
  set -- $(grep '^cpu ' /proc/stat 2>/dev/null)
  total=0
  idle=0
  index=0
  shift
  for value in "$@"; do
    total=$((total + value))
    index=$((index + 1))
    if [ "$index" -eq 4 ] || [ "$index" -eq 5 ]; then
      idle=$((idle + value))
    fi
  done
  printf '%s %s\n' "$total" "$idle"
}

cpu1=$(read_cpu)
sleep 1
cpu2=$(read_cpu)
cpu_percent=$(awk -v a="$cpu1" -v b="$cpu2" 'BEGIN{split(a,x," ");split(b,y," ");dt=y[1]-x[1];di=y[2]-x[2];if(dt>0){printf "%.1f",(1-(di/dt))*100}else{printf "null"}}')
memory_percent=$(awk '/MemTotal/{t=$2}/MemAvailable/{a=$2} END{if(t>0){printf "%.1f",(t-a)*100/t}else{printf "null"}}' /proc/meminfo 2>/dev/null)
uptime_days=$(awk '{printf "%.2f",$1/86400}' /proc/uptime 2>/dev/null)
cores=$(getconf _NPROCESSORS_ONLN 2>/dev/null || echo 1)
load_per_core=$(awk -v c="$cores" '{if(c>0){printf "%.2f",$1/c}else{printf "null"}}' /proc/loadavg 2>/dev/null)
disks=$(df -P -x tmpfs -x devtmpfs 2>/dev/null | awk 'NR>1{mount=$6;used=$5;gsub("%","",used);if(mount !~ "^/(proc|sys|run)" && mount ~ /^[A-Za-z0-9_\\.\\/:=-]+$/){if(n>0){printf ","};printf "{\"mount\":\"%s\",\"used_percent\":%s}",mount,used;n++}}')
processes=$(ps -eo pid,user,stat,pcpu,pmem,comm --sort=-pcpu 2>/dev/null | awk 'NR>1 && n<12{pid=$1;user=$2;state=$3;cpu=$4;mem=$5;name=$6;gsub(/["\\]/,"",user);gsub(/["\\]/,"",state);gsub(/["\\]/,"",name);if(pid ~ /^[0-9]+$/){if(n>0){printf ","};printf "{\"pid\":%d,\"user\":\"%s\",\"state\":\"%s\",\"cpu_percent\":%.1f,\"memory_percent\":%.1f,\"name\":\"%s\"}",pid,user,state,cpu,mem,name;n++}}')
auth_security='{"status":"not_enabled","window_days":__AUTH_LOG_WINDOW_DAYS__,"method":"ssh"}'
if [ "__AUTH_SECURITY_ENABLED__" = "1" ]; then
  auth_log=""
  if command -v journalctl >/dev/null 2>&1; then
    auth_log="$(journalctl --since "__AUTH_LOG_WINDOW_DAYS__ days ago" -u sshd --no-pager 2>/dev/null | tail -n 5000)"
  fi
  for log in /var/log/auth.log /var/log/secure; do
    if [ -r "$log" ]; then
      auth_log="$auth_log
$(tail -n 5000 "$log" 2>/dev/null)"
    fi
  done
  failed=$(printf '%s\n' "$auth_log" | grep -Eic 'Failed password|authentication failure' 2>/dev/null || echo 0)
  success=$(printf '%s\n' "$auth_log" | grep -Eic 'Accepted (password|publickey)' 2>/dev/null || echo 0)
  source_count=$(printf '%s\n' "$auth_log" | sed -nE 's/.* from ([0-9A-Fa-f:.]+) .*/\1/p;s/.*rhost=([^ ]+).*/\1/p' | sort -u | grep -c . 2>/dev/null || echo 0)
  password_login=null
  root_login=null
  max_auth=null
  if command -v sshd >/dev/null 2>&1; then
    sshd_t="$(sshd -T 2>/dev/null)"
    pa="$(printf '%s\n' "$sshd_t" | awk '$1=="passwordauthentication"{print $2;exit}')"
    pr="$(printf '%s\n' "$sshd_t" | awk '$1=="permitrootlogin"{print $2;exit}')"
    ma="$(printf '%s\n' "$sshd_t" | awk '$1=="maxauthtries"{print $2;exit}')"
    [ "$pa" = "yes" ] && password_login=true
    [ "$pa" = "no" ] && password_login=false
    case "$pr" in no|prohibit-password|forced-commands-only) root_login=false ;; yes|without-password) root_login=true ;; esac
    [ -n "$ma" ] && max_auth="$ma"
  fi
  lockout=false
  [ -r /etc/security/faillock.conf ] || grep -R "pam_faillock\\|pam_tally2" /etc/pam.d >/dev/null 2>&1
  [ "$?" -eq 0 ] && lockout=true
  fail2ban=false
  if command -v systemctl >/dev/null 2>&1 && systemctl is-active fail2ban >/dev/null 2>&1; then fail2ban=true; fi
  status=unavailable
  [ -n "$auth_log" ] && status=collected
  auth_security=$(printf '{"status":"%s","window_days":__AUTH_LOG_WINDOW_DAYS__,"method":"ssh","log_evidence":{"failed_login_count":%s,"successful_login_count":%s,"unique_source_count":%s,"top_sources":[],"top_users":[],"patterns":[]},"protection_config":{"password_login_enabled":%s,"root_login_enabled":%s,"max_auth_tries":%s,"account_lockout_enabled":%s,"fail2ban_enabled":%s},"password_strength":{"status":"not_enabled","weak_count":0,"medium_count":0,"strong_count":0,"reuse_group_count":0,"findings":[]}}' "$status" "$failed" "$success" "$source_count" "$password_login" "$root_login" "$max_auth" "$lockout" "$fail2ban")
fi
[ -n "$memory_percent" ] || memory_percent=null
[ -n "$uptime_days" ] || uptime_days=null
[ -n "$load_per_core" ] || load_per_core=null
printf '{"metrics":{"cpu_percent":%s,"memory_percent":%s,"load_per_core":%s,"uptime_days":%s,"disks":[%s]},"services":[],"processes":[%s],"auth_security":%s}\n' "$cpu_percent" "$memory_percent" "$load_per_core" "$uptime_days" "$disks" "$processes" "$auth_security"
'''

LINUX_REMOTE_SH = LINUX_REMOTE_SH_FALLBACK

WINDOWS_METRICS_PS = r'''
$os=Get-CimInstance Win32_OperatingSystem
$cpu=(Get-CimInstance Win32_Processor|Measure-Object LoadPercentage -Average).Average
$t=[double]$os.TotalVisibleMemorySize
$f=[double]$os.FreePhysicalMemory
$d=Get-CimInstance Win32_LogicalDisk -Filter 'DriveType=3'|%{$s=[double]$_.Size;$fr=[double]$_.FreeSpace;@{mount=$_.DeviceID;used_percent=[math]::Round((($s-$fr)*100/$s),1)}}
$svc=Get-Service WinRM -EA SilentlyContinue
$p=@{}
Get-Process|Sort-Object CPU -Descending|Select-Object -First 8|%{$p[$_.Id]=$null;$p[$_.Id]=$_}
Get-Process|Sort-Object WorkingSet64 -Descending|Select-Object -First 8|%{$p[$_.Id]=$null;$p[$_.Id]=$_}
$procs=$p.Values|Sort-Object CPU -Descending|Select-Object -First 12|%{@{pid=$_.Id;name=$_.ProcessName;state='running';cpu_seconds=[math]::Round([double]($_.CPU),1);memory_mb=[math]::Round([double]($_.WorkingSet64)/1MB,1)}}
$auth=@{status='not_enabled';window_days=__AUTH_LOG_WINDOW_DAYS__;method='winrm'}
# AUTH_SECURITY_START
if ('__AUTH_SECURITY_ENABLED__' -eq '1') {
  $auth=@{status='unavailable';window_days=__AUTH_LOG_WINDOW_DAYS__;method='winrm';log_evidence=@{failed_login_count=0;successful_login_count=0;unique_source_count=0;top_sources=@();top_users=@();patterns=@()};protection_config=@{password_login_enabled=$null;root_login_enabled=$null;max_auth_tries=$null;account_lockout_enabled=$null;fail2ban_enabled=$null;rdp_nla_enabled=$null;winrm_basic_enabled=$null;audit_policy_enabled=$null};password_strength=@{status='not_enabled';weak_count=0;medium_count=0;strong_count=0;reuse_group_count=0;findings=@()}}
  try {
    $since=(Get-Date).AddDays(-1 * __AUTH_LOG_WINDOW_DAYS__)
    $events=Get-WinEvent -FilterHashtable @{LogName='Security';Id=4624,4625,4771,4776,4740;StartTime=$since} -EA Stop | Select-Object -First 5000
    $failed=($events|?{$_.Id -in 4625,4771,4776}).Count
    $success=($events|?{$_.Id -eq 4624}).Count
    $sources=@{}
    $users=@{}
    foreach($e in $events){
      try{$x=[xml]$e.ToXml()}catch{continue}
      $src=''
      $user=''
      foreach($n in $x.Event.EventData.Data){
        if($n.Name -in @('IpAddress','WorkstationName') -and [string]$n.'#text' -and [string]$n.'#text' -ne '-'){$src=[string]$n.'#text'}
        if($n.Name -in @('TargetUserName','AccountName') -and [string]$n.'#text' -and [string]$n.'#text' -ne '-'){$user=[string]$n.'#text'}
      }
      if($src){$sources[$src]=1+($sources[$src] -as [int])}
      if($user){$u=if($user.Length -le 3){$user.Substring(0,1)+'***'}else{$user.Substring(0,[Math]::Min(3,$user.Length))+'***'};$users[$u]=1+($users[$u] -as [int])}
    }
    $topSources=@($sources.GetEnumerator()|Sort-Object Value -Descending|Select-Object -First 5|%{@{source=$_.Key;count=$_.Value}})
    $topUsers=@($users.GetEnumerator()|Sort-Object Value -Descending|Select-Object -First 5|%{@{user=$_.Key;count=$_.Value}})
    $net=(net accounts 2>$null | Out-String)
    $lockout=$false
    if($net -match '(?i)(Lockout threshold|锁定阈值)[^0-9]*([0-9]+)' -and [int]$Matches[2] -gt 0){$lockout=$true}
    $winrmAuth=(winrm get winrm/config/service/auth 2>$null | Out-String)
    $winrmSvc=(winrm get winrm/config/service 2>$null | Out-String)
    $basic=if($winrmAuth -match '(?i)Basic\s*=\s*true'){$true}else{$false}
    $unencrypted=if($winrmSvc -match '(?i)AllowUnencrypted\s*=\s*true'){$true}else{$false}
    $nla=$null
    try{$nla=((Get-ItemProperty 'HKLM:\SYSTEM\CurrentControlSet\Control\Terminal Server\WinStations\RDP-Tcp' -Name UserAuthentication -EA Stop).UserAuthentication -eq 1)}catch{}
    $patterns=@()
    if(($events|?{$_.Id -eq 4740}).Count -gt 0){$patterns+=@{type='account-lockout';count=($events|?{$_.Id -eq 4740}).Count}}
    if($failed -ge 50){$patterns+=@{type='high-failed-login';count=$failed}}
    $auth.status='collected'
    $auth.log_evidence=@{failed_login_count=$failed;successful_login_count=$success;unique_source_count=$sources.Count;top_sources=$topSources;top_users=$topUsers;patterns=$patterns}
    $auth.protection_config.account_lockout_enabled=$lockout
    $auth.protection_config.rdp_nla_enabled=$nla
    $auth.protection_config.winrm_basic_enabled=$basic
    $auth.protection_config.allow_unencrypted=$unencrypted
    $auth.protection_config.audit_policy_enabled=$true
  } catch {
    $auth.log_evidence.patterns=@(@{type='windows-eventlog-permission-denied'})
  }
}
# AUTH_SECURITY_END
@{metrics=@{cpu_percent=$cpu;memory_percent=[math]::Round((($t-$f)*100/$t),1);uptime_days=[math]::Round(((Get-Date)-$os.LastBootUpTime).TotalDays,2);disks=$d};services=@(@{name='WinRM';status=if($svc.Status -eq 'Running'){'running'}else{''+$svc.Status};expected='running'});processes=$procs;auth_security=$auth}|ConvertTo-Json -Depth 6 -Compress
'''


def valid_ip(address):
    try:
        ipaddress.ip_address(str(address))
        return True
    except Exception:
        return False


def error_item(error_type):
    return {"stage": "host-metrics", "type": error_type}


def method_for(item):
    os_type = str(item.get("os_type") or "").lower()
    ports = {int(port) for port in item.get("ports", []) if str(port).isdigit()}
    if os_type.startswith("win"):
        return "winrm"
    if os_type.startswith("linux") or 22 in ports:
        return "ssh"
    if {3389, 5985, 5986} & ports:
        return "winrm"
    return "not-supported"


def auth_candidates(auth):
    candidates = []
    if isinstance(auth.get("candidates"), list):
        for candidate in auth.get("candidates")[:4]:
            if not isinstance(candidate, dict):
                continue
            username = str(candidate.get("username") or "").strip()
            password = str(candidate.get("password") or "")
            if username and password:
                candidates.append({"username": username, "password": password})
    username = str(auth.get("username") or "").strip()
    password = str(auth.get("password") or "")
    if username and password:
        candidates.insert(0, {"username": username, "password": password})
    deduped = []
    seen = set()
    for candidate in candidates:
        key = (candidate["username"], candidate["password"])
        if key in seen:
            continue
        seen.add(key)
        deduped.append(candidate)
    return deduped[:4]


def auth_for(item):
    auth = item.get("runtime_auth") if isinstance(item.get("runtime_auth"), dict) else {}
    candidates = auth_candidates(auth)
    if candidates:
        first = candidates[0]
        return {
            "username": first["username"],
            "password": first["password"],
            "candidates": candidates,
            "winrm_ports": auth.get("winrm_ports") if isinstance(auth.get("winrm_ports"), list) else [],
        }
    return {}


def run(command, timeout, input_text=None, env=None):
    try:
        return subprocess.run(command, text=True, input=input_text, capture_output=True, timeout=timeout, check=False, env=env)
    except subprocess.TimeoutExpired:
        return None
    except Exception:
        return None


def run_ssh_askpass(command, password, input_text, timeout):
    askpass_dir = tempfile.mkdtemp(prefix="server-health-askpass-")
    askpass_path = os.path.join(askpass_dir, "askpass.sh")
    try:
        with open(askpass_path, "w", encoding="utf-8") as handle:
            handle.write("#!/bin/sh\nprintf '%s\\n' \"$SERVER_HEALTH_SSH_PASSWORD\"\n")
        os.chmod(askpass_path, 0o700)
        env = dict(os.environ)
        env.update(
            {
                "SERVER_HEALTH_SSH_PASSWORD": password,
                "SSH_ASKPASS": askpass_path,
                "SSH_ASKPASS_REQUIRE": "force",
                "DISPLAY": env.get("DISPLAY") or "server-health:0",
            }
        )
        runner = ["setsid"] if shutil.which("setsid") else []
        return run(runner + command, timeout=timeout, input_text=input_text, env=env)
    finally:
        try:
            os.remove(askpass_path)
        except OSError:
            pass
        try:
            os.rmdir(askpass_dir)
        except OSError:
            pass


def run_ssh_password(command, password, input_text, timeout):
    try:
        import pty

        master, slave = pty.openpty()
        proc = subprocess.Popen(command, stdin=slave, stdout=slave, stderr=slave, text=False, close_fds=True)
        os.close(slave)
        output = bytearray()
        sent_password = False
        sent_input = False
        deadline = time.time() + timeout
        password_bytes = (password + "\n").encode("utf-8", errors="ignore")
        input_bytes = input_text.encode("utf-8", errors="ignore")
        while time.time() < deadline:
            try:
                chunk = os.read(master, 4096)
            except BlockingIOError:
                chunk = b""
            except OSError:
                break
            if chunk:
                output.extend(chunk)
                lower = bytes(output[-4096:]).lower()
                if not sent_password and b"password" in lower:
                    os.write(master, password_bytes)
                    sent_password = True
                if sent_password and not sent_input and (b"\n" in chunk or b"\r" in chunk):
                    os.write(master, input_bytes)
                    try:
                        os.close(master)
                    except OSError:
                        pass
                    sent_input = True
                    break
            if proc.poll() is not None:
                break
            time.sleep(0.1)
        try:
            proc.wait(timeout=max(1, min(5, timeout)))
        except subprocess.TimeoutExpired:
            proc.kill()
        stdout = output.decode("utf-8", errors="replace")
        return {"returncode": proc.returncode if proc.returncode is not None else 124, "stdout": stdout}
    except Exception as exc:
        return {"returncode": 127, "stdout": exc.__class__.__name__}


def parse_json(stdout):
    text = str(stdout or "").strip()
    if not text:
        return None
    try:
        return json.loads(text)
    except Exception:
        pass
    for start_char, end_char in (("{", "}"), ("[", "]")):
        start = text.find(start_char)
        end = text.rfind(end_char)
        if start != -1 and end != -1 and end > start:
            try:
                return json.loads(text[start : end + 1])
            except Exception:
                continue
    return None


def clean_payload(payload):
    if not isinstance(payload, dict):
        return {"metrics": {}, "services": [], "processes": [], "auth_security": {"status": "not_enabled"}}
    metrics = payload.get("metrics") if isinstance(payload.get("metrics"), dict) else {}
    services = payload.get("services") if isinstance(payload.get("services"), list) else []
    processes = payload.get("processes") if isinstance(payload.get("processes"), list) else []
    auth_security = payload.get("auth_security") if isinstance(payload.get("auth_security"), dict) else {"status": "not_enabled"}
    clean_metrics = {}
    for key in ("cpu_percent", "memory_percent", "load_per_core", "uptime_days", "disks"):
        if key in metrics:
            clean_metrics[key] = metrics[key]
    clean_services = []
    for service in services:
        if isinstance(service, dict):
            clean_services.append(
                {
                    "name": str(service.get("name") or "service"),
                    "status": str(service.get("status") or service.get("state") or "unknown"),
                    "expected": str(service.get("expected") or "running"),
                }
            )
    clean_processes = []
    for process in processes:
        if isinstance(process, dict):
            item = {
                "pid": process.get("pid"),
                "name": str(process.get("name") or "process")[:80],
                "state": str(process.get("state") or "unknown")[:32],
            }
            if process.get("user") is not None:
                item["user"] = str(process.get("user") or "")[:32]
            for key in ("cpu_percent", "memory_percent", "cpu_seconds", "memory_mb"):
                if isinstance(process.get(key), (int, float)):
                    item[key] = process[key]
            clean_processes.append(item)
    return {"metrics": clean_metrics, "services": clean_services, "processes": clean_processes[:30], "auth_security": auth_security}


def xml_unescape(value):
    return str(value).replace("&quot;", '"').replace("&apos;", "'").replace("&gt;", ">").replace("&lt;", "<").replace("&amp;", "&")


def collect_winrm_basic(address, username, password, port):
    enc = base64.b64encode(WINDOWS_METRICS_PS.encode("utf-16le")).decode("ascii")
    body = '<s:Envelope xmlns:s="http://www.w3.org/2003/05/soap-envelope" xmlns:a="http://schemas.xmlsoap.org/ws/2004/08/addressing" xmlns:w="http://schemas.dmtf.org/wbem/wsman/1/wsman.xsd" xmlns:p="http://schemas.microsoft.com/wbem/wsman/1/windows/shell"><s:Header><a:To>/wsman</a:To><w:ResourceURI s:mustUnderstand="true">http://schemas.microsoft.com/wbem/wsman/1/windows/shell/cmd</w:ResourceURI><a:Action s:mustUnderstand="true">http://schemas.microsoft.com/wbem/wsman/1/windows/shell/Command</a:Action><a:MessageID>uuid:%s</a:MessageID></s:Header><s:Body><p:CommandLine><p:Command>powershell</p:Command><p:Arguments>-NoProfile -EncodedCommand %s</p:Arguments></p:CommandLine></s:Body></s:Envelope>' % (int(time.time() * 1000000), enc)
    auth = base64.b64encode((username + ":" + password).encode("utf-8")).decode("ascii")
    headers = {"Content-Type": "application/soap+xml;charset=UTF-8", "Authorization": "Basic " + auth}
    try:
        if int(port) == 5986:
            conn = http.client.HTTPSConnection(str(address), int(port), timeout=max(3, METRICS_TIMEOUT))
        else:
            conn = http.client.HTTPConnection(str(address), int(port), timeout=max(3, METRICS_TIMEOUT))
        conn.request("POST", "/wsman", body=body.encode("utf-8"), headers=headers)
        response = conn.getresponse()
        response_body = response.read().decode("utf-8", errors="replace")
        conn.close()
    except socket.timeout:
        return {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item("winrm-timeout")]}
    except Exception:
        return {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item("winrm-basic-failed")]}
    if response.status in (401, 403):
        return {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item("winrm-auth-failed")]}
    if response.status >= 400:
        return {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item("winrm-command-failed")]}
    streams = __import__("re").findall(r'<[^>]*Stream[^>]*Name="stdout"[^>]*>(.*?)</[^>]*Stream>', response_body, __import__("re").S)
    payload = parse_json("".join(xml_unescape(s) for s in streams)) or parse_json(response_body)
    if not payload:
        return {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item("winrm-invalid-payload")]}
    clean = clean_payload(payload)
    return {"status": "collected", **clean, "collection_errors": []}


def collect_ssh_once(address, auth):
    if not shutil.which("ssh"):
        return {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item("ssh-tool-missing")]}
    if not auth and not os.path.isfile(SSH_KEY):
        return {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item("ssh-key-missing")]}
    user = auth.get("username") if auth else SSH_USER
    if auth:
        if shutil.which("sshpass"):
            env = dict(os.environ)
            env["SSHPASS"] = auth["password"]
            command = [
                "sshpass",
                "-e",
                "ssh",
                "-o",
                "BatchMode=no",
                "-o",
                "PreferredAuthentications=password,keyboard-interactive",
                "-o",
                "PubkeyAuthentication=no",
                "-o",
                "PasswordAuthentication=yes",
                "-o",
                "KbdInteractiveAuthentication=yes",
                "-o",
                "NumberOfPasswordPrompts=1",
                "-o",
                "HostKeyAlgorithms=+ssh-rsa,ssh-dss",
                "-o",
                "KexAlgorithms=+diffie-hellman-group1-sha1,diffie-hellman-group14-sha1",
                "-o",
                "Ciphers=+aes128-cbc,3des-cbc",
                "-o",
                "MACs=+hmac-sha1",
                "-o",
                "StrictHostKeyChecking=no",
                "-o",
                "UserKnownHostsFile=/dev/null",
                "-o",
                "ConnectTimeout=" + str(max(3, min(METRICS_TIMEOUT, 15))),
                user + "@" + str(address),
                "sh -s",
            ]
            result = run(command, timeout=METRICS_TIMEOUT + 5, input_text=LINUX_REMOTE_SH, env=env)
        elif os.name == "posix" and shutil.which("setsid"):
            command = [
                "ssh",
                "-o",
                "BatchMode=no",
                "-o",
                "PreferredAuthentications=password,keyboard-interactive",
                "-o",
                "PubkeyAuthentication=no",
                "-o",
                "PasswordAuthentication=yes",
                "-o",
                "KbdInteractiveAuthentication=yes",
                "-o",
                "NumberOfPasswordPrompts=1",
                "-o",
                "HostKeyAlgorithms=+ssh-rsa,ssh-dss",
                "-o",
                "KexAlgorithms=+diffie-hellman-group1-sha1,diffie-hellman-group14-sha1",
                "-o",
                "Ciphers=+aes128-cbc,3des-cbc",
                "-o",
                "MACs=+hmac-sha1",
                "-o",
                "StrictHostKeyChecking=no",
                "-o",
                "UserKnownHostsFile=/dev/null",
                "-o",
                "ConnectTimeout=" + str(max(3, min(METRICS_TIMEOUT, 15))),
                user + "@" + str(address),
                "sh -s",
            ]
            result = run_ssh_askpass(command, auth["password"], LINUX_REMOTE_SH, timeout=METRICS_TIMEOUT + 8)
        elif os.name == "posix":
            command = [
                "ssh",
                "-tt",
                "-o",
                "BatchMode=no",
                "-o",
                "PreferredAuthentications=password,keyboard-interactive",
                "-o",
                "PubkeyAuthentication=no",
                "-o",
                "PasswordAuthentication=yes",
                "-o",
                "KbdInteractiveAuthentication=yes",
                "-o",
                "NumberOfPasswordPrompts=1",
                "-o",
                "HostKeyAlgorithms=+ssh-rsa,ssh-dss",
                "-o",
                "KexAlgorithms=+diffie-hellman-group1-sha1,diffie-hellman-group14-sha1",
                "-o",
                "Ciphers=+aes128-cbc,3des-cbc",
                "-o",
                "MACs=+hmac-sha1",
                "-o",
                "StrictHostKeyChecking=no",
                "-o",
                "UserKnownHostsFile=/dev/null",
                "-o",
                "ConnectTimeout=" + str(max(3, min(METRICS_TIMEOUT, 15))),
                user + "@" + str(address),
                "sh -s",
            ]
            pty_result = run_ssh_password(command, auth["password"], LINUX_REMOTE_SH, timeout=METRICS_TIMEOUT + 8)
            result = type("Result", (), {"returncode": pty_result["returncode"], "stdout": pty_result["stdout"]})()
        else:
            return {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item("ssh-password-tool-missing")]}
    else:
        command = [
            "ssh",
            "-i",
            SSH_KEY,
            "-o",
            "BatchMode=yes",
            "-o",
            "StrictHostKeyChecking=no",
            "-o",
            "UserKnownHostsFile=/dev/null",
            "-o",
            "ConnectTimeout=" + str(max(3, min(METRICS_TIMEOUT, 15))),
            user + "@" + str(address),
            "sh -s",
        ]
        result = run(command, timeout=METRICS_TIMEOUT + 5, input_text=LINUX_REMOTE_SH)
    if result is None:
        return {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item("ssh-timeout")]}
    if result.returncode != 0:
        lower_output = (str(getattr(result, "stdout", "") or "") + "\n" + str(getattr(result, "stderr", "") or "")).lower()
        if "permission denied" in lower_output or "authentication failed" in lower_output:
            error_type = "ssh-auth-failed"
        elif "too many authentication failures" in lower_output:
            error_type = "ssh-auth-failed"
        elif "connection timed out" in lower_output or "operation timed out" in lower_output:
            error_type = "ssh-timeout"
        elif "python3" in lower_output and ("not found" in lower_output or "no such file" in lower_output):
            error_type = "ssh-remote-python-missing"
        else:
            error_type = "ssh-command-failed"
        return {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item(error_type)]}
    payload = parse_json(result.stdout)
    if not payload:
        return {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item("ssh-invalid-payload")]}
    clean = clean_payload(payload)
    return {"status": "collected", **clean, "collection_errors": []}


def collect_ssh(address, auth):
    candidates = auth.get("candidates") if isinstance(auth.get("candidates"), list) else []
    if not candidates and auth:
        candidates = [{"username": auth.get("username"), "password": auth.get("password")}]
    if not candidates:
        return collect_ssh_once(address, auth)
    last_result = None
    for candidate in candidates[:4]:
        username = str(candidate.get("username") or "").strip()
        password = str(candidate.get("password") or "")
        if not username or not password:
            continue
        last_result = collect_ssh_once(address, {"username": username, "password": password})
        if last_result.get("status") == "collected":
            return last_result
        error_types = [item.get("type") for item in last_result.get("collection_errors", []) if isinstance(item, dict)]
        if error_types and error_types[0] not in {"ssh-auth-failed", "ssh-command-failed"}:
            return last_result
    return last_result or {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item("ssh-auth-failed")]}


def collect_winrm(address, auth):
    if not os.path.isfile(WINRM_SECRET) and not auth:
        return {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item("winrm-credential-missing")]}
    candidates = auth.get("candidates") if auth and isinstance(auth.get("candidates"), list) else []
    if auth and not candidates:
        candidates = [{"username": auth.get("username"), "password": auth.get("password")}]
    if not auth:
        try:
            secret = json.loads(open(WINRM_SECRET, "r", encoding="utf-8").read())
        except Exception:
            return {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item("winrm-credential-invalid")]}
        candidates = [{"username": str(secret.get("username") or secret.get("user") or ""), "password": str(secret.get("password") or "")}]
    candidates = [
        {"username": str(candidate.get("username") or "").strip(), "password": str(candidate.get("password") or "")}
        for candidate in candidates[:4]
        if isinstance(candidate, dict) and str(candidate.get("username") or "").strip() and str(candidate.get("password") or "")
    ]
    if not candidates:
        return {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item("winrm-credential-invalid")]}
    configured_ports = auth.get("winrm_ports") if auth else []
    ports = configured_ports if configured_ports else [5986, 5985]
    port_errors = []
    open_port = None
    for port in ports:
        try:
            with socket.create_connection((str(address), port), timeout=min(5, max(2, METRICS_TIMEOUT))):
                open_port = int(port)
                break
        except socket.timeout:
            port_errors.append("winrm-port-timeout")
        except OSError:
            port_errors.append("winrm-port-closed")
    else:
        return {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item(port_errors[0] if port_errors else "winrm-port-unreachable")]}
    last_result = None
    for candidate in candidates:
        username = candidate["username"]
        password = candidate["password"]
        if not shutil.which("pwsh"):
            last_result = collect_winrm_basic(address, username, password, open_port)
        else:
            env = dict(os.environ)
            env.update({"ZHEYIN_MONITOR_HOST": str(address), "ZHEYIN_MONITOR_USERNAME": username, "ZHEYIN_MONITOR_PASSWORD": password})
            winrm_ps = (
                "$s=ConvertTo-SecureString $env:ZHEYIN_MONITOR_PASSWORD -AsPlainText -Force;"
                "$c=[pscredential]::new($env:ZHEYIN_MONITOR_USERNAME,$s);"
                "Invoke-Command -ComputerName $env:ZHEYIN_MONITOR_HOST -Credential $c -ScriptBlock {"
                + WINDOWS_METRICS_PS
                + "}"
            )
            result = run(["pwsh", "-NoLogo", "-NoProfile", "-NonInteractive", "-Command", winrm_ps], timeout=METRICS_TIMEOUT + 8, env=env)
            if result is None:
                last_result = {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item("winrm-timeout")]}
            elif result.returncode != 0:
                last_result = {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item("winrm-command-failed")]}
            else:
                payload = parse_json(result.stdout)
                if not payload:
                    last_result = {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item("winrm-invalid-payload")]}
                else:
                    clean = clean_payload(payload)
                    last_result = {"status": "collected", **clean, "collection_errors": []}
        if last_result.get("status") == "collected":
            return last_result
        error_types = [item.get("type") for item in last_result.get("collection_errors", []) if isinstance(item, dict)]
        if error_types and error_types[0] not in {"winrm-auth-failed", "winrm-command-failed"}:
            return last_result
    return last_result or {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item("winrm-auth-failed")]}


def main():
    results = []
    for item in SERVERS:
        address = str(item.get("address") or "")
        method = method_for(item)
        if not valid_ip(address):
            result = {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item("invalid-address")]}
            method = "not-supported"
        elif method == "ssh":
            result = collect_ssh(address, auth_for(item))
        elif method == "winrm":
            result = collect_winrm(address, auth_for(item))
        else:
            result = {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item("os-not-supported")]}
        results.append(
            {
                "index": item.get("index"),
                "address": address,
                "metrics_method": method,
                "status": result.get("status") or "failed",
                "metrics": result.get("metrics") if isinstance(result.get("metrics"), dict) else {},
                "services": result.get("services") if isinstance(result.get("services"), list) else [],
                "collection_errors": result.get("collection_errors") if isinstance(result.get("collection_errors"), list) else [],
            }
        )
    payload = {
        "metrics_target": HOST_METRICS_TARGET,
        "profile": HOST_METRICS_PROFILE,
        "checked_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "servers": results,
    }
    encoded = base64.b64encode(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")).decode("ascii")
    print("SERVER_HEALTH_B64:" + encoded)


try:
    main()
except Exception as exc:
    payload = {
        "metrics_target": HOST_METRICS_TARGET,
        "profile": HOST_METRICS_PROFILE,
        "error": exc.__class__.__name__,
        "error_detail": str(exc)[:160] if exc.__class__.__name__ in {"NameError", "ImportError", "ModuleNotFoundError"} else "",
        "servers": [],
    }
    encoded = base64.b64encode(json.dumps(payload, ensure_ascii=False).encode("utf-8")).decode("ascii")
    print("SERVER_HEALTH_B64:" + encoded)
"""

REMOTE_SNMP_METRICS_PYTHON = r"""
import base64
import ipaddress
import json
import os
import re
import shutil
import subprocess
import time


SERVERS = json.loads(base64.b64decode("__SERVERS_B64__").decode("utf-8"))
SNMP_TARGET = "__SNMP_TARGET__"
SNMP_PROFILE = "__SNMP_PROFILE__"
SNMP_TIMEOUT = int(float("__SNMP_TIMEOUT__"))
SNMP_SECRET = os.environ.get("SERVER_HEALTH_SNMP_SECRET", "/opt/server-health-monitor/secrets/snmp_zheyin_monitor.json")

OID_SYS_UPTIME = "1.3.6.1.2.1.1.3.0"
OID_HR_SYSTEM_UPTIME = "1.3.6.1.2.1.25.1.1.0"
OID_CPU_IDLE = "1.3.6.1.4.1.2021.11.11.0"
OID_MEM_TOTAL = "1.3.6.1.4.1.2021.4.5.0"
OID_MEM_AVAIL = "1.3.6.1.4.1.2021.4.6.0"
OID_HR_PROCESSOR_LOAD = "1.3.6.1.2.1.25.3.3.1.2"
OID_HR_STORAGE = "1.3.6.1.2.1.25.2.3.1"
OID_HR_STORAGE_DESCR = OID_HR_STORAGE + ".3"
OID_HR_STORAGE_ALLOC = OID_HR_STORAGE + ".4"
OID_HR_STORAGE_SIZE = OID_HR_STORAGE + ".5"
OID_HR_STORAGE_USED = OID_HR_STORAGE + ".6"


def valid_ip(address):
    try:
        ipaddress.ip_address(str(address))
        return True
    except Exception:
        return False


def error_item(error_type):
    return {"stage": "snmp-metrics", "type": error_type}


def failed(item, error_type):
    return {
        "index": item.get("index"),
        "address": str(item.get("address") or ""),
        "metrics_method": "snmpv3",
        "status": "failed",
        "metrics": {},
        "services": [],
        "collection_errors": [error_item(error_type)],
    }


def parse_number(value):
    text = str(value or "").strip()
    if not text:
        return None
    match = re.search(r"\(([-+]?\d+(?:\.\d+)?)\)", text)
    if not match:
        match = re.search(r"[-+]?\d+(?:\.\d+)?", text)
    if not match:
        return None
    try:
        return float(match.group(1) if match.lastindex else match.group(0))
    except Exception:
        return None


def parse_text(value):
    text = str(value or "").strip()
    if len(text) >= 2 and text[0] == '"' and text[-1] == '"':
        return text[1:-1]
    return text


def classify_failure(stdout, stderr):
    lower = (str(stdout or "") + "\n" + str(stderr or "")).lower()
    if any(token in lower for token in ("authentication", "wrong digest", "unknown user", "decryption", "authorization")):
        return "snmp-auth-failed"
    if "timeout" in lower:
        return "snmp-timeout"
    if "no response" in lower:
        return "snmp-no-response"
    return "snmp-invalid-payload"


def run(command, timeout):
    try:
        return subprocess.run(command, text=True, capture_output=True, timeout=timeout, check=False)
    except subprocess.TimeoutExpired:
        return None
    except Exception:
        return None


def load_secret():
    if not os.path.isfile(SNMP_SECRET):
        return None, "snmp-credential-missing"
    try:
        data = json.loads(open(SNMP_SECRET, "r", encoding="utf-8").read())
    except Exception:
        return None, "snmp-credential-missing"
    secret = {
        "username": str(data.get("username") or data.get("user") or ""),
        "auth_protocol": str(data.get("auth_protocol") or "SHA"),
        "auth_key": str(data.get("auth_key") or ""),
        "priv_protocol": str(data.get("priv_protocol") or "AES"),
        "priv_key": str(data.get("priv_key") or ""),
        "security_level": str(data.get("security_level") or "authPriv"),
    }
    level = secret["security_level"]
    if not secret["username"]:
        return None, "snmp-credential-missing"
    if level in ("authNoPriv", "authPriv") and not secret["auth_key"]:
        return None, "snmp-credential-missing"
    if level == "authPriv" and not secret["priv_key"]:
        return None, "snmp-credential-missing"
    return secret, None


def snmp_base_args(secret):
    args = ["-v3", "-l", secret["security_level"], "-u", secret["username"]]
    if secret["security_level"] in ("authNoPriv", "authPriv"):
        args.extend(["-a", secret["auth_protocol"], "-A", secret["auth_key"]])
    if secret["security_level"] == "authPriv":
        args.extend(["-x", secret["priv_protocol"], "-X", secret["priv_key"]])
    args.extend(["-t", str(max(1, min(SNMP_TIMEOUT, 15))), "-r", "0"])
    return args


def snmpget(secret, address, oids):
    command = [shutil.which("snmpget")] + snmp_base_args(secret) + ["-Oqv", str(address)] + list(oids)
    result = run(command, timeout=max(3, SNMP_TIMEOUT + 2))
    if result is None:
        return None, "snmp-timeout"
    if result.returncode != 0:
        return None, classify_failure(result.stdout, result.stderr)
    return result.stdout.splitlines(), None


def snmpwalk(secret, address, oid):
    command = [shutil.which("snmpwalk")] + snmp_base_args(secret) + ["-On", "-Oq", str(address), oid]
    result = run(command, timeout=max(4, SNMP_TIMEOUT + 3))
    if result is None:
        return None, "snmp-timeout"
    if result.returncode != 0:
        return None, classify_failure(result.stdout, result.stderr)
    return result.stdout.splitlines(), None


def collect_processor_load(secret, address):
    rows, error = snmpwalk(secret, address, OID_HR_PROCESSOR_LOAD)
    if error or not rows:
        return None
    values = []
    for line in rows:
        value = parse_number(line.rsplit(" ", 1)[-1])
        if value is not None:
            values.append(value)
    if not values:
        return None
    return round(sum(values) / len(values), 1)


def hr_storage_rows(secret, address):
    rows, error = snmpwalk(secret, address, OID_HR_STORAGE)
    if error or not rows:
        return {}
    table = {}
    for line in rows:
        parts = line.split(" ", 1)
        if len(parts) != 2:
            continue
        oid, value = parts
        for base, field in (
            (OID_HR_STORAGE_DESCR, "descr"),
            (OID_HR_STORAGE_ALLOC, "allocation_units"),
            (OID_HR_STORAGE_SIZE, "size"),
            (OID_HR_STORAGE_USED, "used"),
        ):
            prefix = "." + base + "."
            if oid.startswith(prefix):
                index = oid[len(prefix) :]
                table.setdefault(index, {})[field] = value
                break
    return table


def derive_memory_from_storage(table):
    for row in table.values():
        descr = parse_text(row.get("descr")).lower()
        if "memory" not in descr:
            continue
        size = parse_number(row.get("size"))
        used = parse_number(row.get("used"))
        if size and used is not None:
            return round(used * 100 / size, 1)
    return None


def derive_disks_from_storage(table):
    disks = []
    for row in table.values():
        descr = parse_text(row.get("descr"))
        lower = descr.lower()
        if not descr or "memory" in lower or "swap" in lower or "virtual" in lower:
            continue
        size = parse_number(row.get("size"))
        used = parse_number(row.get("used"))
        if not size or used is None:
            continue
        if not (descr.startswith("/") or ":" in descr or "fixed" in lower or "disk" in lower):
            continue
        disks.append({"mount": descr[:80], "used_percent": round(used * 100 / size, 1)})
    return disks[:30]


def collect_snmp(address, secret):
    scalar_oids = [OID_SYS_UPTIME, OID_CPU_IDLE, OID_MEM_TOTAL, OID_MEM_AVAIL]
    lines, scalar_error = snmpget(secret, address, scalar_oids)
    metrics = {}
    if lines and len(lines) >= 4:
        uptime_ticks = parse_number(lines[0])
        if uptime_ticks is not None:
            metrics["uptime_days"] = round(uptime_ticks / 100 / 86400, 2)
        idle = parse_number(lines[1])
        if idle is not None:
            metrics["cpu_percent"] = round(max(0, min(100, 100 - idle)), 1)
        total = parse_number(lines[2])
        available = parse_number(lines[3])
        if total and available is not None:
            metrics["memory_percent"] = round((total - available) * 100 / total, 1)
    elif scalar_error in {"snmp-auth-failed", "snmp-timeout", "snmp-no-response"}:
        return {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item(scalar_error)]}

    if "cpu_percent" not in metrics:
        cpu = collect_processor_load(secret, address)
        if cpu is not None:
            metrics["cpu_percent"] = cpu
    table = hr_storage_rows(secret, address)
    if "memory_percent" not in metrics:
        memory = derive_memory_from_storage(table)
        if memory is not None:
            metrics["memory_percent"] = memory
    disks = derive_disks_from_storage(table)
    if disks:
        metrics["disks"] = disks

    if not metrics:
        return {"status": "failed", "metrics": {}, "services": [], "collection_errors": [error_item(scalar_error or "snmp-invalid-payload")]}
    return {"status": "collected", "metrics": metrics, "services": [], "collection_errors": []}


def main():
    snmpget_path = shutil.which("snmpget")
    snmpwalk_path = shutil.which("snmpwalk")
    secret, secret_error = load_secret()
    results = []
    for item in SERVERS:
        address = str(item.get("address") or "")
        if not valid_ip(address):
            result = failed(item, "invalid-address")
        elif not snmpget_path or not snmpwalk_path:
            result = failed(item, "snmp-tool-missing")
        elif secret_error:
            result = failed(item, secret_error)
        else:
            collected = collect_snmp(address, secret)
            result = {
                "index": item.get("index"),
                "address": address,
                "metrics_method": "snmpv3",
                "status": collected.get("status") or "failed",
                "metrics": collected.get("metrics") if isinstance(collected.get("metrics"), dict) else {},
                "services": [],
                "collection_errors": collected.get("collection_errors") if isinstance(collected.get("collection_errors"), list) else [],
            }
        results.append(result)
    payload = {
        "metrics_target": SNMP_TARGET,
        "profile": SNMP_PROFILE,
        "checked_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "servers": results,
    }
    encoded = base64.b64encode(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")).decode("ascii")
    print("SERVER_HEALTH_B64:" + encoded)


try:
    main()
except Exception as exc:
    payload = {"metrics_target": SNMP_TARGET, "profile": SNMP_PROFILE, "error": exc.__class__.__name__, "servers": []}
    encoded = base64.b64encode(json.dumps(payload, ensure_ascii=False).encode("utf-8")).decode("ascii")
    print("SERVER_HEALTH_B64:" + encoded)
"""


def now_local() -> str:
    return dt.datetime.now(dt.timezone.utc).astimezone().isoformat(timespec="seconds")


def normalize_key(value: str) -> str:
    return re.sub(r"[\s_-]+", "", value.strip().lower())


def resolve_environment(value: str | None) -> dict[str, Any]:
    if not value:
        return {"key": "", "target": "", "display_name": "Server environment", "managed": False}
    normalized = normalize_key(value)
    for key, meta in MANAGED_ENVIRONMENTS.items():
        aliases = {normalize_key(alias) for alias in meta["aliases"]}
        if normalized == normalize_key(key) or normalized in aliases:
            resolved = dict(meta)
            resolved["key"] = key
            resolved["managed"] = True
            return resolved
    return {"key": value, "target": value, "display_name": value, "managed": False}


def resolve_cli_candidate(value: str) -> str | None:
    candidate = Path(value).expanduser()
    if candidate.is_file():
        return str(candidate)
    found = shutil.which(value)
    if found:
        return found
    return None


def find_doops(explicit: str | None = None) -> str:
    for source, value in (("--doops", explicit), ("DOOPS_BIN", os.environ.get("DOOPS_BIN"))):
        if not value:
            continue
        resolved = resolve_cli_candidate(value)
        if resolved:
            return resolved
        raise FileNotFoundError(f"doops CLI from {source} not found: {value}")

    for name in ("doops", "doops.exe"):
        found = shutil.which(name)
        if found:
            return found

    for candidate in (
        Path.home() / ".local" / "bin" / "doops",
        Path.home() / ".local" / "bin" / "doops.exe",
    ):
        if candidate.is_file():
            return str(candidate)
    raise FileNotFoundError("doops CLI not found via --doops, DOOPS_BIN, PATH, or ~/.local/bin")


def find_doops_config() -> Path | None:
    candidates: list[Path] = []
    env_config = os.environ.get("DOOPS_CONFIG")
    if env_config:
        candidates.append(Path(env_config).expanduser())
    candidates.extend(
        [
            Path.cwd() / ".agent" / "skills" / "doops" / "config.json",
            Path.home() / ".config" / "doops" / "config.json",
        ]
    )
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    return None


def config_target_names(config_path: Path | None) -> list[str]:
    if not config_path:
        return []
    try:
        data = json.loads(config_path.read_text(encoding="utf-8"))
    except Exception:
        return []
    names: set[str] = set()
    targets = data.get("targets") if isinstance(data, dict) else None
    if isinstance(targets, dict):
        names.update(str(key) for key in targets)
    elif isinstance(targets, list):
        for item in targets:
            if isinstance(item, dict):
                for key in ("name", "target", "id", "alias"):
                    if item.get(key):
                        names.add(str(item[key]))
            elif item:
                names.add(str(item))
    return sorted(names)


def self_node_script_template(auth_security: bool = False) -> str:
    script = REMOTE_PYTHON.strip()
    if not auth_security:
        script = script.replace("import re\n", "")
        script = script.replace(
            '\nAUTH_SECURITY_ENABLED = "__AUTH_SECURITY_ENABLED__" == "1"\nAUTH_LOG_WINDOW_DAYS = int("__AUTH_LOG_WINDOW_DAYS__")\n',
            "\n",
        )
        start_index = script.find("\ndef mask_identifier(")
        end_index = script.find("\ndef main():", start_index)
        if start_index != -1 and end_index != -1:
            script = script[:start_index] + script[end_index:]
        script = script.replace('        "auth_security": collect_auth_security(),\n', "")
    return script


def build_remote_command(auth_security: bool = False, auth_log_window_days: int = 7) -> str:
    script = (
        self_node_script_template(auth_security)
        .replace("__AUTH_SECURITY_ENABLED__", "1" if auth_security else "0")
        .replace("__AUTH_LOG_WINDOW_DAYS__", str(max(1, int(auth_log_window_days or 7))))
    )
    script_b64 = base64.b64encode(gzip.compress(script.encode("utf-8"))).decode("ascii")
    fallback_b64 = base64.b64encode(
        json.dumps(
            {
                "hostname": "unknown",
                "platform": "unknown",
                "os_type": "linux",
                "metrics": {},
                "services": [],
                "notes": "python-unavailable",
            },
            ensure_ascii=False,
        ).encode("utf-8")
    ).decode("ascii")
    return " ".join(
        [
            "set +e;",
            f"echo {START_MARKER};",
            f"if command -v python3 >/dev/null 2>&1; then printf {script_b64} | base64 -d | gzip -dc | python3; "
            f"elif command -v python >/dev/null 2>&1; then printf {script_b64} | base64 -d | gzip -dc | python; "
            f"else printf {fallback_b64} | base64 -d; fi;",
            f"echo {END_MARKER}",
        ]
    )


def run_command(args: list[str], timeout: int) -> subprocess.CompletedProcess[str]:
    env = {**os.environ, "DOOPS_ALLOW_INSECURE_GATEWAY": os.environ.get("DOOPS_ALLOW_INSECURE_GATEWAY", "1")}
    return subprocess.run(
        args,
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        timeout=timeout,
        check=False,
        env=env,
    )


def run_doops_command(args: list[str], timeout: int) -> dict[str, Any]:
    started = time.perf_counter()
    result = run_command(args, timeout=timeout)
    return {
        "returncode": result.returncode,
        "stdout": result.stdout,
        "stderr": result.stderr,
        "duration_ms": int((time.perf_counter() - started) * 1000),
    }


def sanitize_doops_result(result: dict[str, Any], limit: int = 4000) -> dict[str, Any]:
    return {
        "returncode": int(result.get("returncode", 1)),
        "stdout_excerpt": redact(str(result.get("stdout") or ""), limit=limit),
        "stderr_excerpt": redact(str(result.get("stderr") or ""), limit=limit),
        "duration_ms": int(result.get("duration_ms") or 0),
    }


def run_doops_targets(doops: str, target: str, timeout: int) -> dict[str, Any]:
    return run_doops_command([doops, "targets", "--target", target], timeout=min(timeout, 30))


def run_doops_info(doops: str, target: str, session: str, timeout: int) -> dict[str, Any]:
    return run_doops_command([doops, "-session", session, "info", "--target", target], timeout=min(timeout, 30))


def run_doops_exec(doops: str, target: str, session: str, command: str, timeout: int) -> dict[str, Any]:
    return run_doops_command([doops, "-session", session, "exec", "--target", target, "--cmd", command], timeout=timeout)


def doops_local_src_path(src: str) -> str:
    if os.name != "nt":
        return src
    text = str(Path(src).resolve())
    match = re.match(r"^([A-Za-z]):[\\/](.*)$", text)
    if not match:
        return text.replace("\\", "/")
    drive = match.group(1).lower()
    rest = match.group(2).replace("\\", "/")
    return f"/mnt/{drive}/{rest}"


def run_doops_push(doops: str, target: str, session: str, src: str, timeout: int) -> dict[str, Any]:
    return run_doops_command([doops, "-session", session, "push", "--target", target, "--src", doops_local_src_path(src)], timeout=min(timeout, 60))


def run_doops_read(doops: str, target: str, session: str, path: str, timeout: int) -> dict[str, Any]:
    return run_doops_command([doops, "-session", session, "read", "--target", target, "--path", path], timeout=min(timeout, 60))


def run_doops_write(doops: str, target: str, session: str, path: str, content: str, timeout: int) -> dict[str, Any]:
    return run_doops_command([doops, "-session", session, "write", "--target", target, "--path", path, content], timeout=min(timeout, 60))


def run_doops_clean(doops: str, target: str, session: str, timeout: int) -> dict[str, Any]:
    return run_doops_command([doops, "-session", session, "clean", "--target", target, "--workspace", session], timeout=min(timeout, 60))


def extract_payload(stdout: str) -> dict[str, Any]:
    start = stdout.rfind(START_MARKER)
    end = stdout.rfind(END_MARKER)
    if start == -1 or end == -1 or end <= start:
        raise ValueError("doops output did not contain health JSON markers")
    json_text = stdout[start + len(START_MARKER) : end].strip()
    if not json_text:
        raise ValueError("doops health JSON payload is empty")
    for line in json_text.splitlines():
        line = line.strip()
        if PAYLOAD_PREFIX in line:
            encoded = line.split(PAYLOAD_PREFIX, 1)[1].strip()
            return json.loads(base64.b64decode(encoded).decode("utf-8"))
    return json.loads(json_text)


def redact(text: str, limit: int = 4000) -> str:
    clean = re.sub(r"(?i)(token|password|secret|cookie|api[_-]?key|auth[_-]?key|priv[_-]?key|community)(\s*[:=]\s*)\S+", r"\1\2<redacted>", text)
    clean = re.sub(r"(?i)snmp[-_ ]?community\s+\S+", "snmp-community <redacted>", clean)
    clean = re.sub(r"(?i)(authorization)(\s*[:=]\s*)\S+", r"\1\2<redacted>", clean)
    clean = re.sub(r"(?i)\bbearer\s+[A-Za-z0-9._~+/=-]+", "Bearer <redacted>", clean)
    clean = clean.replace(chr(0xFFFD), "\\ufffd")
    clean = "".join(
        character if character in "\n\r\t" or (ord(character) >= 32 and ord(character) != 127 and not 0x80 <= ord(character) <= 0x9F) else " "
        for character in clean
    )
    return clean[-limit:]


def parse_boolish(value: Any) -> bool | None:
    text = str(value or "").strip().lower()
    if text in {"true", "yes", "y", "1", "busy"}:
        return True
    if text in {"false", "no", "n", "0", "idle"}:
        return False
    return None


def last_seen_age_seconds(value: str) -> int | None:
    token = str(value or "").split()[0] if str(value or "").strip() else ""
    if not token:
        return None
    try:
        parsed = dt.datetime.fromisoformat(token.replace("Z", "+00:00"))
    except ValueError:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=dt.datetime.now().astimezone().tzinfo)
    return max(0, int((dt.datetime.now(dt.timezone.utc) - parsed.astimezone(dt.timezone.utc)).total_seconds()))


def parse_doops_targets(stdout: str) -> dict[str, Any]:
    rows: list[dict[str, Any]] = []
    for raw_line in str(stdout or "").splitlines():
        line = raw_line.strip()
        if not line or set(line) <= {"-", " "}:
            continue
        lower = line.lower()
        if "cluster" in lower and "instance" in lower:
            continue
        parts = line.split()
        if len(parts) >= 4:
            busy = parse_boolish(parts[2])
            last_seen = redact(" ".join(parts[3:]), limit=200)
            rows.append(
                {
                    "cluster": parts[0],
                    "instance": parts[1],
                    "busy": bool(busy) if busy is not None else False,
                    "last_seen": last_seen,
                }
            )
        elif "online" in lower:
            rows.append({"cluster": parts[0] if parts else "", "instance": "", "busy": False, "last_seen": "", "online": True})
    return {
        "targets": rows,
        "target_online": bool(rows) or "online" in str(stdout or "").lower(),
        "target_busy": any(bool(row.get("busy")) for row in rows),
        "last_seen": rows[0].get("last_seen", "") if rows else "",
        "last_seen_age_seconds": last_seen_age_seconds(str(rows[0].get("last_seen", ""))) if rows else None,
        "gateway_mode_detected": any(row.get("cluster") and row.get("instance") for row in rows),
        "raw_excerpt_redacted": redact(str(stdout or "")),
    }


def target_preflight(doops: str, target: str, session: str, timeout: int) -> tuple[dict[str, Any], dict[str, Any]]:
    targets_result = run_doops_targets(doops, target, timeout)
    gateway = parse_doops_targets(str(targets_result.get("stdout") or ""))
    evidence = {
        "target": target,
        "session": session,
        "collected_at": now_local(),
        "targets_returncode": targets_result["returncode"],
        "targets_stdout_excerpt": redact(str(targets_result.get("stdout") or "")),
        "targets_stderr_excerpt": redact(str(targets_result.get("stderr") or "")),
        "doops_gateway": gateway,
        "target_online": bool(gateway.get("target_online")) if targets_result["returncode"] == 0 else False,
        "target_busy": bool(gateway.get("target_busy")),
        "last_seen_age_seconds": gateway.get("last_seen_age_seconds"),
        "gateway_mode_detected": bool(gateway.get("gateway_mode_detected")),
    }
    if targets_result["returncode"] != 0:
        raise RuntimeError(f"doops targets check failed for target {target}: {json.dumps(evidence, ensure_ascii=False)}")
    info_result = run_doops_info(doops, target, session, timeout)
    evidence["collector_info_status"] = "collected" if info_result["returncode"] == 0 else "failed"
    evidence["collector_info"] = sanitize_doops_result(info_result, limit=2000)
    return evidence, gateway


def remote_artifact_dir(session: str, requested: str | None = None) -> str:
    template = requested or "/root/ws/<session>/server-health"
    return template.replace("<session>", session).replace("$SESSION", session).rstrip("/")


def workspace_runner_script() -> str:
    return f"""#!/usr/bin/env bash
set +e
cd "$(dirname "$0")" || exit 2
PYTHON_BIN="${{PYTHON:-python3}}"
if ! command -v "$PYTHON_BIN" >/dev/null 2>&1; then
  PYTHON_BIN=python
fi
"$PYTHON_BIN" collector.py > collector-output.txt 2> collector-log.txt
COLLECTOR_STATUS=$?
export COLLECTOR_STATUS
"$PYTHON_BIN" - <<'PY'
import base64
import json
import os
from pathlib import Path

output = Path("collector-output.txt").read_text(encoding="utf-8", errors="replace") if Path("collector-output.txt").exists() else ""
payload = None
for line in output.splitlines():
    if "{PAYLOAD_PREFIX}" in line:
        encoded = line.split("{PAYLOAD_PREFIX}", 1)[1].strip()
        payload = json.loads(base64.b64decode(encoded).decode("utf-8"))
        break
if payload is None:
    text = output.strip()
    start = text.find("{{")
    end = text.rfind("}}")
    if start != -1 and end != -1 and end > start:
        try:
            payload = json.loads(text[start:end + 1])
        except Exception:
            payload = None
if payload is None:
    payload = {{"error": "collector-invalid-payload", "servers": []}}
Path("result.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
Path("evidence.json").write_text(json.dumps({{"remote_collector_status": int(os.environ.get("COLLECTOR_STATUS") or 0)}}, ensure_ascii=False, indent=2), encoding="utf-8")
PY
echo {START_MARKER}
cat result.json 2>/dev/null || printf '{{"error":"result-missing"}}'
echo
echo {END_MARKER}
exit "$COLLECTOR_STATUS"
"""


def create_workspace_package(package_root: Path, script: str, manifest: dict[str, Any]) -> None:
    remote_dir = package_root / "server-health"
    remote_dir.mkdir(parents=True, exist_ok=True)
    (remote_dir / "collector.py").write_text(script, encoding="utf-8")
    (remote_dir / "run.sh").write_text(workspace_runner_script(), encoding="utf-8", newline="\n")
    (remote_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")


def parse_payload_from_read(stdout: str) -> dict[str, Any] | None:
    parsed = parse_json(stdout)
    return parsed if isinstance(parsed, dict) else None


def execute_workspace_payload(
    doops: str,
    target: str,
    session: str,
    timeout: int,
    script: str,
    manifest: dict[str, Any],
    requested_remote_artifact_dir: str | None,
    keep_remote_workspace: bool,
    allow_write_fallback: bool = True,
) -> tuple[dict[str, Any], dict[str, Any]]:
    remote_dir = remote_artifact_dir(session, requested_remote_artifact_dir)
    result_path = f"{remote_dir}/result.json"
    evidence_path = f"{remote_dir}/evidence.json"
    log_path = f"{remote_dir}/collector-log.txt"
    trace: dict[str, Any] = {
        "transport_mode": "workspace-push",
        "remote_result_path": result_path,
        "cleanup_status": "skipped" if keep_remote_workspace else "pending",
    }
    with tempfile.TemporaryDirectory(prefix="server-health-doops-") as tmp:
        root = Path(tmp)
        create_workspace_package(root, script, manifest)
        push_result = run_doops_push(doops, target, session, str(root), timeout)
        trace["push"] = sanitize_doops_result(push_result, limit=1200)
        if push_result["returncode"] != 0:
            if not allow_write_fallback:
                raise RuntimeError("doops push failed")
            return execute_write_payload(
                doops,
                target,
                session,
                timeout,
                script,
                manifest,
                requested_remote_artifact_dir,
                keep_remote_workspace,
                push_result,
            )
        exec_result = run_doops_exec(doops, target, session, f"bash {shlex.quote(remote_dir + '/run.sh')}", timeout)
        trace["exec"] = sanitize_doops_result(exec_result, limit=1200)
        read_result = run_doops_read(doops, target, session, result_path, timeout)
        trace["read"] = sanitize_doops_result(read_result, limit=1200)
        payload = parse_payload_from_read(str(read_result.get("stdout") or "")) if read_result["returncode"] == 0 else None
        if payload is None:
            payload = extract_payload(str(exec_result.get("stdout") or ""))
            trace["result_source"] = "exec-stdout"
        else:
            trace["result_source"] = "doops-read"
        evidence_result = run_doops_read(doops, target, session, evidence_path, timeout)
        if evidence_result["returncode"] == 0:
            trace["remote_evidence"] = parse_json(evidence_result["stdout"]) or sanitize_doops_result(evidence_result, limit=1200)
        log_result = run_doops_read(doops, target, session, log_path, timeout)
        if log_result["returncode"] == 0:
            trace["remote_log_excerpt"] = redact(str(log_result.get("stdout") or ""), limit=1200)
    if not keep_remote_workspace:
        clean_result = run_doops_clean(doops, target, session, timeout)
        trace["clean"] = sanitize_doops_result(clean_result, limit=1200)
        trace["cleanup_status"] = "cleaned" if clean_result["returncode"] == 0 else "failed"
    return payload, trace


def execute_write_payload(
    doops: str,
    target: str,
    session: str,
    timeout: int,
    script: str,
    manifest: dict[str, Any],
    requested_remote_artifact_dir: str | None,
    keep_remote_workspace: bool,
    push_result: dict[str, Any] | None = None,
) -> tuple[dict[str, Any], dict[str, Any]]:
    remote_dir = remote_artifact_dir(session, requested_remote_artifact_dir)
    result_path = f"{remote_dir}/result.json"
    trace: dict[str, Any] = {
        "transport_mode": "file-write",
        "remote_result_path": result_path,
        "cleanup_status": "skipped" if keep_remote_workspace else "pending",
    }
    if push_result:
        trace["push"] = sanitize_doops_result(push_result, limit=1200)
    files = {
        f"{remote_dir}/collector.py": script,
        f"{remote_dir}/run.sh": workspace_runner_script(),
        f"{remote_dir}/manifest.json": json.dumps(manifest, ensure_ascii=False, indent=2),
    }
    for path, content in files.items():
        write_result = run_doops_write(doops, target, session, path, content, timeout)
        trace.setdefault("write", []).append(sanitize_doops_result(write_result, limit=1200))
        if write_result["returncode"] != 0:
            raise RuntimeError(f"doops write failed for {path}")
    exec_result = run_doops_exec(doops, target, session, f"bash {shlex.quote(remote_dir + '/run.sh')}", timeout)
    trace["exec"] = sanitize_doops_result(exec_result, limit=1200)
    read_result = run_doops_read(doops, target, session, result_path, timeout)
    trace["read"] = sanitize_doops_result(read_result, limit=1200)
    payload = parse_payload_from_read(str(read_result.get("stdout") or "")) if read_result["returncode"] == 0 else None
    if payload is None:
        payload = extract_payload(str(exec_result.get("stdout") or ""))
        trace["result_source"] = "exec-stdout"
    else:
        trace["result_source"] = "doops-read"
    if not keep_remote_workspace:
        clean_result = run_doops_clean(doops, target, session, timeout)
        trace["clean"] = sanitize_doops_result(clean_result, limit=1200)
        trace["cleanup_status"] = "cleaned" if clean_result["returncode"] == 0 else "failed"
    return payload, trace


def execute_inline_payload(
    doops: str,
    target: str,
    session: str,
    timeout: int,
    command: str,
) -> tuple[dict[str, Any], dict[str, Any]]:
    exec_result = run_doops_exec(doops, target, session, command, timeout)
    trace = {
        "transport_mode": "inline-exec",
        "exec_returncode": exec_result["returncode"],
        "exec_stderr_excerpt": redact(str(exec_result.get("stderr") or "")),
        "cleanup_status": "not-applicable",
    }
    if exec_result["returncode"] != 0:
        trace["exec_stdout_excerpt"] = redact(str(exec_result.get("stdout") or ""))
        raise RuntimeError(f"doops exec failed for target {target}: {json.dumps(trace, ensure_ascii=False)}")
    return extract_payload(str(exec_result.get("stdout") or "")), trace


def execute_payload(
    doops: str,
    target: str,
    session: str,
    timeout: int,
    inline_command: str,
    workspace_script: str,
    manifest: dict[str, Any],
    workspace_mode: str,
    keep_remote_workspace: bool,
    requested_remote_artifact_dir: str | None,
) -> tuple[dict[str, Any], dict[str, Any]]:
    if workspace_mode == "auto":
        mode = "workspace" if manifest.get("kind") in {"doops-live", "network-probe", "host-metrics", "snmp-metrics"} else "inline"
    else:
        mode = workspace_mode if workspace_mode in {"inline", "workspace"} else "inline"
    if mode == "workspace":
        try:
            return execute_workspace_payload(
                doops,
                target,
                session,
                timeout,
                workspace_script,
                manifest,
                requested_remote_artifact_dir,
                keep_remote_workspace,
            )
        except Exception as exc:
            payload, trace = execute_inline_payload(doops, target, session, timeout, inline_command)
            trace["workspace_fallback_reason"] = exc.__class__.__name__
            return payload, trace
    return execute_inline_payload(doops, target, session, timeout, inline_command)


def collect_doops_payload(
    doops: str,
    target: str,
    session: str,
    timeout: int,
    workspace_mode: str = "auto",
    keep_remote_workspace: bool = False,
    requested_remote_artifact_dir: str | None = None,
    auth_security: bool = False,
    auth_log_window_days: int = 7,
) -> tuple[dict[str, Any], dict[str, Any]]:
    evidence, _gateway = target_preflight(doops, target, session, timeout)
    remote_script = (
        self_node_script_template(auth_security)
        .replace("__AUTH_SECURITY_ENABLED__", "1" if auth_security else "0")
        .replace("__AUTH_LOG_WINDOW_DAYS__", str(max(1, int(auth_log_window_days or 7))))
    )
    payload, trace = execute_payload(
        doops,
        target,
        session,
        timeout,
        build_remote_command(auth_security, auth_log_window_days),
        remote_script,
        {"kind": "doops-live", "target": target, "session": session, "created_at": now_local(), "auth_security": bool(auth_security)},
        workspace_mode,
        keep_remote_workspace,
        requested_remote_artifact_dir,
    )
    evidence.update(trace)
    evidence["remote_payload"] = payload
    return payload, evidence


def collection_trace_from_evidence(evidence: dict[str, Any], target: str, session: str) -> dict[str, Any]:
    return {
        "doops_target": target,
        "session": session,
        "transport_mode": str(evidence.get("transport_mode") or "inline-exec"),
        "remote_result_path": str(evidence.get("remote_result_path") or ""),
        "cleanup_status": str(evidence.get("cleanup_status") or "not-applicable"),
    }


def doops_channel_from_evidence(evidence: dict[str, Any], target: str, session: str) -> dict[str, Any]:
    gateway = evidence.get("doops_gateway") if isinstance(evidence.get("doops_gateway"), dict) else {}
    return {
        "target": target,
        "session": session,
        "transport_mode": str(evidence.get("transport_mode") or "inline-exec"),
        "target_online": bool(evidence.get("target_online", True)),
        "target_busy": bool(evidence.get("target_busy", False)),
        "last_seen": str(gateway.get("last_seen") or ""),
        "last_seen_age_seconds": evidence.get("last_seen_age_seconds"),
        "gateway_mode_detected": bool(evidence.get("gateway_mode_detected")),
        "collector_info_status": str(evidence.get("collector_info_status") or "unknown"),
        "cleanup_status": str(evidence.get("cleanup_status") or "not-applicable"),
    }


def run_doops_check(
    doops: str,
    target: str,
    session: str,
    namespace: str,
    deployment: str,
    expected_image: str,
    timeout: int,
) -> dict[str, Any]:
    return run_doops_command(
        [
            doops,
            "-session",
            session,
            "check",
            "--target",
            target,
            "--namespace",
            namespace,
            "--deployment",
            deployment,
            "--image",
            expected_image,
        ],
        timeout=min(timeout, 120),
    )


def app_check_target(server: dict[str, Any], check: dict[str, Any]) -> str:
    collect = server.get("collect") if isinstance(server.get("collect"), dict) else {}
    for candidate in (
        check.get("doops_target"),
        server.get("doops_target"),
        collect.get("probe_target"),
        collect.get("metrics_probe_target"),
    ):
        if candidate:
            return str(candidate)
    auth_ref = str(server.get("auth_ref") or "")
    if auth_ref.startswith("doops:"):
        return auth_ref.split(":", 1)[1]
    return ""


def observe_application_check(doops: str, server: dict[str, Any], check: dict[str, Any], session: str, timeout: int, index: int) -> dict[str, Any]:
    check_type = str(check.get("type") or "")
    target = app_check_target(server, check)
    namespace = str(check.get("namespace") or "")
    deployment = str(check.get("deployment") or "")
    expected_image = str(check.get("expected_image") or check.get("image") or "")
    observation = {
        "type": check_type,
        "doops_target": target,
        "namespace": namespace,
        "deployment": deployment,
        "expected_image": expected_image,
        "session": f"{session}-app-{index:02d}",
        "status": "failed",
        "message": "",
    }
    if check_type != "k8s-deployment-image" or not (target and namespace and deployment and expected_image):
        observation["message"] = "application check fields incomplete"
        return observation
    result = run_doops_check(doops, target, observation["session"], namespace, deployment, expected_image, timeout)
    combined = f"{result.get('stdout') or ''}\n{result.get('stderr') or ''}"
    observation["returncode"] = result["returncode"]
    observation["duration_ms"] = result["duration_ms"]
    observation["message"] = redact(combined, limit=1500)
    if result["returncode"] == 0:
        observation["status"] = "matched"
    elif re.search(r"(?i)mismatch|not\s+match|different|image", combined):
        observation["status"] = "mismatch"
    else:
        observation["status"] = "failed"
    return observation


def apply_application_checks(doops: str, inventory: dict[str, Any], session: str, timeout: int) -> tuple[dict[str, Any], dict[str, Any]]:
    servers = inventory.get("servers") if isinstance(inventory.get("servers"), list) else []
    evidence_items: list[dict[str, Any]] = []
    for server_index, server in enumerate(servers, start=1):
        if not isinstance(server, dict):
            continue
        checks = server.get("application_checks") if isinstance(server.get("application_checks"), list) else []
        if not checks:
            continue
        observations = []
        for check_index, check in enumerate(checks, start=1):
            if not isinstance(check, dict):
                continue
            observation = observe_application_check(doops, server, check, f"{session}-{server_index:02d}", timeout, check_index)
            observations.append(observation)
            evidence_items.append({"server": str(server.get("name") or f"server-{server_index:02d}"), **observation})
        server["application_observation"] = observations
    return inventory, {"session": session, "collected_at": now_local(), "checks": evidence_items}


def source_environment(source: dict[str, Any]) -> dict[str, Any]:
    environment = source.get("environment") if isinstance(source.get("environment"), dict) else {}
    clean = {
        "name": str(environment.get("name") or "Server environment"),
        "inspection_window": now_local(),
    }
    for key in ("owner", "notes"):
        if environment.get(key):
            clean[key] = str(environment[key])
    return clean


def clean_cell(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def valid_ip(value: str) -> bool:
    try:
        ipaddress.ip_address(value)
        return True
    except Exception:
        return False


def parse_json(stdout: str) -> Any:
    text = str(stdout or "").strip()
    if not text:
        return None
    try:
        return json.loads(text)
    except Exception:
        pass
    for start_char, end_char in (("{", "}"), ("[", "]")):
        start = text.find(start_char)
        end = text.rfind(end_char)
        if start != -1 and end != -1 and end > start:
            try:
                return json.loads(text[start : end + 1])
            except Exception:
                continue
    return None


def infer_os_type(raw_os: str, auth_text: str) -> str:
    text = f"{raw_os} {auth_text}".lower()
    if any(marker in text for marker in ("windows", "win", "administrator")):
        return "windows"
    if any(marker in text for marker in ("linux", "centos", "ubuntu", "debian", "drcom", "root/")):
        return "linux"
    return "unknown"


def infer_port_obligations(name: str, role: str, os_type: str, address: str) -> list[dict[str, Any]]:
    if not valid_ip(address):
        return []
    obligations: dict[int, dict[str, Any]] = {}
    if os_type == "windows":
        obligations[3389] = {"port": 3389, "usage": "management", "source": "excel-inferred", "required": True, "label": "Windows RDP"}
    elif os_type == "linux":
        obligations[22] = {"port": 22, "usage": "management", "source": "excel-inferred", "required": True, "label": "SSH"}
    if "ftp" in f"{name} {role}".lower():
        obligations[21] = {"port": 21, "usage": "business", "source": "role-inferred", "required": True, "label": "FTP"}
    return [obligations[port] for port in sorted(obligations)]


def infer_ports(name: str, role: str, os_type: str, address: str) -> list[int]:
    return [item["port"] for item in infer_port_obligations(name, role, os_type, address)]


def limit_servers(servers: list[dict[str, Any]], limit: int | None) -> list[dict[str, Any]]:
    if limit is None or limit <= 0:
        return servers
    return servers[:limit]


def lifecycle_from_value(value: Any, field_name: str = "") -> dict[str, str]:
    text = clean_cell(value)
    field = clean_cell(field_name).lower()
    combined = f"{field} {text}".lower()
    inactive_markers = (
        "停用",
        "已停",
        "下线",
        "退役",
        "废弃",
        "注销",
        "不再使用",
        "未使用",
        "inactive",
        "disabled",
        "retired",
        "decommissioned",
        "offline",
        "否",
        "no",
        "false",
        "0",
    )
    active_markers = (
        "仍在使用",
        "在用",
        "启用",
        "使用中",
        "存活",
        "在线",
        "active",
        "enabled",
        "yes",
        "true",
        "1",
        "是",
    )
    if text:
        if any(marker in combined for marker in inactive_markers):
            return {"status": "inactive", "source": "explicit", "value": text, "reason": "清单字段标记为停用或不存活"}
        if any(marker in combined for marker in active_markers):
            return {"status": "active", "source": "explicit", "value": text, "reason": "清单字段标记为仍在使用"}
    return {"status": "active", "source": "assumed", "value": text, "reason": "清单未明确标记停用，按仍在使用纳入复核"}


def lifecycle_candidate_fields(server: dict[str, Any]) -> list[tuple[str, Any]]:
    names = (
        "lifecycle_status",
        "asset_status",
        "inventory_status",
        "usage_status",
        "server_status",
        "status",
        "state",
        "enabled",
        "disabled",
        "active",
        "decommissioned",
        "retired",
        "是否存活",
        "是否停用",
        "使用状态",
        "资产状态",
        "运行状态",
        "存活状态",
    )
    return [(name, server.get(name)) for name in names if name in server]


def ensure_lifecycle(server: dict[str, Any]) -> dict[str, Any]:
    if server.get("lifecycle_status") in {"active", "inactive"}:
        return server
    lifecycle = {"status": "active", "source": "assumed", "value": "", "reason": "清单未明确标记停用，按仍在使用纳入复核"}
    for name, value in lifecycle_candidate_fields(server):
        candidate = lifecycle_from_value(value, name)
        if candidate["source"] == "explicit":
            lifecycle = candidate
            break
        if clean_cell(value) and lifecycle["source"] != "explicit":
            lifecycle = candidate
    server["lifecycle_status"] = lifecycle["status"]
    server["lifecycle_source"] = lifecycle["source"]
    server["lifecycle_value"] = lifecycle["value"]
    server["lifecycle_reason"] = lifecycle["reason"]
    return server


def classify_source_inventory(source: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    servers = source.get("servers") if isinstance(source.get("servers"), list) else []
    active = []
    inactive = []
    for index, item in enumerate(servers, start=1):
        if not isinstance(item, dict):
            continue
        server = ensure_lifecycle(dict(item))
        server["inventory_index"] = int(server.get("inventory_index") or index)
        if server.get("lifecycle_status") == "inactive":
            inactive.append(server)
        else:
            active.append(server)
    active_source = dict(source)
    active_source["servers"] = active
    groups = {
        "total_servers": len(active) + len(inactive),
        "active_servers": len(active),
        "inactive_servers": len(inactive),
        "monitorable_servers": 0,
        "unmonitorable_servers": 0,
        "mode": "inventory-classification",
        "active": active,
        "inactive": inactive,
    }
    return active_source, groups


def strip_auth_label(value: str) -> str:
    text = clean_cell(value)
    text = re.sub(r"^(?:账号|帐号|用户名|用户|登录名|密码|口令|username|user|login|password|pass|pwd)\s*[:：=]\s*", "", text, flags=re.I)
    if any(marker in text for marker in ("：", ":")) and re.search(r"[\u4e00-\u9fff]", text.split("/", 1)[0]):
        text = re.split(r"[:：]", text)[-1]
    return text.strip()


def split_auth_usernames(username_text: str) -> list[str]:
    text = strip_auth_label(username_text)
    parts = re.split(r"[,，、;；\s]+", text)
    usernames: list[str] = []
    for part in parts:
        username = strip_auth_label(part)
        if username and username not in usernames:
            usernames.append(username)
    return usernames[:4]


def parse_runtime_auth(auth_text: str) -> dict[str, Any]:
    text = clean_cell(auth_text)
    if not text:
        return {}
    labeled_user = re.search(r"(?:账号|帐号|用户名|用户|登录名|username|user|login)\s*[:：=]\s*([^\r\n/]+)", text, flags=re.I)
    labeled_password = re.search(r"(?:密码|口令|password|pass|pwd)\s*[:：=]\s*([^\r\n]+)", text, flags=re.I)
    if labeled_user and labeled_password:
        usernames = split_auth_usernames(labeled_user.group(1))
        password = strip_auth_label(labeled_password.group(1))
        if usernames and password:
            candidates = [{"username": username, "password": password} for username in usernames]
            return {"username": candidates[0]["username"], "password": password, "candidates": candidates}
    segments = [segment.strip() for segment in re.split(r"[\r\n]+", text) if segment.strip()]
    if not segments:
        segments = [text]
    for segment in segments:
        if "/" not in segment:
            continue
        username_text, password_text = segment.split("/", 1)
        usernames = split_auth_usernames(username_text)
        password = strip_auth_label(password_text)
        if usernames and password:
            candidates = [{"username": username, "password": password} for username in usernames]
            return {"username": candidates[0]["username"], "password": password, "candidates": candidates}
    return {}


def mask_identifier(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    if len(text) <= 2:
        return text[0] + "***"
    if len(text) <= 4:
        return text[:1] + "***"
    return text[:3] + "***"


def default_auth_security(status: str = "not_enabled", window_days: int = 7, method: str = "unavailable") -> dict[str, Any]:
    return {
        "status": status,
        "window_days": int(window_days or 7),
        "method": method,
        "log_evidence": {
            "failed_login_count": 0,
            "successful_login_count": 0,
            "unique_source_count": 0,
            "top_sources": [],
            "top_users": [],
            "patterns": [],
        },
        "protection_config": {
            "password_login_enabled": None,
            "root_login_enabled": None,
            "max_auth_tries": None,
            "account_lockout_enabled": None,
            "fail2ban_enabled": None,
            "rdp_nla_enabled": None,
            "winrm_basic_enabled": None,
            "audit_policy_enabled": None,
        },
        "password_strength": {
            "status": "not_enabled",
            "weak_count": 0,
            "medium_count": 0,
            "strong_count": 0,
            "reuse_group_count": 0,
            "findings": [],
        },
    }


def _top_counts(counter: dict[str, int], key_name: str, mask: bool = False, limit: int = 5) -> list[dict[str, Any]]:
    rows = []
    for key, count in sorted(counter.items(), key=lambda item: (-item[1], item[0]))[:limit]:
        rows.append({key_name: mask_identifier(key) if mask else key, "count": count})
    return rows


def parse_linux_auth_log(text: str) -> dict[str, Any]:
    failed = 0
    success = 0
    sources: dict[str, int] = {}
    users: dict[str, int] = {}
    failed_pairs: dict[tuple[str, str], int] = {}
    successful_pairs: set[tuple[str, str]] = set()
    for raw_line in str(text or "").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        lower = line.lower()
        source = ""
        user = ""
        failed_match = re.search(r"Failed password for (?:invalid user )?(?P<user>\S+) from (?P<source>[0-9a-fA-F:.]+)", line)
        accepted_match = re.search(r"Accepted (?:password|publickey) for (?P<user>\S+) from (?P<source>[0-9a-fA-F:.]+)", line)
        pam_match = re.search(r"authentication failure;.*rhost=(?P<source>\S+).*user=(?P<user>\S+)", line)
        sudo_failure = "sudo" in lower and "authentication failure" in lower
        if failed_match:
            failed += 1
            user = failed_match.group("user")
            source = failed_match.group("source")
        elif pam_match:
            failed += 1
            user = pam_match.group("user")
            source = pam_match.group("source")
        elif sudo_failure:
            failed += 1
            user_match = re.search(r"user=(\S+)", line)
            host_match = re.search(r"rhost=(\S+)", line)
            user = user_match.group(1) if user_match else "sudo"
            source = host_match.group(1) if host_match else "local"
        elif accepted_match:
            success += 1
            user = accepted_match.group("user")
            source = accepted_match.group("source")
            successful_pairs.add((source, user))
        else:
            continue
        if source:
            sources[source] = sources.get(source, 0) + 1
        if user:
            users[user] = users.get(user, 0) + 1
        if failed_match or pam_match or sudo_failure:
            failed_pairs[(source or "unknown", user or "unknown")] = failed_pairs.get((source or "unknown", user or "unknown"), 0) + 1
    patterns = []
    for source, user in successful_pairs:
        count = failed_pairs.get((source, user), 0)
        if count:
            patterns.append({"type": "failed-then-success", "source": source, "user": mask_identifier(user), "failed_count": count})
    for source, count in sources.items():
        source_users = {user for (item_source, user), pair_count in failed_pairs.items() if item_source == source and pair_count > 0}
        if len(source_users) >= 5:
            patterns.append({"type": "source-password-spray", "source": source, "user_count": len(source_users), "failed_count": count})
    for user in {pair[1] for pair in failed_pairs}:
        user_sources = {source for (source, item_user), pair_count in failed_pairs.items() if item_user == user and pair_count > 0}
        if len(user_sources) >= 5:
            patterns.append({"type": "distributed-account-attack", "user": mask_identifier(user), "source_count": len(user_sources)})
    return {
        "failed_login_count": failed,
        "successful_login_count": success,
        "unique_source_count": len([source for source in sources if source and source != "local"]),
        "top_sources": _top_counts(sources, "source", mask=False),
        "top_users": _top_counts(users, "user", mask=True),
        "patterns": patterns[:20],
    }


def parse_windows_auth_events(text: str) -> dict[str, Any]:
    failed = success = 0
    lockouts = 0
    sources: dict[str, int] = {}
    users: dict[str, int] = {}
    for raw_line in str(text or "").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        event_match = re.search(r"\b(4625|4624|4771|4776|4740)\b", line)
        if not event_match:
            continue
        event_id = event_match.group(1)
        source_match = re.search(r"(?:Source(?: Network)? Address|IpAddress|WorkstationName)[=:]\s*([^\s,;]+)", line, flags=re.I)
        user_match = re.search(r"(?:TargetUserName|Account Name|User)[=:]\s*([^\s,;]+)", line, flags=re.I)
        source = source_match.group(1) if source_match else "unknown"
        user = user_match.group(1) if user_match else "unknown"
        if event_id in {"4625", "4771", "4776"}:
            failed += 1
        elif event_id == "4624":
            success += 1
        elif event_id == "4740":
            lockouts += 1
        sources[source] = sources.get(source, 0) + 1
        users[user] = users.get(user, 0) + 1
    patterns = []
    if failed >= 50:
        patterns.append({"type": "high-failed-login-volume", "failed_count": failed})
    if lockouts:
        patterns.append({"type": "account-lockout-observed", "count": lockouts})
    return {
        "failed_login_count": failed,
        "successful_login_count": success,
        "unique_source_count": len([source for source in sources if source and source.lower() != "unknown"]),
        "top_sources": _top_counts(sources, "source", mask=False),
        "top_users": _top_counts(users, "user", mask=True),
        "patterns": patterns[:20],
    }


def password_complexity_classes(password: str) -> int:
    return sum(
        [
            bool(re.search(r"[a-z]", password)),
            bool(re.search(r"[A-Z]", password)),
            bool(re.search(r"\d", password)),
            bool(re.search(r"[^A-Za-z0-9]", password)),
        ]
    )


def score_password_strength(password: str, username: str = "", context_terms: list[str] | None = None) -> dict[str, Any]:
    value = str(password or "")
    normalized = value.lower()
    user = str(username or "").strip().lower()
    terms = [str(item or "").strip().lower() for item in context_terms or [] if str(item or "").strip()]
    weak_terms = ["123456", "password", "admin", "qwerty", "welcome", "p@ssw0rd", "zheyin", "hdu"]
    rules = []
    if not value:
        rules.append("空密码")
    if len(value) < 10:
        rules.append("长度不足")
    if user and user in normalized:
        rules.append("包含用户名")
    if user and normalized == user:
        rules.append("与用户名相同")
    if any(term and term in normalized for term in weak_terms):
        rules.append("常见弱口令模式")
    if re.search(r"(0123|1234|2345|3456|4567|5678|6789|abcd|qwer|asdf)", normalized):
        rules.append("连续字符或键盘序列")
    if re.search(r"(.)\1{3,}", value):
        rules.append("重复字符过多")
    if re.search(r"(19|20)\d{2}$", value):
        rules.append("年份后缀")
    if any(len(term) >= 3 and term in normalized for term in terms):
        rules.append("包含服务器或业务名称")
    classes = password_complexity_classes(value)
    if classes < 3:
        rules.append("字符类型不足")
    severity = "strong"
    if rules or len(value) < 10:
        severity = "weak"
    elif len(value) < 14 or classes < 4:
        severity = "medium"
    return {"severity": severity, "rules": sorted(set(rules)) or ["未命中弱口令规则"], "length": len(value), "classes": classes}


def local_auth_candidates(auth: dict[str, Any]) -> list[dict[str, str]]:
    candidates: list[dict[str, str]] = []
    if isinstance(auth.get("candidates"), list):
        for candidate in auth.get("candidates", [])[:4]:
            if not isinstance(candidate, dict):
                continue
            username = str(candidate.get("username") or "").strip()
            password = str(candidate.get("password") or "")
            if username and password:
                candidates.append({"username": username, "password": password})
    username = str(auth.get("username") or "").strip()
    password = str(auth.get("password") or "")
    if username and password:
        candidates.insert(0, {"username": username, "password": password})
    deduped: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for candidate in candidates:
        key = (candidate["username"], candidate["password"])
        if key in seen:
            continue
        seen.add(key)
        deduped.append(candidate)
    return deduped[:4]


def auth_candidates_from_server(server: dict[str, Any]) -> list[dict[str, str]]:
    auth = server.get("runtime_auth") if isinstance(server.get("runtime_auth"), dict) else {}
    return local_auth_candidates(auth)


def source_server_lookup(source: dict[str, Any]) -> dict[str, dict[str, Any]]:
    lookup: dict[str, dict[str, Any]] = {}
    for index, server in enumerate(source.get("servers", []) if isinstance(source.get("servers"), list) else [], start=1):
        if not isinstance(server, dict):
            continue
        for key in {str(index), str(server.get("inventory_index") or index), str(server.get("address") or ""), str(server.get("name") or "")}:
            if key:
                lookup[key] = server
    return lookup


def merge_auth_security(existing: Any, window_days: int = 7) -> dict[str, Any]:
    merged = default_auth_security(window_days=window_days)
    if isinstance(existing, dict):
        merged.update({key: existing.get(key, merged.get(key)) for key in ("status", "window_days", "method")})
        if isinstance(existing.get("log_evidence"), dict):
            merged["log_evidence"].update(existing["log_evidence"])
        if isinstance(existing.get("protection_config"), dict):
            merged["protection_config"].update(existing["protection_config"])
        if isinstance(existing.get("password_strength"), dict):
            merged["password_strength"].update(existing["password_strength"])
    return merged


def apply_password_strength_inventory(
    inventory: dict[str, Any],
    source: dict[str, Any],
    enabled: bool = False,
    source_name: str = "runtime-auth",
) -> dict[str, Any]:
    inventory = dict(inventory)
    servers = [dict(server) for server in inventory.get("servers", []) if isinstance(server, dict)]
    lookup = source_server_lookup(source)
    all_passwords: dict[str, int] = {}
    if enabled:
        for source_server in lookup.values():
            for candidate in auth_candidates_from_server(source_server):
                password = candidate.get("password") or ""
                if password:
                    all_passwords[password] = all_passwords.get(password, 0) + 1
    for index, server in enumerate(servers, start=1):
        source_server = (
            lookup.get(str(server.get("inventory_index") or index))
            or lookup.get(str(server.get("address") or ""))
            or lookup.get(str(server.get("name") or ""))
            or server
        )
        if not enabled:
            if server.get("auth_security"):
                auth_security = merge_auth_security(server.get("auth_security"))
                auth_security["password_strength"] = default_auth_security()["password_strength"]
                server["auth_security"] = auth_security
            server.pop("runtime_auth", None)
            continue
        else:
            auth_security = merge_auth_security(server.get("auth_security"))
            candidates = auth_candidates_from_server(source_server)
            findings = []
            counts = {"weak": 0, "medium": 0, "strong": 0}
            reuse_groups = 0
            context_terms = [server.get("name"), server.get("role"), server.get("business_system"), server.get("address")]
            for candidate in candidates:
                password = candidate.get("password") or ""
                username = candidate.get("username") or ""
                score = score_password_strength(password, username, context_terms)
                severity = score["severity"]
                counts[severity] = counts.get(severity, 0) + 1
                reused = bool(password and all_passwords.get(password, 0) > 1)
                if reused:
                    reuse_groups += 1
                if severity in {"weak", "medium"} or reused:
                    rules = list(score["rules"])
                    if reused:
                        rules.append("多服务器或多账号复用")
                    findings.append({"account": mask_identifier(username), "severity": severity, "rule": "、".join(sorted(set(rules)))})
            auth_security["password_strength"] = {
                "status": "scored" if candidates else "unavailable",
                "weak_count": counts.get("weak", 0),
                "medium_count": counts.get("medium", 0),
                "strong_count": counts.get("strong", 0),
                "reuse_group_count": reuse_groups,
                "findings": findings[:20],
                "source": source_name,
            }
        server["auth_security"] = auth_security
        server.pop("runtime_auth", None)
    inventory["servers"] = servers
    return inventory


def apply_auth_security_boundary_inventory(inventory: dict[str, Any], enabled: bool, window_days: int = 7) -> dict[str, Any]:
    inventory = dict(inventory)
    servers = []
    for server in inventory.get("servers", []) if isinstance(inventory.get("servers"), list) else []:
        if not isinstance(server, dict):
            continue
        clean_server = dict(server)
        if enabled:
            existing = merge_auth_security(clean_server.get("auth_security"), window_days=window_days)
            if existing.get("status") == "not_enabled":
                existing = default_auth_security("unavailable", window_days=window_days, method="unavailable")
            clean_server["auth_security"] = existing
        elif clean_server.get("auth_security"):
            existing = merge_auth_security(clean_server.get("auth_security"), window_days=window_days)
            if existing.get("status") != "not_enabled":
                clean_server["auth_security"] = existing
            else:
                clean_server.pop("auth_security", None)
        clean_server.pop("runtime_auth", None)
        servers.append(clean_server)
    inventory["servers"] = servers
    return inventory


def load_excel_inventory(path: Path, limit: int | None = None, include_runtime_auth: bool = False) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise RuntimeError("openpyxl is required to read Excel inventories") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.worksheets[0]
    servers: list[dict[str, Any]] = []
    headers: list[str] = []
    for excel_row, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        if excel_row == 1:
            headers = [clean_cell(value) for value in row]
            continue
        if limit is not None and limit > 0 and len(servers) >= limit:
            break
        cells = list(row[:7])
        if len(cells) < 7:
            cells.extend([None] * (7 - len(cells)))
        if not any(clean_cell(value) for value in cells):
            continue
        name = clean_cell(cells[0]) or f"server-{excel_row:02d}"
        address = clean_cell(cells[1])
        raw_os = clean_cell(cells[2])
        auth_text = clean_cell(cells[3])
        role = clean_cell(cells[4])
        lifecycle_field = headers[6] if len(headers) > 6 and headers[6] else "是否存活"
        lifecycle = lifecycle_from_value(cells[6], lifecycle_field)
        os_type = infer_os_type(raw_os, auth_text)
        port_obligations = infer_port_obligations(name, role, os_type, address)
        ports = [item["port"] for item in port_obligations]
        server = {
                "name": name,
                "address": address if valid_ip(address) else "",
                "raw_address": address,
                "os_type": os_type,
                "role": role,
                "owner": "",
                "ports": ports,
                "port_obligations": port_obligations,
                "inventory_index": excel_row - 1,
                "lifecycle_status": lifecycle["status"],
                "lifecycle_source": lifecycle["source"],
                "lifecycle_value": lifecycle["value"],
                "lifecycle_reason": lifecycle["reason"],
                "auth_ref": f"excel-row-{excel_row}-credential-redacted" if auth_text else "",
                "collect": {"network": True, "ssh": False, "winrm": False},
                "notes": f"Excel row {excel_row}; raw_os={raw_os or 'unknown'}; credential-redacted.",
            }
        if include_runtime_auth:
            runtime_auth = parse_runtime_auth(auth_text)
            if runtime_auth:
                server["runtime_auth"] = runtime_auth
        servers.append(server)
    return {
        "environment": {
            "name": path.stem,
            "inspection_window": now_local(),
            "notes": "Excel inventory imported with credential values redacted.",
        },
        "servers": servers,
    }


def load_source_inventory(path: Path, limit: int | None = None, include_runtime_auth: bool = False) -> dict[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(f"inventory not found: {path}")
    if path.suffix.lower() == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data, list):
            return {"environment": {"name": path.stem}, "servers": limit_servers(data, limit)}
        if isinstance(data, dict) and isinstance(data.get("servers"), list):
            data = dict(data)
            data["servers"] = limit_servers(data["servers"], limit)
            return data
        raise ValueError("JSON inventory must be an object with servers or an array")
    if path.suffix.lower() in {".csv", ".tsv"}:
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle, delimiter=delimiter))
        return {"environment": {"name": path.stem}, "servers": limit_servers(rows, limit)}
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return load_excel_inventory(path, limit, include_runtime_auth=include_runtime_auth)
    raise ValueError("inventory must be .json, .csv, .tsv, .xlsx, or .xlsm")


def doops_target_for_server(server: dict[str, Any]) -> str:
    collect = server.get("collect") if isinstance(server.get("collect"), dict) else {}
    candidates = [
        server.get("doops_target"),
        server.get("doopsTarget"),
        server.get("target"),
        collect.get("doops_target"),
        collect.get("target"),
    ]
    auth_ref = str(server.get("auth_ref") or "")
    if auth_ref.startswith("doops:"):
        candidates.append(auth_ref.split(":", 1)[1])
    if collect.get("doops") and server.get("address"):
        candidates.append(server.get("address"))
    for candidate in candidates:
        if candidate:
            return str(candidate).strip()
    raise ValueError(f"server {server.get('name') or '<unnamed>'} is missing doops_target")


def build_server(
    target: str,
    session: str,
    payload: dict[str, Any],
    source_server: dict[str, Any] | None = None,
    evidence: dict[str, Any] | None = None,
) -> dict[str, Any]:
    source_server = source_server or {}
    hostname = str(payload.get("hostname") or target)
    name = str(source_server.get("name") or hostname)
    role = str(source_server.get("role") or f"doops target {target}")
    owner = str(source_server.get("owner") or "")
    server = {
        "name": name,
        "address": target,
        "os_type": str(source_server.get("os_type") or payload.get("os_type") or "linux").lower(),
        "role": role,
        "owner": owner,
        "ports": source_server.get("ports", []),
        "port_obligations": source_server.get("port_obligations", []) if isinstance(source_server.get("port_obligations"), list) else [],
        "application_checks": source_server.get("application_checks", []) if isinstance(source_server.get("application_checks"), list) else [],
        "application_observation": source_server.get("application_observation", []) if isinstance(source_server.get("application_observation"), list) else [],
        "collection_trace": collection_trace_from_evidence(evidence or {}, target, session)
        if evidence
        else source_server.get("collection_trace", {}) if isinstance(source_server.get("collection_trace"), dict) else {},
        "auth_ref": f"doops:{target}",
        "collect": {"doops": True, "network": False, "self_node_metrics": True, "doops_target": target},
        "metrics": payload.get("metrics") if isinstance(payload.get("metrics"), dict) else {},
        "services": payload.get("services") if isinstance(payload.get("services"), list) else [],
        "processes": payload.get("processes") if isinstance(payload.get("processes"), list) else [],
        "notes": f"Remote hostname: {hostname}; Platform: {payload.get('platform') or 'unknown'}; collected_at: {payload.get('collected_at') or 'unknown'}; session: {session}",
    }
    auth_security = merge_auth_security(payload.get("auth_security")) if payload.get("auth_security") else None
    if auth_security and auth_security.get("status") != "not_enabled":
        server["auth_security"] = auth_security
    return server


def build_inventory(environment: dict[str, Any], target: str, session: str, payload: dict[str, Any]) -> dict[str, Any]:
    display_name = str(environment.get("display_name") or target)
    return {
        "environment": {
            "name": f"{display_name} - doops live inspection",
            "inspection_window": now_local(),
            "notes": f"Read-only evidence collected through doops target {target} with session {session}.",
        },
        "servers": [build_server(target, session, payload)],
    }


def collect_inventory(
    doops: str,
    source: dict[str, Any],
    session: str,
    timeout: int,
    workspace_mode: str = "auto",
    keep_remote_workspace: bool = False,
    requested_remote_artifact_dir: str | None = None,
    auth_security: bool = False,
    auth_log_window_days: int = 7,
) -> tuple[dict[str, Any], dict[str, Any]]:
    servers = source.get("servers") if isinstance(source.get("servers"), list) else []
    if not servers:
        raise ValueError("inventory must contain at least one server")
    collected_servers = []
    evidence_items = []
    for index, raw_server in enumerate(servers, start=1):
        if not isinstance(raw_server, dict):
            raise ValueError(f"server #{index} must be an object")
        target = doops_target_for_server(raw_server)
        payload, evidence = collect_doops_payload(
            doops,
            target,
            f"{session}-{index:02d}",
            timeout,
            workspace_mode,
            keep_remote_workspace,
            requested_remote_artifact_dir,
            auth_security,
            auth_log_window_days,
        )
        collected_servers.append(build_server(target, f"{session}-{index:02d}", payload, raw_server, evidence))
        evidence["source_server"] = {
            "name": str(raw_server.get("name") or f"server-{index:02d}"),
            "role": str(raw_server.get("role") or ""),
            "owner": str(raw_server.get("owner") or ""),
        }
        evidence_items.append(evidence)
    environment = source_environment(source)
    environment["notes"] = (
        (environment.get("notes") + " " if environment.get("notes") else "")
        + f"Read-only evidence collected through doops for {len(collected_servers)} supplied server(s) with session {session}."
    )
    inventory = {"environment": environment, "servers": collected_servers}
    if evidence_items:
        inventory["doops_channel"] = doops_channel_from_evidence(evidence_items[0], str(evidence_items[0].get("target") or ""), session)
    return inventory, {
        "session": session,
        "collected_at": now_local(),
        "server_count": len(collected_servers),
        "servers": evidence_items,
    }


def build_probe_script(servers: list[dict[str, Any]], probe_target: str, probe_timeout: float) -> str:
    probe_items = [
        {
            "index": int(server.get("probe_index") or index),
            "name": str(server.get("name") or f"server-{index:02d}"),
            "address": str(server.get("address") or ""),
            "ports": [int(port) for port in server.get("ports", []) if str(port).isdigit()],
        }
        for index, server in enumerate(servers, start=1)
    ]
    servers_b64 = base64.b64encode(json.dumps(probe_items, ensure_ascii=False).encode("utf-8")).decode("ascii")
    return (
        REMOTE_PROBE_PYTHON.replace("__SERVERS_B64__", servers_b64)
        .replace("__PROBE_TARGET__", probe_target)
        .replace("__PROBE_TIMEOUT__", str(probe_timeout))
    )


def build_probe_command(servers: list[dict[str, Any]], probe_target: str, probe_timeout: float) -> str:
    script = build_probe_script(servers, probe_target, probe_timeout)
    script_b64 = base64.b64encode(script.encode("utf-8")).decode("ascii")
    return " ".join(
        [
            "set +e;",
            f"echo {START_MARKER};",
            f"printf {script_b64} | base64 -d | python3;",
            f"echo {END_MARKER}",
        ]
    )


def chunk_probe_servers(
    servers: list[dict[str, Any]],
    probe_target: str,
    probe_timeout: float,
    max_command_length: int = 7600,
) -> list[list[dict[str, Any]]]:
    chunks: list[list[dict[str, Any]]] = []
    current: list[dict[str, Any]] = []
    for server in servers:
        candidate = current + [server]
        if current and len(build_probe_command(candidate, probe_target, probe_timeout)) > max_command_length:
            chunks.append(current)
            current = [server]
        else:
            current = candidate
    if current:
        chunks.append(current)
    return chunks


def collect_probe_payload(
    doops: str,
    probe_target: str,
    servers: list[dict[str, Any]],
    session: str,
    timeout: int,
    probe_timeout: float,
    workspace_mode: str = "auto",
    keep_remote_workspace: bool = False,
    requested_remote_artifact_dir: str | None = None,
) -> tuple[dict[str, Any], dict[str, Any]]:
    evidence, _gateway = target_preflight(doops, probe_target, session, timeout)
    evidence.update({"probe_target": probe_target, "server_count": len(servers)})
    payload, trace = execute_payload(
        doops,
        probe_target,
        session,
        timeout,
        build_probe_command(servers, probe_target, probe_timeout),
        build_probe_script(servers, probe_target, probe_timeout),
        {"kind": "network-probe", "target": probe_target, "session": session, "server_count": len(servers), "created_at": now_local()},
        workspace_mode,
        keep_remote_workspace,
        requested_remote_artifact_dir,
    )
    evidence.update(trace)
    evidence["remote_payload"] = payload
    return payload, evidence


def collect_probe_inventory(
    doops: str,
    source: dict[str, Any],
    probe_target: str,
    session: str,
    timeout: int,
    probe_timeout: float,
    workspace_mode: str = "auto",
    keep_remote_workspace: bool = False,
    requested_remote_artifact_dir: str | None = None,
) -> tuple[dict[str, Any], dict[str, Any]]:
    servers = source.get("servers") if isinstance(source.get("servers"), list) else []
    if not servers:
        raise ValueError("inventory must contain at least one server")
    indexed_servers = []
    for index, server in enumerate(servers, start=1):
        if not isinstance(server, dict):
            raise ValueError(f"server #{index} must be an object")
        indexed = dict(server)
        indexed["probe_index"] = index
        indexed_servers.append(indexed)
    payload_servers: list[dict[str, Any]] = []
    chunk_evidence = []
    for chunk_index, chunk in enumerate(chunk_probe_servers(indexed_servers, probe_target, probe_timeout), start=1):
        chunk_payload, chunk_item = collect_probe_payload(
            doops,
            probe_target,
            chunk,
            f"{session}-probe-{chunk_index:02d}",
            timeout,
            probe_timeout,
            workspace_mode,
            keep_remote_workspace,
            requested_remote_artifact_dir,
        )
        payload_servers.extend(item for item in chunk_payload.get("servers", []) if isinstance(item, dict))
        chunk_evidence.append(chunk_item)
    payload = {
        "probe_target": probe_target,
        "checked_at": now_local(),
        "servers": payload_servers,
    }
    evidence = {
        "probe_target": probe_target,
        "session": session,
        "collected_at": now_local(),
        "server_count": len(servers),
        "chunks": chunk_evidence,
        "remote_payload": payload,
    }
    observations = {
        int(item.get("index")): item
        for item in payload.get("servers", [])
        if isinstance(item, dict) and str(item.get("index") or "").isdigit()
    }
    collected_servers = []
    for index, raw_server in enumerate(servers, start=1):
        if not isinstance(raw_server, dict):
            raise ValueError(f"server #{index} must be an object")
        observation = observations.get(
            index,
            {
                "index": index,
                "address": raw_server.get("address") or "",
                "reachable": False,
                "ping": {"status": "unknown", "latency_ms": None, "error": "missing-observation"},
                "ports": [],
            },
        )
        collect = raw_server.get("collect") if isinstance(raw_server.get("collect"), dict) else {}
        clean_server = {
            "name": str(raw_server.get("name") or f"server-{index:02d}"),
            "address": str(raw_server.get("address") or ""),
            "os_type": str(raw_server.get("os_type") or raw_server.get("os") or "unknown").lower(),
            "role": str(raw_server.get("role") or ""),
            "owner": str(raw_server.get("owner") or ""),
            "department": str(raw_server.get("department") or ""),
            "business_system": str(raw_server.get("business_system") or ""),
            "business_domain": str(raw_server.get("business_domain") or ""),
            "user_group": str(raw_server.get("user_group") or ""),
            "criticality": str(raw_server.get("criticality") or ""),
            "has_redundancy": raw_server.get("has_redundancy"),
            "impact_note": str(raw_server.get("impact_note") or ""),
            "inventory_index": raw_server.get("inventory_index") or index,
            "lifecycle_status": str(raw_server.get("lifecycle_status") or "active"),
            "lifecycle_source": str(raw_server.get("lifecycle_source") or "assumed"),
            "lifecycle_value": str(raw_server.get("lifecycle_value") or ""),
            "lifecycle_reason": str(raw_server.get("lifecycle_reason") or ""),
            "ports": raw_server.get("ports", []),
            "port_obligations": raw_server.get("port_obligations", []) if isinstance(raw_server.get("port_obligations", []), list) else [],
            "application_checks": raw_server.get("application_checks", []) if isinstance(raw_server.get("application_checks"), list) else [],
            "application_observation": raw_server.get("application_observation", []) if isinstance(raw_server.get("application_observation"), list) else [],
            "auth_ref": raw_server.get("auth_ref") or "",
            "collect": {**collect, "doops_probe": True, "probe_target": probe_target, "network": False, "ssh": False, "winrm": False},
            "metrics": raw_server.get("metrics", {}) if isinstance(raw_server.get("metrics"), dict) else {},
            "services": raw_server.get("services", []) if isinstance(raw_server.get("services"), list) else [],
            "processes": raw_server.get("processes", []) if isinstance(raw_server.get("processes"), list) else [],
            "auth_security": raw_server.get("auth_security", {}) if isinstance(raw_server.get("auth_security"), dict) else {},
            "network_observation": {"mode": "doops-network-probe", "probe_target": probe_target, **observation},
            "collection_trace": collection_trace_from_evidence(chunk_evidence[0] if chunk_evidence else {}, probe_target, session),
            "notes": f"{raw_server.get('notes') or ''} doops network probe through {probe_target} session {session}.".strip(),
        }
        if isinstance(raw_server.get("runtime_auth"), dict):
            clean_server["runtime_auth"] = raw_server["runtime_auth"]
        collected_servers.append(clean_server)
    environment = source_environment(source)
    environment["notes"] = (
        (environment.get("notes") + " " if environment.get("notes") else "")
        + f"Read-only network evidence collected from doops target {probe_target} for {len(collected_servers)} server(s)."
    )
    inventory = {"environment": environment, "servers": collected_servers}
    if chunk_evidence:
        inventory["doops_channel"] = doops_channel_from_evidence(chunk_evidence[0], probe_target, session)
    return inventory, evidence


def build_host_metrics_command(
    servers: list[dict[str, Any]],
    metrics_target: str,
    profile: str,
    metrics_timeout: int,
    auth_security: bool = False,
    auth_log_window_days: int = 7,
) -> str:
    script = build_host_metrics_script(servers, metrics_target, profile, metrics_timeout, auth_security, auth_log_window_days)
    script_b64 = base64.b64encode(gzip.compress(script.encode("utf-8"))).decode("ascii")
    return " ".join(
        [
            "set +e;",
            f"echo {START_MARKER};",
            f"HOST_METRICS_PROFILE={shlex.quote(profile)};",
            f"if command -v python3 >/dev/null 2>&1; then printf {script_b64} | base64 -d | gzip -dc | python3; else printf '%s\\n' SERVER_HEALTH_B64:eyJlcnJvciI6ImNvbGxlY3Rvci1weXRob24tdW5hdmFpbGFibGUiLCJzZXJ2ZXJzIjpbXX0=; fi;",
            f"echo {END_MARKER}",
        ]
    )


def build_host_metrics_script(
    servers: list[dict[str, Any]],
    metrics_target: str,
    profile: str,
    metrics_timeout: int,
    auth_security: bool = False,
    auth_log_window_days: int = 7,
) -> str:
    metric_items = [
        {
            "index": int(server.get("metrics_index") or index),
            "name": str(server.get("name") or f"server-{index:02d}"),
            "address": str(server.get("address") or ""),
            "os_type": str(server.get("os_type") or "unknown").lower(),
            "ports": [int(port) for port in server.get("ports", []) if str(port).isdigit()],
            "runtime_auth": server.get("runtime_auth") if isinstance(server.get("runtime_auth"), dict) else {},
        }
        for index, server in enumerate(servers, start=1)
    ]
    servers_b64 = base64.b64encode(json.dumps(metric_items, ensure_ascii=False).encode("utf-8")).decode("ascii")
    methods = {local_metrics_method(item) for item in metric_items}
    template = remote_host_metrics_template(methods, auth_security)
    script = (
        template.replace("__SERVERS_B64__", servers_b64)
        .replace("__HOST_METRICS_TARGET__", metrics_target)
        .replace("__HOST_METRICS_PROFILE__", profile)
        .replace("__METRICS_TIMEOUT__", str(metrics_timeout))
        .replace("__AUTH_SECURITY_ENABLED__", "1" if auth_security else "0")
        .replace("__AUTH_LOG_WINDOW_DAYS__", str(max(1, int(auth_log_window_days or 7))))
    )
    return script


def local_metrics_method(server: dict[str, Any]) -> str:
    os_type = str(server.get("os_type") or "").lower()
    ports = {int(port) for port in server.get("ports", []) if str(port).isdigit()}
    if os_type.startswith("win"):
        return "winrm"
    if os_type.startswith("linux") or 22 in ports:
        return "ssh"
    if {3389, 5985, 5986} & ports:
        return "winrm"
    return "not-supported"


def remove_between(text: str, start: str, end: str) -> str:
    start_index = text.find(start)
    end_index = text.find(end)
    if start_index == -1 or end_index == -1 or end_index <= start_index:
        return text
    return text[:start_index] + text[end_index:]


def strip_host_auth_security(script: str) -> str:
    compact_windows_metrics = (
        "WINDOWS_METRICS_PS = r'''$os=Get-CimInstance Win32_OperatingSystem;"
        "$cpu=(Get-CimInstance Win32_Processor|Measure-Object LoadPercentage -Average).Average;"
        "$t=[double]$os.TotalVisibleMemorySize;$f=[double]$os.FreePhysicalMemory;"
        "$d=Get-CimInstance Win32_LogicalDisk -Filter 'DriveType=3'|%{$s=[double]$_.Size;$fr=[double]$_.FreeSpace;@{mount=$_.DeviceID;used_percent=[math]::Round((($s-$fr)*100/$s),1)}};"
        "$svc=Get-Service WinRM -EA SilentlyContinue;$p=@{};"
        "Get-Process|Sort-Object CPU -Descending|Select-Object -First 8|%{$p[$_.Id]=$null;$p[$_.Id]=$_};"
        "Get-Process|Sort-Object WorkingSet64 -Descending|Select-Object -First 8|%{$p[$_.Id]=$null;$p[$_.Id]=$_};"
        "$procs=$p.Values|Sort-Object CPU -Descending|Select-Object -First 12|%{@{pid=$_.Id;name=$_.ProcessName;state='running';cpu_seconds=[math]::Round([double]($_.CPU),1);memory_mb=[math]::Round([double]($_.WorkingSet64)/1MB,1)}};"
        "@{metrics=@{cpu_percent=$cpu;memory_percent=[math]::Round((($t-$f)*100/$t),1);uptime_days=[math]::Round(((Get-Date)-$os.LastBootUpTime).TotalDays,2);disks=$d};services=@(@{name='WinRM';status=if($svc.Status -eq 'Running'){'running'}else{''+$svc.Status};expected='running'});processes=$procs}|ConvertTo-Json -Depth 6 -Compress'''"
    )
    windows_start = script.find("WINDOWS_METRICS_PS = r'''")
    windows_end = script.find("\n\n\ndef valid_ip", windows_start)
    if windows_start != -1 and windows_end != -1:
        script = script[:windows_start] + compact_windows_metrics + script[windows_end:]
    script = script.replace(
        'AUTH_SECURITY_ENABLED = "__AUTH_SECURITY_ENABLED__" == "1"\nAUTH_LOG_WINDOW_DAYS = int("__AUTH_LOG_WINDOW_DAYS__")\n',
        "",
    )
    ps_start = script.find("\n# AUTH_SECURITY_START")
    ps_end = script.find("\n# AUTH_SECURITY_END", ps_start)
    if ps_start != -1 and ps_end != -1:
        ps_end_line = script.find("\n", ps_end + 1)
        script = script[:ps_start] + (script[ps_end_line:] if ps_end_line != -1 else "")
    script = script.replace(";auth_security=$auth", "")
    start_index = script.find('\nauth_security=\'{"status":"not_enabled"')
    end_index = script.find('\n[ -n "$memory_percent" ]', start_index)
    if start_index != -1 and end_index != -1:
        script = script[:start_index] + script[end_index:]
    lines = []
    for line in script.splitlines():
        if line.startswith("printf '{\"metrics\"") and '"auth_security":%s' in line:
            line = line.replace(',"auth_security":%s', "").replace(' "$auth_security"', "")
        lines.append(line)
    return "\n".join(lines)


def remote_host_metrics_template(methods: set[str], auth_security: bool = False) -> str:
    script = REMOTE_HOST_METRICS_PYTHON
    if not auth_security:
        script = strip_host_auth_security(script)
    if "ssh" in methods:
        script = remove_between(script, "LINUX_REMOTE_PYTHON = r'''", "\n\nLINUX_REMOTE_SH_FALLBACK = r'''")
    if methods and methods <= {"ssh", "not-supported"}:
        script = script.replace("import http.client\n", "")
        script = script.replace("import socket\n", "")
        script = remove_between(script, "WINDOWS_METRICS_PS = r'''", "\n\n\ndef valid_ip")
        script = remove_between(script, "def xml_unescape", "\n\ndef collect_ssh")
        script = remove_between(script, "def collect_winrm", "\n\ndef main")
    elif methods and methods <= {"winrm", "not-supported"}:
        script = remove_between(script, "LINUX_REMOTE_PYTHON = r'''", "WINDOWS_METRICS_PS = r'''")
        script = remove_between(script, "def run_ssh_password", "\n\ndef parse_json")
        script = remove_between(script, "def collect_ssh", "\n\ndef collect_winrm")
    return script


def chunk_host_metrics_servers(
    servers: list[dict[str, Any]],
    metrics_target: str,
    profile: str,
    metrics_timeout: int,
    max_command_length: int = 7200,
) -> list[list[dict[str, Any]]]:
    chunks: list[list[dict[str, Any]]] = []
    for server in servers:
        chunks.append([server])
    return chunks


def build_snmp_metrics_command(
    servers: list[dict[str, Any]],
    snmp_target: str,
    profile: str,
    snmp_timeout: int,
) -> str:
    script = build_snmp_metrics_script(servers, snmp_target, profile, snmp_timeout)
    script_b64 = base64.b64encode(gzip.compress(script.encode("utf-8"))).decode("ascii")
    return " ".join(
        [
            "set +e;",
            f"echo {START_MARKER};",
            f"SNMP_METRICS_PROFILE={shlex.quote(profile)};",
            f"if command -v python3 >/dev/null 2>&1; then printf {script_b64} | base64 -d | gzip -dc | python3; else printf '%s\\n' SERVER_HEALTH_B64:eyJlcnJvciI6ImNvbGxlY3Rvci1weXRob24tdW5hdmFpbGFibGUiLCJzZXJ2ZXJzIjpbXX0=; fi;",
            f"echo {END_MARKER}",
        ]
    )


def build_snmp_metrics_script(
    servers: list[dict[str, Any]],
    snmp_target: str,
    profile: str,
    snmp_timeout: int,
) -> str:
    metric_items = [
        {
            "index": int(server.get("metrics_index") or index),
            "name": str(server.get("name") or f"server-{index:02d}"),
            "address": str(server.get("address") or ""),
            "os_type": str(server.get("os_type") or "unknown").lower(),
            "ports": [int(port) for port in server.get("ports", []) if str(port).isdigit()],
        }
        for index, server in enumerate(servers, start=1)
    ]
    servers_b64 = base64.b64encode(json.dumps(metric_items, ensure_ascii=False).encode("utf-8")).decode("ascii")
    script = (
        REMOTE_SNMP_METRICS_PYTHON.replace("__SERVERS_B64__", servers_b64)
        .replace("__SNMP_TARGET__", snmp_target)
        .replace("__SNMP_PROFILE__", profile)
        .replace("__SNMP_TIMEOUT__", str(snmp_timeout))
    )
    return script


def chunk_snmp_metrics_servers(
    servers: list[dict[str, Any]],
    snmp_target: str,
    profile: str,
    snmp_timeout: int,
    max_command_length: int = 7200,
) -> list[list[dict[str, Any]]]:
    chunks: list[list[dict[str, Any]]] = []
    for server in servers:
        chunks.append([server])
    return chunks


def sanitize_collection_errors(errors: Any) -> list[dict[str, str]]:
    clean_errors = []
    if not isinstance(errors, list):
        return clean_errors
    for item in errors:
        if not isinstance(item, dict):
            continue
        stage = re.sub(r"[^A-Za-z0-9_.-]+", "-", str(item.get("stage") or "host-metrics")).strip("-") or "host-metrics"
        error_type = re.sub(r"[^A-Za-z0-9_.-]+", "-", str(item.get("type") or item.get("error") or "unknown")).strip("-") or "unknown"
        clean_errors.append({"stage": stage[:80], "type": error_type[:120]})
    return clean_errors


def sanitize_metric_payload(item: dict[str, Any]) -> dict[str, Any]:
    metrics = item.get("metrics") if isinstance(item.get("metrics"), dict) else {}
    clean_metrics: dict[str, Any] = {}
    for key in ("cpu_percent", "memory_percent", "load_per_core", "uptime_days"):
        value = metrics.get(key)
        if isinstance(value, (int, float)):
            clean_metrics[key] = value
    disks = []
    for disk in metrics.get("disks", []) if isinstance(metrics.get("disks", []), list) else []:
        if not isinstance(disk, dict):
            continue
        used = disk.get("used_percent")
        if isinstance(used, (int, float)):
            disks.append({"mount": str(disk.get("mount") or disk.get("name") or "disk")[:80], "used_percent": used})
    if disks:
        clean_metrics["disks"] = disks[:30]

    services = []
    for service in item.get("services", []) if isinstance(item.get("services", []), list) else []:
        if not isinstance(service, dict):
            continue
        services.append(
            {
                "name": str(service.get("name") or "service")[:120],
                "status": str(service.get("status") or service.get("state") or "unknown")[:80],
                "expected": str(service.get("expected") or "running")[:80],
            }
        )
    processes = []
    for process in item.get("processes", []) if isinstance(item.get("processes", []), list) else []:
        if not isinstance(process, dict):
            continue
        clean_process = {
            "name": str(process.get("name") or "process")[:80],
            "state": str(process.get("state") or "unknown")[:32],
        }
        pid = process.get("pid")
        if isinstance(pid, int):
            clean_process["pid"] = pid
        if process.get("user") is not None:
            clean_process["user"] = str(process.get("user") or "")[:32]
        for key in ("cpu_percent", "memory_percent", "cpu_seconds", "memory_mb"):
            value = process.get(key)
            if isinstance(value, (int, float)):
                clean_process[key] = value
        processes.append(clean_process)
    clean_payload = {
        "metrics": clean_metrics,
        "services": services[:50],
        "processes": processes[:30],
        "collection_errors": sanitize_collection_errors(item.get("collection_errors")),
        "status": str(item.get("status") or "failed"),
        "metrics_method": str(item.get("metrics_method") or "not-supported"),
    }
    auth_security = merge_auth_security(item.get("auth_security")) if item.get("auth_security") else None
    if auth_security and auth_security.get("status") != "not_enabled":
        clean_payload["auth_security"] = auth_security
    return clean_payload


def base_inventory_server(raw_server: dict[str, Any], index: int) -> dict[str, Any]:
    server = {
        "name": str(raw_server.get("name") or f"server-{index:02d}"),
        "address": str(raw_server.get("address") or ""),
        "raw_address": str(raw_server.get("raw_address") or raw_server.get("address") or ""),
        "os_type": str(raw_server.get("os_type") or raw_server.get("os") or "unknown").lower(),
        "role": str(raw_server.get("role") or ""),
        "owner": str(raw_server.get("owner") or ""),
        "department": str(raw_server.get("department") or ""),
        "business_system": str(raw_server.get("business_system") or ""),
        "business_domain": str(raw_server.get("business_domain") or ""),
        "user_group": str(raw_server.get("user_group") or ""),
        "criticality": str(raw_server.get("criticality") or ""),
        "has_redundancy": raw_server.get("has_redundancy"),
        "impact_note": str(raw_server.get("impact_note") or ""),
        "inventory_index": raw_server.get("inventory_index") or index,
        "lifecycle_status": str(raw_server.get("lifecycle_status") or "active"),
        "lifecycle_source": str(raw_server.get("lifecycle_source") or "assumed"),
        "lifecycle_value": str(raw_server.get("lifecycle_value") or ""),
        "lifecycle_reason": str(raw_server.get("lifecycle_reason") or ""),
        "ports": raw_server.get("ports", []),
        "port_obligations": raw_server.get("port_obligations", []) if isinstance(raw_server.get("port_obligations", []), list) else [],
        "application_checks": raw_server.get("application_checks", []) if isinstance(raw_server.get("application_checks"), list) else [],
        "application_observation": raw_server.get("application_observation", []) if isinstance(raw_server.get("application_observation"), list) else [],
        "collection_trace": raw_server.get("collection_trace", {}) if isinstance(raw_server.get("collection_trace"), dict) else {},
        "auth_ref": raw_server.get("auth_ref") or "",
        "collect": raw_server.get("collect") if isinstance(raw_server.get("collect"), dict) else {},
        "metrics": raw_server.get("metrics", {}) if isinstance(raw_server.get("metrics"), dict) else {},
        "services": raw_server.get("services", []) if isinstance(raw_server.get("services"), list) else [],
        "processes": raw_server.get("processes", []) if isinstance(raw_server.get("processes"), list) else [],
        "auth_security": raw_server.get("auth_security", {}) if isinstance(raw_server.get("auth_security"), dict) else {},
        "network_observation": raw_server.get("network_observation", {}) if isinstance(raw_server.get("network_observation"), dict) else {},
        "notes": str(raw_server.get("notes") or ""),
    }
    if isinstance(raw_server.get("port_obligations"), list):
        server["port_obligations"] = raw_server["port_obligations"]
    return server


def metric_monitorable(server: dict[str, Any]) -> bool:
    collect = server.get("collect") if isinstance(server.get("collect"), dict) else {}
    metrics = server.get("metrics") if isinstance(server.get("metrics"), dict) else {}
    if collect.get("host_metrics_status") != "collected":
        return False
    return any(isinstance(metrics.get(key), (int, float)) for key in ("cpu_percent", "memory_percent", "load_per_core", "uptime_days"))


def first_collection_error(server: dict[str, Any]) -> str:
    errors = server.get("collection_errors") if isinstance(server.get("collection_errors"), list) else []
    for item in errors:
        if isinstance(item, dict):
            return str(item.get("type") or item.get("error") or "")
    return ""


def classification_csv_row(server: dict[str, Any], group: str) -> dict[str, Any]:
    collect = server.get("collect") if isinstance(server.get("collect"), dict) else {}
    metrics = server.get("metrics") if isinstance(server.get("metrics"), dict) else {}
    return {
        "group": group,
        "inventory_index": server.get("inventory_index") or "",
        "name": server.get("name") or "",
        "address": server.get("address") or server.get("raw_address") or "",
        "os_type": server.get("os_type") or server.get("os") or "",
        "role": server.get("role") or "",
        "owner": server.get("owner") or "",
        "department": server.get("department") or "",
        "business_system": server.get("business_system") or "",
        "lifecycle_status": server.get("lifecycle_status") or "",
        "lifecycle_source": server.get("lifecycle_source") or "",
        "lifecycle_value": server.get("lifecycle_value") or "",
        "lifecycle_reason": server.get("lifecycle_reason") or "",
        "host_metrics_status": collect.get("host_metrics_status") or "",
        "metrics_method": collect.get("metrics_method") or "",
        "cpu_percent": metrics.get("cpu_percent", ""),
        "memory_percent": metrics.get("memory_percent", ""),
        "error": first_collection_error(server),
    }


def write_group_csv(path: Path, rows: list[dict[str, Any]], group: str) -> None:
    fieldnames = [
        "group",
        "inventory_index",
        "name",
        "address",
        "os_type",
        "role",
        "owner",
        "department",
        "business_system",
        "lifecycle_status",
        "lifecycle_source",
        "lifecycle_value",
        "lifecycle_reason",
        "host_metrics_status",
        "metrics_method",
        "cpu_percent",
        "memory_percent",
        "error",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for server in rows:
            writer.writerow(classification_csv_row(server, group))


def finalize_inventory_classification(
    inventory: dict[str, Any],
    groups: dict[str, Any],
    out_path: Path,
    metrics_attempted: bool,
) -> dict[str, Any]:
    active_servers = inventory.get("servers") if isinstance(inventory.get("servers"), list) else []
    monitorable = [server for server in active_servers if isinstance(server, dict) and metric_monitorable(server)]
    unmonitorable = [server for server in active_servers if isinstance(server, dict) and not metric_monitorable(server)]
    groups = dict(groups)
    groups["monitorable"] = monitorable if metrics_attempted else []
    groups["unmonitorable"] = unmonitorable if metrics_attempted else active_servers
    groups["monitorable_servers"] = len(groups["monitorable"])
    groups["unmonitorable_servers"] = len(groups["unmonitorable"])
    groups["metric_classification_attempted"] = bool(metrics_attempted)
    csv_dir = out_path.parent
    csv_dir.mkdir(parents=True, exist_ok=True)
    files = {
        "active_servers_csv": str((csv_dir / "inventory-active-servers.csv").resolve()),
        "inactive_servers_csv": str((csv_dir / "inventory-inactive-servers.csv").resolve()),
        "monitorable_servers_csv": str((csv_dir / "inventory-monitorable-servers.csv").resolve()),
        "unmonitorable_servers_csv": str((csv_dir / "inventory-unmonitorable-servers.csv").resolve()),
    }
    write_group_csv(csv_dir / "inventory-active-servers.csv", groups.get("active", []), "active")
    write_group_csv(csv_dir / "inventory-inactive-servers.csv", groups.get("inactive", []), "inactive")
    write_group_csv(csv_dir / "inventory-monitorable-servers.csv", groups["monitorable"], "monitorable")
    write_group_csv(csv_dir / "inventory-unmonitorable-servers.csv", groups["unmonitorable"], "unmonitorable")
    summary = {
        "mode": groups.get("mode", "inventory-classification"),
        "total_servers": int(groups.get("total_servers") or 0),
        "active_servers": int(groups.get("active_servers") or 0),
        "inactive_servers": int(groups.get("inactive_servers") or 0),
        "monitorable_servers": int(groups.get("monitorable_servers") or 0),
        "unmonitorable_servers": int(groups.get("unmonitorable_servers") or 0),
        "metric_classification_attempted": bool(groups.get("metric_classification_attempted")),
        "files": files,
    }
    inventory = dict(inventory)
    inventory["servers"] = monitorable if metrics_attempted else active_servers
    inventory["inventory_groups"] = summary
    environment = source_environment(inventory)
    environment["notes"] = (
        (environment.get("notes") + " " if environment.get("notes") else "")
        + f"清单分层：共 {summary['total_servers']} 台，仍在使用 {summary['active_servers']} 台，已停用 {summary['inactive_servers']} 台；"
        + f"可正常监测主机指标 {summary['monitorable_servers']} 台，不可正常监测 {summary['unmonitorable_servers']} 台。"
    )
    inventory["environment"] = environment
    return inventory


def collect_host_metrics_payload(
    doops: str,
    metrics_target: str,
    servers: list[dict[str, Any]],
    session: str,
    timeout: int,
    profile: str,
    metrics_timeout: int,
    workspace_mode: str = "auto",
    keep_remote_workspace: bool = False,
    requested_remote_artifact_dir: str | None = None,
    auth_security: bool = False,
    auth_log_window_days: int = 7,
) -> tuple[dict[str, Any], dict[str, Any]]:
    evidence, _gateway = target_preflight(doops, metrics_target, session, timeout)
    evidence.update({"metrics_target": metrics_target, "profile": profile, "server_count": len(servers)})
    command = build_host_metrics_command(servers, metrics_target, profile, metrics_timeout, auth_security, auth_log_window_days)
    try:
        payload, trace = execute_payload(
            doops,
            metrics_target,
            session,
            timeout,
            command,
            build_host_metrics_script(servers, metrics_target, profile, metrics_timeout, auth_security, auth_log_window_days),
            {"kind": "host-metrics", "target": metrics_target, "session": session, "profile": profile, "server_count": len(servers), "created_at": now_local(), "auth_security": bool(auth_security)},
            workspace_mode,
            keep_remote_workspace,
            requested_remote_artifact_dir,
        )
        evidence.update(trace)
    except Exception as exc:
        fallback_servers = []
        for server in servers:
            fallback_servers.append(
                {
                    "index": int(server.get("metrics_index") or server.get("probe_index") or 0),
                    "address": str(server.get("address") or ""),
                    "metrics_method": "not-supported",
                    "status": "failed",
                    "metrics": {},
                    "services": [],
                    "processes": [],
                    "collection_errors": [{"stage": "host-metrics", "type": "host-metrics-payload-missing"}],
                }
            )
        clean_payload = {
            "metrics_target": metrics_target,
            "profile": profile,
            "checked_at": now_local(),
            "servers": fallback_servers,
        }
        evidence["payload_error"] = str(exc)
        evidence["transport_mode"] = "inline-exec"
        evidence["remote_payload"] = clean_payload
        return clean_payload, evidence
    clean_servers = []
    for item in payload.get("servers", []) if isinstance(payload.get("servers", []), list) else []:
        if isinstance(item, dict):
            clean_item = sanitize_metric_payload(item)
            clean_item["index"] = item.get("index")
            clean_item["address"] = str(item.get("address") or "")
            clean_servers.append(clean_item)
    clean_payload = {
        "metrics_target": str(payload.get("metrics_target") or metrics_target),
        "profile": str(payload.get("profile") or profile),
        "checked_at": str(payload.get("checked_at") or now_local()),
        "servers": clean_servers,
    }
    evidence["remote_payload"] = clean_payload
    return clean_payload, evidence


def apply_host_metrics_inventory(
    doops: str,
    source: dict[str, Any],
    metrics_target: str,
    session: str,
    timeout: int,
    profile: str,
    metrics_timeout: int,
    workspace_mode: str = "auto",
    keep_remote_workspace: bool = False,
    requested_remote_artifact_dir: str | None = None,
    auth_security: bool = False,
    auth_log_window_days: int = 7,
) -> tuple[dict[str, Any], dict[str, Any]]:
    servers = source.get("servers") if isinstance(source.get("servers"), list) else []
    if not servers:
        raise ValueError("inventory must contain at least one server")
    indexed_servers = []
    for index, server in enumerate(servers, start=1):
        if not isinstance(server, dict):
            raise ValueError(f"server #{index} must be an object")
        indexed = dict(server)
        indexed["metrics_index"] = index
        indexed_servers.append(indexed)

    payload_servers: list[dict[str, Any]] = []
    chunk_evidence = []
    for chunk_index, chunk in enumerate(chunk_host_metrics_servers(indexed_servers, metrics_target, profile, metrics_timeout), start=1):
        chunk_payload, chunk_item = collect_host_metrics_payload(
            doops,
            metrics_target,
            chunk,
            f"{session}-metrics-{chunk_index:02d}",
            timeout,
            profile,
            metrics_timeout,
            workspace_mode,
            keep_remote_workspace,
            requested_remote_artifact_dir,
            auth_security,
            auth_log_window_days,
        )
        payload_servers.extend(item for item in chunk_payload.get("servers", []) if isinstance(item, dict))
        chunk_evidence.append(chunk_item)
    observations = {
        int(item.get("index")): item
        for item in payload_servers
        if isinstance(item, dict) and str(item.get("index") or "").isdigit()
    }
    collected_servers = []
    for index, raw_server in enumerate(servers, start=1):
        if not isinstance(raw_server, dict):
            raise ValueError(f"server #{index} must be an object")
        observed = observations.get(
            index,
            {
                "index": index,
                "address": raw_server.get("address") or "",
                "metrics_method": "not-supported",
                "status": "failed",
                "metrics": {},
                "services": [],
                "processes": [],
                "collection_errors": [{"stage": "host-metrics", "type": "missing-observation"}],
            },
        )
        clean_observed = sanitize_metric_payload(observed)
        collect = raw_server.get("collect") if isinstance(raw_server.get("collect"), dict) else {}
        errors = clean_observed["collection_errors"]
        status = "collected" if clean_observed["status"] == "collected" and clean_observed["metrics"] else "failed"
        clean_server = base_inventory_server(raw_server, index)
        clean_server["collect"] = {
            **(clean_server.get("collect") if isinstance(clean_server.get("collect"), dict) else {}),
            "host_metrics": True,
            "metrics_method": clean_observed["metrics_method"],
            "metrics_probe_target": metrics_target,
            "host_metrics_profile": profile,
            "host_metrics_status": status,
        }
        clean_server["metrics"] = clean_observed["metrics"]
        clean_server["services"] = clean_observed["services"]
        clean_server["processes"] = clean_observed["processes"]
        if clean_observed.get("auth_security"):
            clean_server["auth_security"] = merge_auth_security(clean_observed.get("auth_security"))
        elif clean_server.get("auth_security"):
            clean_server["auth_security"] = merge_auth_security(clean_server.get("auth_security"))
        clean_server["collection_errors"] = errors
        clean_server["collection_trace"] = collection_trace_from_evidence(chunk_evidence[0] if chunk_evidence else {}, metrics_target, session)
        clean_server["notes"] = f"{raw_server.get('notes') or ''} host metrics through {metrics_target} profile {profile}.".strip()
        collected_servers.append(clean_server)
    inventory = dict(source)
    environment = source_environment(source)
    environment["notes"] = (
        (environment.get("notes") + " " if environment.get("notes") else "")
        + f"Read-only host metrics attempted from doops target {metrics_target} for {len(collected_servers)} server(s)."
    )
    inventory["environment"] = environment
    inventory["servers"] = collected_servers
    if chunk_evidence:
        inventory["doops_channel"] = doops_channel_from_evidence(chunk_evidence[0], metrics_target, session)
    evidence = {
        "metrics_target": metrics_target,
        "profile": profile,
        "session": session,
        "collected_at": now_local(),
        "server_count": len(servers),
        "chunks": chunk_evidence,
        "remote_payload": {"metrics_target": metrics_target, "profile": profile, "servers": payload_servers},
    }
    return inventory, evidence


def collect_snmp_metrics_payload(
    doops: str,
    snmp_target: str,
    servers: list[dict[str, Any]],
    session: str,
    timeout: int,
    profile: str,
    snmp_timeout: int,
    workspace_mode: str = "auto",
    keep_remote_workspace: bool = False,
    requested_remote_artifact_dir: str | None = None,
) -> tuple[dict[str, Any], dict[str, Any]]:
    evidence, _gateway = target_preflight(doops, snmp_target, session, timeout)
    evidence.update({"metrics_target": snmp_target, "profile": profile, "server_count": len(servers)})
    command = build_snmp_metrics_command(servers, snmp_target, profile, snmp_timeout)
    try:
        payload, trace = execute_payload(
            doops,
            snmp_target,
            session,
            timeout,
            command,
            build_snmp_metrics_script(servers, snmp_target, profile, snmp_timeout),
            {"kind": "snmp-metrics", "target": snmp_target, "session": session, "profile": profile, "server_count": len(servers), "created_at": now_local()},
            workspace_mode,
            keep_remote_workspace,
            requested_remote_artifact_dir,
        )
        evidence.update(trace)
    except Exception as exc:
        fallback_servers = []
        for server in servers:
            fallback_servers.append(
                {
                    "index": int(server.get("metrics_index") or server.get("probe_index") or 0),
                    "address": str(server.get("address") or ""),
                    "metrics_method": "snmpv3",
                    "status": "failed",
                    "metrics": {},
                    "services": [],
                    "processes": [],
                    "collection_errors": [{"stage": "snmp-metrics", "type": "snmp-invalid-payload"}],
                }
            )
        clean_payload = {
            "metrics_target": snmp_target,
            "profile": profile,
            "checked_at": now_local(),
            "servers": fallback_servers,
        }
        evidence["payload_error"] = str(exc)
        evidence["transport_mode"] = "inline-exec"
        evidence["remote_payload"] = clean_payload
        return clean_payload, evidence
    clean_servers = []
    for item in payload.get("servers", []) if isinstance(payload.get("servers", []), list) else []:
        if isinstance(item, dict):
            clean_item = sanitize_metric_payload(item)
            clean_item["index"] = item.get("index")
            clean_item["address"] = str(item.get("address") or "")
            if clean_item["metrics_method"] in {"", "not-supported"}:
                clean_item["metrics_method"] = "snmpv3"
            clean_servers.append(clean_item)
    clean_payload = {
        "metrics_target": str(payload.get("metrics_target") or snmp_target),
        "profile": str(payload.get("profile") or profile),
        "checked_at": str(payload.get("checked_at") or now_local()),
        "servers": clean_servers,
    }
    evidence["remote_payload"] = clean_payload
    return clean_payload, evidence


def apply_snmp_metrics_inventory(
    doops: str,
    source: dict[str, Any],
    snmp_target: str,
    session: str,
    timeout: int,
    profile: str,
    snmp_timeout: int,
    workspace_mode: str = "auto",
    keep_remote_workspace: bool = False,
    requested_remote_artifact_dir: str | None = None,
) -> tuple[dict[str, Any], dict[str, Any]]:
    servers = source.get("servers") if isinstance(source.get("servers"), list) else []
    if not servers:
        raise ValueError("inventory must contain at least one server")
    indexed_servers = []
    for index, server in enumerate(servers, start=1):
        if not isinstance(server, dict):
            raise ValueError(f"server #{index} must be an object")
        indexed = dict(server)
        indexed["metrics_index"] = index
        indexed_servers.append(indexed)

    payload_servers: list[dict[str, Any]] = []
    chunk_evidence = []
    for chunk_index, chunk in enumerate(chunk_snmp_metrics_servers(indexed_servers, snmp_target, profile, snmp_timeout), start=1):
        chunk_payload, chunk_item = collect_snmp_metrics_payload(
            doops,
            snmp_target,
            chunk,
            f"{session}-snmp-{chunk_index:02d}",
            timeout,
            profile,
            snmp_timeout,
            workspace_mode,
            keep_remote_workspace,
            requested_remote_artifact_dir,
        )
        payload_servers.extend(item for item in chunk_payload.get("servers", []) if isinstance(item, dict))
        chunk_evidence.append(chunk_item)
    observations = {
        int(item.get("index")): item
        for item in payload_servers
        if isinstance(item, dict) and str(item.get("index") or "").isdigit()
    }
    collected_servers = []
    for index, raw_server in enumerate(servers, start=1):
        if not isinstance(raw_server, dict):
            raise ValueError(f"server #{index} must be an object")
        observed = observations.get(
            index,
            {
                "index": index,
                "address": raw_server.get("address") or "",
                "metrics_method": "snmpv3",
                "status": "failed",
                "metrics": {},
                "services": [],
                "processes": [],
                "collection_errors": [{"stage": "snmp-metrics", "type": "missing-observation"}],
            },
        )
        clean_observed = sanitize_metric_payload(observed)
        errors = clean_observed["collection_errors"]
        status = "collected" if clean_observed["status"] == "collected" and clean_observed["metrics"] else "failed"
        clean_server = base_inventory_server(raw_server, index)
        clean_server["collect"] = {
            **(clean_server.get("collect") if isinstance(clean_server.get("collect"), dict) else {}),
            "host_metrics": True,
            "metrics_method": "snmpv3",
            "metrics_probe_target": snmp_target,
            "host_metrics_profile": profile,
            "host_metrics_status": status,
        }
        clean_server["metrics"] = clean_observed["metrics"]
        clean_server["services"] = clean_observed["services"]
        clean_server["processes"] = clean_observed["processes"]
        clean_server["collection_errors"] = errors
        clean_server["collection_trace"] = collection_trace_from_evidence(chunk_evidence[0] if chunk_evidence else {}, snmp_target, session)
        clean_server["notes"] = f"{raw_server.get('notes') or ''} SNMPv3 host metrics through {snmp_target} profile {profile}.".strip()
        collected_servers.append(clean_server)
    inventory = dict(source)
    environment = source_environment(source)
    environment["notes"] = (
        (environment.get("notes") + " " if environment.get("notes") else "")
        + f"Read-only SNMPv3 host metrics attempted from doops target {snmp_target} for {len(collected_servers)} server(s)."
    )
    inventory["environment"] = environment
    inventory["servers"] = collected_servers
    if chunk_evidence:
        inventory["doops_channel"] = doops_channel_from_evidence(chunk_evidence[0], snmp_target, session)
    evidence = {
        "metrics_target": snmp_target,
        "profile": profile,
        "session": session,
        "collected_at": now_local(),
        "server_count": len(servers),
        "chunks": chunk_evidence,
        "remote_payload": {"metrics_target": snmp_target, "profile": profile, "servers": payload_servers},
    }
    return inventory, evidence


def build_preflight(args: argparse.Namespace) -> dict[str, Any]:
    environment = resolve_environment(args.environment or args.target)
    target = str(args.target or environment.get("target") or "")
    config_path = find_doops_config()
    payload: dict[str, Any] = {
        "ok": False,
        "checked_at": now_local(),
        "environment_key": environment.get("key", ""),
        "environment_name": environment.get("display_name", environment.get("key", "")),
        "target": target,
        "doops_path": None,
        "config_path": str(config_path) if config_path else None,
        "configured_targets": config_target_names(config_path),
        "targets_returncode": None,
        "targets_stdout_excerpt": "",
        "targets_stderr_excerpt": "",
        "error": "",
    }
    try:
        doops = find_doops(args.doops)
    except Exception as exc:
        payload["error"] = str(exc)
        return payload
    payload["doops_path"] = doops
    if not target:
        payload["error"] = "provide --environment zheyin|hdu or --target <doops-target>"
        return payload
    try:
        result = run_command([doops, "targets", "--target", target], timeout=min(args.timeout, 30))
    except Exception as exc:
        payload["error"] = f"doops targets check failed: {exc.__class__.__name__}"
        return payload
    payload["targets_returncode"] = result.returncode
    payload["targets_stdout_excerpt"] = redact(result.stdout)
    payload["targets_stderr_excerpt"] = redact(result.stderr)
    if result.returncode == 0:
        payload["ok"] = True
    else:
        payload["error"] = f"doops targets check failed for target {target}"
    return payload


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Collect a server-health inventory through doops.")
    parser.add_argument("--inventory", help="User-supplied JSON, CSV, TSV, or Excel server list.")
    parser.add_argument("--environment", "--env", help="Managed environment name, e.g. zheyin/浙音 or hdu/杭电.")
    parser.add_argument("--target", help="Explicit doops target. Defaults from --environment when known.")
    parser.add_argument("--probe-target", help="Use one doops target as an internal network probe for every inventory server.")
    parser.add_argument("--host-metrics", action="store_true", help="Collect read-only host metrics for inventory servers through SSH/WinRM from a doops target.")
    parser.add_argument("--host-metrics-target", default="zheyin", help="doops target used as the internal SSH/WinRM host-metrics collector.")
    parser.add_argument("--host-metrics-profile", default="zheyin-monitor", help="Named read-only credential profile available on the doops collector node.")
    parser.add_argument("--metrics-timeout", type=int, default=20, help="Per-server SSH/WinRM metrics collection timeout in seconds.")
    parser.add_argument("--snmp-metrics", action="store_true", help="Collect read-only SNMPv3 host metrics for inventory servers from a doops target.")
    parser.add_argument("--snmp-target", default="zheyin", help="doops target used as the internal SNMPv3 metrics collector.")
    parser.add_argument("--snmp-profile", default="zheyin-snmpv3-monitor", help="Named SNMPv3 read-only profile available on the doops collector node.")
    parser.add_argument("--snmp-timeout", type=int, default=5, help="Per-server SNMPv3 metrics collection timeout in seconds.")
    parser.add_argument("--use-excel-runtime-credentials", action="store_true", help="Use Excel credentials in memory for this run only; never writes them to output artifacts.")
    parser.add_argument("--auth-security", action="store_true", help="Collect read-only authentication log and protection configuration evidence.")
    parser.add_argument("--auth-log-window-days", type=int, default=7, help="Authentication log lookback window in days.")
    parser.add_argument("--password-strength-audit", action="store_true", help="Score supplied runtime credentials locally without writing plaintext passwords.")
    parser.add_argument("--password-strength-source", default="runtime-auth", help="Password strength source label. Default: runtime-auth.")
    parser.add_argument("--classify-inventory", action="store_true", help="Split inventory into active/inactive and monitorable/unmonitorable CSVs; final inventory keeps only monitorable servers when metrics are attempted.")
    parser.add_argument("--limit", type=int, help="Limit imported inventory rows, useful for Excel samples.")
    parser.add_argument("--session", help="doops session name. Defaults to server-health-<timestamp>.")
    parser.add_argument("--doops", help="Path to doops CLI. Defaults to DOOPS_BIN, PATH, or ~/.local/bin/doops.")
    parser.add_argument("--preflight", action="store_true", help="Print sanitized doops diagnostics as JSON and exit.")
    parser.add_argument("--doops-workspace-mode", choices=("auto", "inline", "workspace"), default="auto", help="Remote collector transport: auto, inline, or push/read workspace.")
    parser.add_argument("--keep-remote-workspace", action="store_true", help="Keep the doops remote workspace for diagnostics instead of cleaning it.")
    parser.add_argument("--remote-artifact-dir", default="/root/ws/<session>/server-health", help="Remote artifact directory template for workspace mode.")
    parser.add_argument("--out", help="Output inventory JSON path. Required unless --preflight is used.")
    parser.add_argument("--evidence-out", help="Optional sanitized doops evidence JSON path.")
    parser.add_argument("--timeout", type=int, default=120, help="doops exec timeout in seconds.")
    parser.add_argument("--probe-timeout", type=float, default=1.2, help="Per-host TCP/ICMP probe timeout in seconds.")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    session = args.session or f"server-health-{dt.datetime.now().strftime('%Y%m%d-%H%M%S')}"
    if args.preflight:
        payload = build_preflight(args)
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return 0 if payload.get("ok") else 1
    if not args.out:
        raise ValueError("provide --out <inventory.json> unless --preflight is used")
    doops = find_doops(args.doops)
    classification_groups = None
    if args.inventory:
        source = load_source_inventory(
            Path(args.inventory),
            args.limit,
            include_runtime_auth=bool(args.use_excel_runtime_credentials or args.password_strength_audit),
        )
        if args.classify_inventory:
            source, classification_groups = classify_source_inventory(source)
        if args.classify_inventory and not source.get("servers"):
            inventory = {"environment": source_environment(source), "servers": []}
            evidence = {"classification": "no-active-servers", "server_count": 0}
        elif args.probe_target:
            inventory, evidence = collect_probe_inventory(
                doops,
                source,
                args.probe_target,
                session,
                args.timeout,
                args.probe_timeout,
                args.doops_workspace_mode,
                bool(args.keep_remote_workspace),
                args.remote_artifact_dir,
            )
            if args.host_metrics:
                inventory, metrics_evidence = apply_host_metrics_inventory(
                    doops,
                    inventory,
                    args.host_metrics_target,
                    session,
                    args.timeout,
                    args.host_metrics_profile,
                    args.metrics_timeout,
                    args.doops_workspace_mode,
                    bool(args.keep_remote_workspace),
                    args.remote_artifact_dir,
                    bool(args.auth_security),
                    args.auth_log_window_days,
                )
                evidence = {"network_probe": evidence, "host_metrics": metrics_evidence}
            if args.snmp_metrics:
                inventory, snmp_evidence = apply_snmp_metrics_inventory(
                    doops,
                    inventory,
                    args.snmp_target,
                    session,
                    args.timeout,
                    args.snmp_profile,
                    args.snmp_timeout,
                    args.doops_workspace_mode,
                    bool(args.keep_remote_workspace),
                    args.remote_artifact_dir,
                )
                evidence = {**({"network_probe": evidence} if "network_probe" not in evidence else evidence), "snmp_metrics": snmp_evidence}
        elif args.host_metrics:
            inventory, evidence = apply_host_metrics_inventory(
                doops,
                source,
                args.host_metrics_target,
                session,
                args.timeout,
                args.host_metrics_profile,
                args.metrics_timeout,
                args.doops_workspace_mode,
                bool(args.keep_remote_workspace),
                args.remote_artifact_dir,
                bool(args.auth_security),
                args.auth_log_window_days,
            )
            if args.snmp_metrics:
                inventory, snmp_evidence = apply_snmp_metrics_inventory(
                    doops,
                    inventory,
                    args.snmp_target,
                    session,
                    args.timeout,
                    args.snmp_profile,
                    args.snmp_timeout,
                    args.doops_workspace_mode,
                    bool(args.keep_remote_workspace),
                    args.remote_artifact_dir,
                )
                evidence = {"host_metrics": evidence, "snmp_metrics": snmp_evidence}
        elif args.snmp_metrics:
            inventory, evidence = apply_snmp_metrics_inventory(
                doops,
                source,
                args.snmp_target,
                session,
                args.timeout,
                args.snmp_profile,
                args.snmp_timeout,
                args.doops_workspace_mode,
                bool(args.keep_remote_workspace),
                args.remote_artifact_dir,
            )
        else:
            inventory, evidence = collect_inventory(
                doops,
                source,
                session,
                args.timeout,
                args.doops_workspace_mode,
                bool(args.keep_remote_workspace),
                args.remote_artifact_dir,
                bool(args.auth_security),
                args.auth_log_window_days,
            )
    else:
        environment = resolve_environment(args.environment or args.target)
        target = args.target or environment.get("target")
        if not target:
            raise ValueError("provide --inventory <server-list>, --environment zheyin|hdu, or --target <doops-target>")
        payload, evidence = collect_doops_payload(
            doops,
            str(target),
            session,
            args.timeout,
            args.doops_workspace_mode,
            bool(args.keep_remote_workspace),
            args.remote_artifact_dir,
            bool(args.auth_security),
            args.auth_log_window_days,
        )
        inventory = build_inventory(environment, str(target), session, payload)
        inventory["doops_channel"] = doops_channel_from_evidence(evidence, str(target), session)
        if inventory.get("servers"):
            inventory["servers"][0]["collection_trace"] = collection_trace_from_evidence(evidence, str(target), session)

    inventory, app_evidence = apply_application_checks(doops, inventory, session, args.timeout)
    if app_evidence.get("checks"):
        evidence = {"collection": evidence, "application_checks": app_evidence}
    if args.password_strength_audit:
        inventory = apply_password_strength_inventory(inventory, source if args.inventory else inventory, enabled=True, source_name=args.password_strength_source)
    else:
        inventory = apply_password_strength_inventory(inventory, source if args.inventory else inventory, enabled=False, source_name=args.password_strength_source)
    inventory = apply_auth_security_boundary_inventory(inventory, enabled=bool(args.auth_security), window_days=args.auth_log_window_days)

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    if args.classify_inventory and classification_groups is not None:
        inventory = finalize_inventory_classification(
            inventory,
            classification_groups,
            out_path,
            metrics_attempted=bool(args.host_metrics or args.snmp_metrics),
        )
        evidence = {"collection": evidence, "inventory_groups": inventory.get("inventory_groups", {})}
    out_path.write_text(json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8")

    evidence_path = Path(args.evidence_out) if args.evidence_out else out_path.with_name(f"{out_path.stem}.evidence.json")
    evidence_path.write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"inventory written to {out_path}")
    print(f"evidence written to {evidence_path}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"error: {exc}", file=sys.stderr)
        raise SystemExit(1)
