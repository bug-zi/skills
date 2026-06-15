#!/usr/bin/env python3
import argparse
import datetime as dt
import http.cookiejar
import ipaddress
import json
import pathlib
import re
import shutil
import ssl
import subprocess
import sys
import tempfile
import textwrap
import time
import urllib.parse
import urllib.request

from openpyxl import load_workbook


ROOT = pathlib.Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = ROOT / "config" / "connection.local.json"
BUILTIN_NSG_FIREWALL = {'base_url': 'https://172.16.15.6:8080',
 'doops_cwd': '/Users/wwyz/Documents/网络安全智能体',
 'doops_target': 'zy',
 'insecure_tls': True,
 'password': '',
 'username': 'sysapi'}
TARGET_IP_OBJECT = "翻墙风险IP"
TARGET_DOMAIN_OBJECT = "翻墙风险域名"
REMOTE_SCRIPT = r'''#!/usr/bin/env python3
import argparse
import datetime as dt
import http.cookiejar
import ipaddress
import json
import pathlib
import re
import ssl
import sys
import time
import urllib.parse
import urllib.request

TARGET_IP_OBJECT = "翻墙风险IP"
TARGET_DOMAIN_OBJECT = "翻墙风险域名"


def add_cookie(cookie_jar, url, name, value):
    parsed = urllib.parse.urlparse(url)
    cookie_jar.set_cookie(
        http.cookiejar.Cookie(
            version=0,
            name=name,
            value=value,
            port=None,
            port_specified=False,
            domain=parsed.hostname or "",
            domain_specified=False,
            domain_initial_dot=False,
            path="/",
            path_specified=True,
            secure=parsed.scheme == "https",
            expires=None,
            discard=True,
            comment=None,
            comment_url=None,
            rest={},
            rfc2109=False,
        )
    )


def make_opener(insecure_tls):
    context = ssl._create_unverified_context() if insecure_tls else ssl.create_default_context()
    cookies = http.cookiejar.CookieJar()
    opener = urllib.request.build_opener(
        urllib.request.HTTPCookieProcessor(cookies),
        urllib.request.HTTPSHandler(context=context),
    )
    return opener, cookies


def post_json(opener, url, body, tries=6, timeout=25):
    payload = json.dumps(body, ensure_ascii=False).encode("utf-8")
    last_error = None
    for _ in range(tries):
        try:
            req = urllib.request.Request(
                url,
                data=payload,
                headers={
                    "Content-Type": "application/json",
                    "Accept": "application/json",
                    "Connection": "close",
                },
                method="POST",
            )
            with opener.open(req, timeout=timeout) as resp:
                raw = resp.read().decode("utf-8", "replace")
                return json.loads(raw) if raw else None
        except Exception as exc:
            last_error = exc
            time.sleep(2)
    raise last_error


def login(cfg):
    base_url = cfg["base_url"].rstrip("/")
    opener, cookies = make_opener(bool(cfg.get("insecure_tls", True)))
    data = post_json(
        opener,
        base_url + "/v1.0/login",
        {"username": cfg["username"], "password": cfg["password"]},
        tries=8,
    )
    token = data.get("result", {}).get("token") if isinstance(data, dict) else None
    if not token:
        raise RuntimeError(f"Login failed: {data}")
    add_cookie(cookies, base_url + "/v1.0/rest/", "token", token)
    return opener


def rest_call(opener, cfg, module, function, body, page_index=1, page_size=20, tries=6):
    payload = [
        {
            "head": {
                "module": module,
                "function": function,
                "page_index": page_index,
                "page_size": page_size,
            },
            "body": body,
        }
    ]
    return post_json(opener, cfg["base_url"].rstrip("/") + "/v1.0/rest/", payload, tries=tries)


def iter_address_objects(opener, cfg):
    first = rest_call(
        opener,
        cfg,
        "obj_address",
        "get_obj_addr_list",
        {"obj_addr": [{"is_detail": False}]},
        page_index=1,
        page_size=20,
    )
    head = first.get("head", {}) if isinstance(first, dict) else {}
    if head.get("error_code") not in (0, "0"):
        raise RuntimeError(f"get_obj_addr_list failed: {head}")
    total = int(head.get("total") or 0)
    page_count = (total + 19) // 20 if total > 0 else 1
    pages = [first]
    for page in range(2, page_count + 1):
        pages.append(
            rest_call(
                opener,
                cfg,
                "obj_address",
                "get_obj_addr_list",
                {"obj_addr": [{"is_detail": False}]},
                page_index=page,
                page_size=20,
            )
        )
    for page in pages:
        for item in page.get("data") or []:
            yield item


def read_target_objects(opener, cfg):
    targets = {}
    duplicates = []
    for item in iter_address_objects(opener, cfg):
        name = item.get("name")
        if name in (TARGET_IP_OBJECT, TARGET_DOMAIN_OBJECT):
            if name in targets:
                duplicates.append(name)
            targets[name] = item
    if duplicates:
        raise RuntimeError(f"Duplicate target address objects: {', '.join(sorted(set(duplicates)))}")
    missing = [name for name in (TARGET_IP_OBJECT, TARGET_DOMAIN_OBJECT) if name not in targets]
    if missing:
        raise RuntimeError(f"Missing target address objects: {', '.join(missing)}")
    return targets


def collect_strings(value):
    if isinstance(value, dict):
        out = []
        for item in value.values():
            out.extend(collect_strings(item))
        return out
    if isinstance(value, list):
        out = []
        for item in value:
            out.extend(collect_strings(item))
        return out
    if value is None:
        return []
    return [str(value)]


def read_policy_objects(opener, cfg):
    data = rest_call(
        opener,
        cfg,
        "sec_policy",
        "get_sec_policy",
        {"sec_policy": [{"name": "", "is_detail": True}]},
        page_index=1,
        page_size=20,
    )
    head = data.get("head", {}) if isinstance(data, dict) else {}
    if head.get("error_code") not in (0, "0"):
        raise RuntimeError(f"get_sec_policy failed: {head}")
    policies = {}
    for item in data.get("data") or []:
        if item.get("name") in (TARGET_IP_OBJECT, TARGET_DOMAIN_OBJECT):
            policies[item["name"]] = item
    return policies


def validate_correct_targets(targets, policies):
    errors = []
    checks = {}
    ip_record = targets[TARGET_IP_OBJECT]
    domain_record = targets[TARGET_DOMAIN_OBJECT]

    for member in ip_record.get("include") or []:
        if member.get("addr_type") != "host" or not member.get("ip"):
            errors.append(f"{TARGET_IP_OBJECT} contains non-host member: {member}")
            continue
        try:
            ipaddress.ip_address(str(member["ip"]))
        except ValueError:
            errors.append(f"{TARGET_IP_OBJECT} contains invalid IP member: {member}")

    domain_re = re.compile(r"^[a-z0-9.-]+\.[a-z0-9-]+$", re.I)
    for member in domain_record.get("include") or []:
        if member.get("addr_type") != "domain" or not member.get("domain"):
            errors.append(f"{TARGET_DOMAIN_OBJECT} contains non-domain member: {member}")
            continue
        domain = str(member["domain"]).lower().rstrip(".")
        if not domain_re.match(domain):
            errors.append(f"{TARGET_DOMAIN_OBJECT} contains invalid domain member: {member}")

    for name in (TARGET_IP_OBJECT, TARGET_DOMAIN_OBJECT):
        policy = policies.get(name)
        if not policy:
            errors.append(f"Missing security policy with expected name: {name}")
            checks[name] = {"object_found": True, "policy_found": False, "policy_references_object": False}
            continue
        text = "\n".join(collect_strings(policy))
        references = name in text
        if not references:
            errors.append(f"Security policy {name} does not reference address object {name}")
        checks[name] = {
            "object_found": True,
            "policy_found": True,
            "policy_references_object": references,
            "reference": targets[name].get("reference"),
            "action": policy.get("action") or policy.get("act") or policy.get("policy_action"),
        }

    if errors:
        raise RuntimeError("Preflight target verification failed: " + " | ".join(errors))
    return checks


def existing_values(record, kind):
    values = []
    for item in record.get("include") or []:
        if kind == "ip" and item.get("addr_type") == "host" and item.get("ip"):
            values.append(str(ipaddress.ip_address(str(item["ip"]))))
        elif kind == "domain" and item.get("addr_type") == "domain" and item.get("domain"):
            values.append(str(item["domain"]).lower().rstrip("."))
    return values


def ip_sort_key(item):
    return tuple(int(part) for part in item.split("."))


def update_address_object(opener, cfg, before, name, include):
    body = {
        "obj_addr": [
            {
                "name": name,
                "oldname": name,
                "desc": before.get("desc", "") or "",
                "include": include,
                "exclude": before.get("exclude") or [],
            }
        ]
    }
    data = rest_call(opener, cfg, "obj_address", "set_obj_addr_conf", body)
    head = data.get("head", {}) if isinstance(data, dict) else {}
    if head.get("error_code") not in (0, "0"):
        raise RuntimeError(f"set_obj_addr_conf {name} failed: {head}")
    return data


def publish(cfg, targets_payload, apply=False):
    excel_ips = targets_payload["ips"]
    excel_domains = targets_payload["domains"]
    opener = login(cfg)
    try:
        before = read_target_objects(opener, cfg)
        policies = read_policy_objects(opener, cfg)
        preflight = validate_correct_targets(before, policies)

        current_ips = existing_values(before[TARGET_IP_OBJECT], "ip")
        current_domains = existing_values(before[TARGET_DOMAIN_OBJECT], "domain")
        merged_ips = sorted(dict.fromkeys(current_ips + excel_ips), key=ip_sort_key)
        merged_domains = sorted(dict.fromkeys(current_domains + excel_domains))
        add_ips = [item for item in excel_ips if item not in set(current_ips)]
        add_domains = [item for item in excel_domains if item not in set(current_domains)]

        if len(merged_ips) > 512 or len(merged_domains) > 512:
            raise RuntimeError(
                f"Object include limit exceeded: {TARGET_IP_OBJECT}={len(merged_ips)}, "
                f"{TARGET_DOMAIN_OBJECT}={len(merged_domains)}"
            )

        report = {
            "time": dt.datetime.now().isoformat(timespec="seconds"),
            "base_url": cfg["base_url"],
            "transport": "doops-remote",
            "excel": targets_payload.get("excel", {}),
            "preflight": preflight,
            "before": {
                TARGET_IP_OBJECT: len(current_ips),
                TARGET_DOMAIN_OBJECT: len(current_domains),
            },
            "plan": {
                "add_ip_count": len(add_ips),
                "add_domain_count": len(add_domains),
                "final_ip_count": len(merged_ips),
                "final_domain_count": len(merged_domains),
                "add_domains": add_domains,
            },
            "applied": False,
        }

        if apply and (add_ips or add_domains):
            if add_ips:
                update_address_object(
                    opener,
                    cfg,
                    before[TARGET_IP_OBJECT],
                    TARGET_IP_OBJECT,
                    [{"ip": item, "addr_type": "host"} for item in merged_ips],
                )
            if add_domains:
                update_address_object(
                    opener,
                    cfg,
                    before[TARGET_DOMAIN_OBJECT],
                    TARGET_DOMAIN_OBJECT,
                    [{"domain": item, "addr_type": "domain"} for item in merged_domains],
                )
            after = read_target_objects(opener, cfg)
            after_policies = read_policy_objects(opener, cfg)
            report["postflight"] = validate_correct_targets(after, after_policies)
            after_ips = existing_values(after[TARGET_IP_OBJECT], "ip")
            after_domains = existing_values(after[TARGET_DOMAIN_OBJECT], "domain")
            report["applied"] = True
            report["after"] = {
                TARGET_IP_OBJECT: len(after_ips),
                TARGET_DOMAIN_OBJECT: len(after_domains),
            }
            report["verification"] = {
                "missing_ips": sorted(set(excel_ips) - set(after_ips), key=ip_sort_key),
                "missing_domains": sorted(set(excel_domains) - set(after_domains)),
            }
        elif apply:
            report["applied"] = True
            report["after"] = report["before"]
            report["verification"] = {"missing_ips": [], "missing_domains": []}
            report["postflight"] = preflight
        return report
    finally:
        try:
            post_json(
                opener,
                cfg["base_url"].rstrip("/") + "/v1.0/out",
                {"username": cfg["username"]},
                tries=1,
            )
        except Exception:
            pass


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    root = pathlib.Path(__file__).resolve().parent
    cfg = json.loads((root / "connection.local.json").read_text(encoding="utf-8"))["nsg_firewall"]
    targets_payload = json.loads((root / "targets.json").read_text(encoding="utf-8"))
    report = publish(cfg, targets_payload, apply=args.apply)
    (root / "result.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    verification = report.get("verification", {})
    if verification.get("missing_ips") or verification.get("missing_domains"):
        raise SystemExit(3)


if __name__ == "__main__":
    main()
'''


def load_config(path):
    cfg = dict(BUILTIN_NSG_FIREWALL)
    if path.exists():
        cfg.update(json.loads(path.read_text(encoding="utf-8")).get("nsg_firewall", {}))
    missing = [key for key in ("base_url", "username", "password") if not cfg.get(key)]
    if missing:
        raise SystemExit(f"Missing config fields: {', '.join(missing)}. Add them to the built-in defaults or config file.")
    return cfg


def add_cookie(cookie_jar, url, name, value):
    parsed = urllib.parse.urlparse(url)
    cookie_jar.set_cookie(
        http.cookiejar.Cookie(
            version=0,
            name=name,
            value=value,
            port=None,
            port_specified=False,
            domain=parsed.hostname or "",
            domain_specified=False,
            domain_initial_dot=False,
            path="/",
            path_specified=True,
            secure=parsed.scheme == "https",
            expires=None,
            discard=True,
            comment=None,
            comment_url=None,
            rest={},
            rfc2109=False,
        )
    )


def make_opener(insecure_tls):
    context = ssl._create_unverified_context() if insecure_tls else ssl.create_default_context()
    cookies = http.cookiejar.CookieJar()
    opener = urllib.request.build_opener(
        urllib.request.HTTPCookieProcessor(cookies),
        urllib.request.HTTPSHandler(context=context),
    )
    return opener, cookies


def post_json(opener, url, body, tries=6, timeout=25):
    payload = json.dumps(body, ensure_ascii=False).encode("utf-8")
    last_error = None
    for _ in range(tries):
        try:
            req = urllib.request.Request(
                url,
                data=payload,
                headers={
                    "Content-Type": "application/json",
                    "Accept": "application/json",
                    "Connection": "close",
                },
                method="POST",
            )
            with opener.open(req, timeout=timeout) as resp:
                raw = resp.read().decode("utf-8", "replace")
                return json.loads(raw) if raw else None
        except Exception as exc:
            last_error = exc
            time.sleep(2)
    raise last_error


def login(cfg):
    base_url = cfg["base_url"].rstrip("/")
    opener, cookies = make_opener(bool(cfg.get("insecure_tls", True)))
    data = post_json(
        opener,
        base_url + "/v1.0/login",
        {"username": cfg["username"], "password": cfg["password"]},
        tries=8,
    )
    token = data.get("result", {}).get("token") if isinstance(data, dict) else None
    if not token:
        raise RuntimeError(f"Login failed: {data}")
    add_cookie(cookies, base_url + "/v1.0/rest/", "token", token)
    return opener


def rest_call(opener, cfg, module, function, body, page_index=1, page_size=20, tries=6):
    payload = [
        {
            "head": {
                "module": module,
                "function": function,
                "page_index": page_index,
                "page_size": page_size,
            },
            "body": body,
        }
    ]
    return post_json(opener, cfg["base_url"].rstrip("/") + "/v1.0/rest/", payload, tries=tries)


def rows_by_header(ws):
    rows = ws.iter_rows(values_only=True)
    header = [str(cell).strip() if cell is not None else "" for cell in next(rows)]
    for row in rows:
        yield dict(zip(header, row))


def normalize_domain(value):
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    parsed = urllib.parse.urlparse(text if "://" in text else "http://" + text)
    host = (parsed.hostname or text.split("/", 1)[0]).strip().lower().rstrip(".")
    if re.match(r"^[a-z0-9.-]+$", host) and "." in host:
        return host
    return None


def ip_sort_key(item):
    return tuple(int(part) for part in item.split("."))


def extract_excel(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    ip_values = []
    for row in rows_by_header(workbook["翻墙风险IP"]):
        raw = str(row.get("外网IP") or "").strip()
        if not raw:
            continue
        ip = ipaddress.ip_address(raw)
        if ip.version == 4:
            ip_values.append(str(ip))

    domain_values = []
    normalizations = []
    for row in rows_by_header(workbook["翻墙风险域名"]):
        raw = row.get("域名")
        domain = normalize_domain(raw)
        if domain:
            domain_values.append(domain)
            normalizations.append({"raw": str(raw), "domain": domain})

    ip_values = sorted(dict.fromkeys(ip_values), key=ip_sort_key)
    domain_values = sorted(dict.fromkeys(domain_values))
    return ip_values, domain_values, normalizations


def iter_address_objects(opener, cfg):
    first = rest_call(
        opener,
        cfg,
        "obj_address",
        "get_obj_addr_list",
        {"obj_addr": [{"is_detail": False}]},
        page_index=1,
        page_size=20,
    )
    head = first.get("head", {}) if isinstance(first, dict) else {}
    if head.get("error_code") not in (0, "0"):
        raise RuntimeError(f"get_obj_addr_list failed: {head}")
    total = int(head.get("total") or 0)
    page_count = (total + 19) // 20 if total > 0 else 1
    pages = [first]
    for page in range(2, page_count + 1):
        pages.append(
            rest_call(
                opener,
                cfg,
                "obj_address",
                "get_obj_addr_list",
                {"obj_addr": [{"is_detail": False}]},
                page_index=page,
                page_size=20,
            )
        )
    for page in pages:
        for item in page.get("data") or []:
            yield item


def read_target_objects(opener, cfg):
    targets = {}
    duplicates = []
    for item in iter_address_objects(opener, cfg):
        name = item.get("name")
        if name in (TARGET_IP_OBJECT, TARGET_DOMAIN_OBJECT):
            if name in targets:
                duplicates.append(name)
            targets[name] = item
    if duplicates:
        raise RuntimeError(f"Duplicate target address objects: {', '.join(sorted(set(duplicates)))}")
    missing = [name for name in (TARGET_IP_OBJECT, TARGET_DOMAIN_OBJECT) if name not in targets]
    if missing:
        raise RuntimeError(f"Missing target address objects: {', '.join(missing)}")
    return targets


def collect_strings(value):
    if isinstance(value, dict):
        out = []
        for item in value.values():
            out.extend(collect_strings(item))
        return out
    if isinstance(value, list):
        out = []
        for item in value:
            out.extend(collect_strings(item))
        return out
    if value is None:
        return []
    return [str(value)]


def read_policy_objects(opener, cfg):
    data = rest_call(
        opener,
        cfg,
        "sec_policy",
        "get_sec_policy",
        {"sec_policy": [{"name": "", "is_detail": True}]},
        page_index=1,
        page_size=20,
    )
    head = data.get("head", {}) if isinstance(data, dict) else {}
    if head.get("error_code") not in (0, "0"):
        raise RuntimeError(f"get_sec_policy failed: {head}")
    policies = {}
    for item in data.get("data") or []:
        if item.get("name") in (TARGET_IP_OBJECT, TARGET_DOMAIN_OBJECT):
            policies[item["name"]] = item
    return policies


def validate_correct_targets(targets, policies):
    errors = []
    checks = {}
    ip_record = targets[TARGET_IP_OBJECT]
    domain_record = targets[TARGET_DOMAIN_OBJECT]

    for member in ip_record.get("include") or []:
        if member.get("addr_type") != "host" or not member.get("ip"):
            errors.append(f"{TARGET_IP_OBJECT} contains non-host member: {member}")
            continue
        try:
            ipaddress.ip_address(str(member["ip"]))
        except ValueError:
            errors.append(f"{TARGET_IP_OBJECT} contains invalid IP member: {member}")

    domain_re = re.compile(r"^[a-z0-9.-]+\.[a-z0-9-]+$", re.I)
    for member in domain_record.get("include") or []:
        if member.get("addr_type") != "domain" or not member.get("domain"):
            errors.append(f"{TARGET_DOMAIN_OBJECT} contains non-domain member: {member}")
            continue
        domain = str(member["domain"]).lower().rstrip(".")
        if not domain_re.match(domain):
            errors.append(f"{TARGET_DOMAIN_OBJECT} contains invalid domain member: {member}")

    for name in (TARGET_IP_OBJECT, TARGET_DOMAIN_OBJECT):
        policy = policies.get(name)
        if not policy:
            errors.append(f"Missing security policy with expected name: {name}")
            checks[name] = {"object_found": True, "policy_found": False, "policy_references_object": False}
            continue
        text = "\n".join(collect_strings(policy))
        references = name in text
        if not references:
            errors.append(f"Security policy {name} does not reference address object {name}")
        checks[name] = {
            "object_found": True,
            "policy_found": True,
            "policy_references_object": references,
            "reference": targets[name].get("reference"),
            "action": policy.get("action") or policy.get("act") or policy.get("policy_action"),
        }

    if errors:
        raise RuntimeError("Preflight target verification failed: " + " | ".join(errors))
    return checks


def existing_values(record, kind):
    values = []
    for item in record.get("include") or []:
        if kind == "ip" and item.get("addr_type") == "host" and item.get("ip"):
            values.append(str(ipaddress.ip_address(str(item["ip"]))))
        elif kind == "domain" and item.get("addr_type") == "domain" and item.get("domain"):
            values.append(str(item["domain"]).lower().rstrip("."))
    return values


def update_address_object(opener, cfg, before, name, include):
    body = {
        "obj_addr": [
            {
                "name": name,
                "oldname": name,
                "desc": before.get("desc", "") or "",
                "include": include,
                "exclude": before.get("exclude") or [],
            }
        ]
    }
    data = rest_call(opener, cfg, "obj_address", "set_obj_addr_conf", body)
    head = data.get("head", {}) if isinstance(data, dict) else {}
    if head.get("error_code") not in (0, "0"):
        raise RuntimeError(f"set_obj_addr_conf {name} failed: {head}")
    return data


def publish_direct(cfg, targets_payload, apply=False):
    excel_ips = targets_payload["ips"]
    excel_domains = targets_payload["domains"]
    opener = login(cfg)
    try:
        before = read_target_objects(opener, cfg)
        policies = read_policy_objects(opener, cfg)
        preflight = validate_correct_targets(before, policies)

        current_ips = existing_values(before[TARGET_IP_OBJECT], "ip")
        current_domains = existing_values(before[TARGET_DOMAIN_OBJECT], "domain")
        merged_ips = sorted(dict.fromkeys(current_ips + excel_ips), key=ip_sort_key)
        merged_domains = sorted(dict.fromkeys(current_domains + excel_domains))
        add_ips = [item for item in excel_ips if item not in set(current_ips)]
        add_domains = [item for item in excel_domains if item not in set(current_domains)]

        if len(merged_ips) > 512 or len(merged_domains) > 512:
            raise RuntimeError(
                f"Object include limit exceeded: {TARGET_IP_OBJECT}={len(merged_ips)}, "
                f"{TARGET_DOMAIN_OBJECT}={len(merged_domains)}"
            )

        report = {
            "time": dt.datetime.now().isoformat(timespec="seconds"),
            "base_url": cfg["base_url"],
            "transport": "direct",
            "excel": targets_payload.get("excel", {}),
            "preflight": preflight,
            "before": {
                TARGET_IP_OBJECT: len(current_ips),
                TARGET_DOMAIN_OBJECT: len(current_domains),
            },
            "plan": {
                "add_ip_count": len(add_ips),
                "add_domain_count": len(add_domains),
                "final_ip_count": len(merged_ips),
                "final_domain_count": len(merged_domains),
                "add_domains": add_domains,
            },
            "applied": False,
        }

        if apply and (add_ips or add_domains):
            if add_ips:
                update_address_object(
                    opener,
                    cfg,
                    before[TARGET_IP_OBJECT],
                    TARGET_IP_OBJECT,
                    [{"ip": item, "addr_type": "host"} for item in merged_ips],
                )
            if add_domains:
                update_address_object(
                    opener,
                    cfg,
                    before[TARGET_DOMAIN_OBJECT],
                    TARGET_DOMAIN_OBJECT,
                    [{"domain": item, "addr_type": "domain"} for item in merged_domains],
                )
            after = read_target_objects(opener, cfg)
            after_policies = read_policy_objects(opener, cfg)
            report["postflight"] = validate_correct_targets(after, after_policies)
            after_ips = existing_values(after[TARGET_IP_OBJECT], "ip")
            after_domains = existing_values(after[TARGET_DOMAIN_OBJECT], "domain")
            report["applied"] = True
            report["after"] = {
                TARGET_IP_OBJECT: len(after_ips),
                TARGET_DOMAIN_OBJECT: len(after_domains),
            }
            report["verification"] = {
                "missing_ips": sorted(set(excel_ips) - set(after_ips), key=ip_sort_key),
                "missing_domains": sorted(set(excel_domains) - set(after_domains)),
            }
        elif apply:
            report["applied"] = True
            report["after"] = report["before"]
            report["verification"] = {"missing_ips": [], "missing_domains": []}
            report["postflight"] = preflight
        return report
    finally:
        try:
            post_json(
                opener,
                cfg["base_url"].rstrip("/") + "/v1.0/out",
                {"username": cfg["username"]},
                tries=1,
            )
        except Exception:
            pass


def run_command(cmd, cwd=None, **kwargs):
    return subprocess.run(cmd, text=True, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, cwd=cwd, **kwargs)


def extract_json(text):
    start = text.find("{")
    end = text.rfind("}")
    if start < 0 or end < start:
        raise RuntimeError(f"No JSON object found in doops output:\n{text}")
    return json.loads(text[start : end + 1])


def publish_doops(cfg, targets_payload, apply=False):
    target = cfg.get("doops_target") or "zy"
    doops_cwd = cfg.get("doops_cwd")
    session = "nsg_publish_" + dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    temp_root = pathlib.Path(tempfile.mkdtemp(prefix="nsg-doops-publish-"))
    remote_created = False
    try:
        (temp_root / "remote_publish.py").write_text(REMOTE_SCRIPT, encoding="utf-8")
        (temp_root / "connection.local.json").write_text(
            json.dumps({"nsg_firewall": cfg}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        (temp_root / "targets.json").write_text(
            json.dumps(targets_payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

        pushed = run_command(["doops", "-session", session, "push", "--target", target, "--src", str(temp_root)], cwd=doops_cwd)
        if pushed.returncode != 0:
            raise RuntimeError(f"doops push failed:\n{pushed.stdout}")
        remote_created = True

        remote_cmd = f"cd /root/ws/{session} && python3 remote_publish.py {'--apply' if apply else ''}"
        executed = run_command(["doops", "-session", session, "exec", "--target", target, "--cmd", remote_cmd], cwd=doops_cwd, timeout=180)
        if executed.returncode != 0:
            raise RuntimeError(f"doops exec failed:\n{executed.stdout}")

        read = run_command(
            ["doops", "-session", session, "read", "--target", target, "--path", f"/root/ws/{session}/result.json"],
            cwd=doops_cwd,
            timeout=60,
        )
        if read.returncode != 0:
            raise RuntimeError(f"doops read failed:\n{read.stdout}")
        report = extract_json(read.stdout)
        report["transport"] = "doops"
        report["doops"] = {"target": target, "session": session}
        return report
    finally:
        if remote_created:
            run_command(
                [
                    "doops",
                    "-session",
                    session,
                    "clean",
                    "--target",
                    target,
                    "--workspace",
                        session,
                    ],
                    cwd=doops_cwd,
                    timeout=60,
                )
        shutil.rmtree(temp_root, ignore_errors=True)


def build_targets_payload(xlsx):
    excel_ips, excel_domains, normalizations = extract_excel(xlsx)
    return {
        "ips": excel_ips,
        "domains": excel_domains,
        "excel": {
            "path": str(xlsx),
            "ip_count": len(excel_ips),
            "domain_count": len(excel_domains),
            "domain_normalizations": normalizations,
        },
    }


def main():
    parser = argparse.ArgumentParser(description="Publish firewall block IP/domain objects from workbook")
    parser.add_argument("xlsx", type=pathlib.Path)
    parser.add_argument("--config", type=pathlib.Path, default=DEFAULT_CONFIG)
    parser.add_argument("--apply", action="store_true", help="write changes; otherwise only print the plan")
    parser.add_argument(
        "--transport",
        choices=("auto", "direct", "doops"),
        default="doops",
        help="connection transport; doops is the local default; auto tries direct first and falls back to doops",
    )
    args = parser.parse_args()

    cfg = load_config(args.config)
    targets_payload = build_targets_payload(args.xlsx)
    report = None
    direct_error = None

    if args.transport in ("auto", "direct"):
        try:
            report = publish_direct(cfg, targets_payload, apply=args.apply)
        except Exception as exc:
            direct_error = str(exc)
            if args.transport == "direct":
                raise

    if report is None:
        report = publish_doops(cfg, targets_payload, apply=args.apply)
        if direct_error:
            report["direct_error"] = direct_error

    print(json.dumps(report, ensure_ascii=False, indent=2))
    verification = report.get("verification", {})
    if verification.get("missing_ips") or verification.get("missing_domains"):
        raise SystemExit(3)


if __name__ == "__main__":
    main()
