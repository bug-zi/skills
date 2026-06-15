from __future__ import annotations

import argparse
import copy
import html
import json
import shutil
import subprocess
import sys
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

from build_adhoc_inventory import build_inventory
from health_common import (
    STATUS_ORDER,
    as_int,
    ensure_dir,
    is_valid_url,
    normalize_url,
    now_iso,
    read_json,
    slugify,
    write_json,
)
from scoring_model import health_status_from_score_and_findings, score_from_finding_clusters

SCRIPT_DIR = Path(__file__).resolve().parent
ROOT = SCRIPT_DIR.parent
DELIVERY_BATCH_DIR = "delivery-batch"
SITE_DELIVERY_HTML = Path("delivery") / "final-report.html"


HEADER_ALIASES = {
    "name": {"网站名称", "站点名称", "系统名称", "名称", "name", "systemname"},
    "url": {"url", "网址", "网站地址", "入口地址", "链接", "link", "websiteurl"},
    "owner": {"负责人", "责任人", "责任部门", "owner"},
    "priority": {"优先级", "重要性", "priority"},
    "remark": {"备注", "说明", "remark", "note"},
    "expected_status": {"期望状态码", "预期状态码", "expectedstatus", "status"},
    "timeout_ms": {"超时时间", "超时时间毫秒", "timeout", "timeoutms"},
}

STANDARD_SHEETS = {"系统清单", "网站页面"}
PRIORITIES = {"critical", "high", "medium", "low"}
STATUS_LABELS = {
    "healthy": "正常",
    "warning": "告警",
    "critical": "严重",
    "unknown": "未知",
}


def resolved_path(path: str | Path) -> Path:
    return Path(path).expanduser().resolve()


def fs_path(path: str | Path) -> str:
    target = Path(path).expanduser().resolve()
    text = str(target)
    if sys.platform != "win32" or text.startswith("\\\\?\\"):
        return text
    if text.startswith("\\\\"):
        return "\\\\?\\UNC\\" + text.lstrip("\\")
    return "\\\\?\\" + text


def remove_tree(path: Path) -> None:
    if path.exists():
        shutil.rmtree(fs_path(path))


def copy_tree(source: Path, target: Path) -> None:
    remove_tree(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copytree(fs_path(source), fs_path(target))


def require_openpyxl():
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as exc:
        raise SystemExit("Missing dependency: openpyxl. Install it with `python -m pip install openpyxl` to read Excel inventories.") from exc
    return load_workbook


def import_standard_excel(path: str | Path) -> tuple[dict, dict]:
    try:
        from import_inventory_excel import import_excel
    except ImportError as exc:
        raise SystemExit("Missing dependency: openpyxl. Install it with `python -m pip install openpyxl` to read Excel inventories.") from exc
    return import_excel(path)


@dataclass
class SiteEntry:
    index: int
    name: str
    url: str
    owner: str = "管理员"
    priority: str = "medium"
    remark: str = ""
    expected_status: int = 200
    timeout_ms: int = 5000
    source: str = "excel"
    error: str | None = None


@dataclass
class SiteRunResult:
    index: int
    slug: str
    name: str
    url: str
    owner: str
    run_dir: Path
    delivery_entry: Path | None
    status: str = "unknown"
    score: int = 0
    findings_count: int = 0
    error: str | None = None
    report: dict | None = None
    quality_status: str = "unknown"


def normalize_header(value) -> str:
    return str(value or "").strip().lower().replace(" ", "").replace("_", "")


def header_lookup(headers: list[str], key: str) -> str | None:
    aliases = {normalize_header(item) for item in HEADER_ALIASES[key]}
    for header in headers:
        if normalize_header(header) in aliases:
            return header
    return None


def row_value(row: dict, header: str | None, default: str = ""):
    if not header:
        return default
    value = row.get(header)
    return default if value is None else value


def first_simple_sheet(workbook) -> tuple[object, list[str]]:
    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        if not rows:
            continue
        headers = [str(cell or "").strip() for cell in rows[0]]
        if header_lookup(headers, "url"):
            return worksheet, headers
    raise SystemExit("未找到包含 URL/网址/网站地址列的简单网站清单 Sheet。")


def load_site_entries_from_excel(excel_path: str | Path) -> list[SiteEntry]:
    load_workbook = require_openpyxl()
    workbook = load_workbook(excel_path, data_only=True)
    worksheet, headers = first_simple_sheet(workbook)
    columns = {key: header_lookup(headers, key) for key in HEADER_ALIASES}
    entries: list[SiteEntry] = []

    for row_number, cells in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        row = {header: value for header, value in zip(headers, cells) if header}
        if not any(value not in (None, "") for value in row.values()):
            continue

        raw_url = str(row_value(row, columns["url"], "") or "").strip()
        raw_name = str(row_value(row, columns["name"], "") or "").strip()
        owner = str(row_value(row, columns["owner"], "管理员") or "管理员").strip()
        remark = str(row_value(row, columns["remark"], "") or "").strip()
        priority = str(row_value(row, columns["priority"], "medium") or "medium").strip().lower()
        if priority not in PRIORITIES:
            priority = "medium"
        expected_status = as_int(row_value(row, columns["expected_status"], 200), 200)
        timeout_ms = as_int(row_value(row, columns["timeout_ms"], 5000), 5000)

        error = None
        normalized_url = ""
        if not raw_url:
            error = "URL 为空"
        else:
            normalized_url = normalize_url(raw_url)
            if not is_valid_url(normalized_url):
                error = f"URL 无效: {raw_url}"

        parsed = urlparse(normalized_url)
        name = raw_name or parsed.netloc or f"第 {row_number} 行网站"
        entries.append(
            SiteEntry(
                index=len(entries) + 1,
                name=name,
                url=normalized_url,
                owner=owner,
                priority=priority,
                remark=remark,
                expected_status=expected_status,
                timeout_ms=timeout_ms,
                error=error,
            )
        )
    return entries


def workbook_looks_standard(excel_path: str | Path) -> bool:
    load_workbook = require_openpyxl()
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        return STANDARD_SHEETS.issubset(set(workbook.sheetnames))
    finally:
        workbook.close()


def site_slug(index: int, name: str, url: str) -> str:
    parsed = urlparse(url)
    base = parsed.netloc.replace(".", "-") if parsed.netloc else name
    return f"{index:03d}-{slugify(base, 'site')}"


def build_inventory_for_entry(entry: SiteEntry, *, discovery: bool = True) -> dict | None:
    if entry.error:
        return None
    inventory = build_inventory(entry.url, name=entry.name, discovery=discovery)
    system = inventory["systems"][0]
    system["priority"] = entry.priority
    system["owner"] = entry.owner
    system["remark"] = entry.remark
    website = system["websites"][0]
    website["expected_status"] = entry.expected_status
    website["timeout_ms"] = entry.timeout_ms
    return inventory


def split_inventory_by_website(inventory: dict) -> list[tuple[SiteEntry, dict | None]]:
    pairs: list[tuple[SiteEntry, dict | None]] = []
    global_config = copy.deepcopy(inventory.get("global", {}))
    source = copy.deepcopy(inventory.get("source", {}))

    for system in inventory.get("systems", []) or []:
        websites = system.get("websites", []) or []
        if not websites:
            entry = SiteEntry(
                index=len(pairs) + 1,
                name=str(system.get("name") or system.get("id") or f"站点 {len(pairs) + 1}"),
                url="",
                owner=str(system.get("owner") or "管理员"),
                priority=str(system.get("priority") or "medium"),
                remark=str(system.get("remark") or ""),
                source="standard-excel",
                error="系统未配置网站页面",
            )
            pairs.append((entry, None))
            continue

        for page in websites:
            page_name = str(page.get("name") or page.get("page_id") or page.get("url") or "").strip()
            system_name = str(system.get("name") or system.get("id") or "").strip()
            if len(websites) == 1:
                name = system_name or page_name
            else:
                name = f"{system_name} - {page_name}" if system_name and page_name else system_name or page_name
            url = str(page.get("url") or "").strip()
            normalized_url = normalize_url(url) if url else ""
            error = None if normalized_url and is_valid_url(normalized_url) else "URL 为空或无效"

            entry = SiteEntry(
                index=len(pairs) + 1,
                name=name or f"站点 {len(pairs) + 1}",
                url=normalized_url,
                owner=str(system.get("owner") or "管理员"),
                priority=str(system.get("priority") or "medium"),
                remark=str(system.get("remark") or ""),
                expected_status=as_int(page.get("expected_status"), 200),
                timeout_ms=as_int(page.get("timeout_ms"), 5000),
                source="standard-excel",
                error=error,
            )
            if error:
                pairs.append((entry, None))
                continue

            system_copy = copy.deepcopy(system)
            system_copy["id"] = slugify(
                f"{system.get('id', '')}-{page.get('page_id', '')}",
                f"site-{entry.index:03d}",
            )
            system_copy["name"] = entry.name
            system_copy["websites"] = [copy.deepcopy(page)]
            site_inventory = {
                "generated_at": now_iso(),
                "source": source,
                "systems": [system_copy],
                "global": copy.deepcopy(global_config),
            }
            pairs.append((entry, site_inventory))
    return pairs


def load_site_inventory_pairs(excel_path: str | Path, checks: str) -> list[tuple[SiteEntry, dict | None]]:
    if workbook_looks_standard(excel_path):
        inventory, report = import_standard_excel(excel_path)
        if report.get("status") != "passed" and not inventory.get("systems"):
            messages = "；".join(str(item.get("message")) for item in report.get("messages", [])[:5])
            raise SystemExit(f"标准资源清单导入失败：{messages}")
        return split_inventory_by_website(inventory)

    discovery = "discovery" in {item.strip() for item in checks.split(",") if item.strip()}
    return [(entry, build_inventory_for_entry(entry, discovery=discovery)) for entry in load_site_entries_from_excel(excel_path)]


def checks_for_inspect(raw: str) -> str:
    requested = {item.strip() for item in raw.split(",") if item.strip()}
    requested.update({"merge", "render"})
    order = ["discovery", "content", "deep-browser", "browser", "host", "merge", "render"]
    ordered = [item for item in order if item in requested]
    ordered.extend(sorted(requested - set(order)))
    return ",".join(ordered)


def command_text(command: list[str]) -> str:
    return " ".join(f'"{item}"' if " " in str(item) else str(item) for item in command)


def run_command(command: list[str], cwd: Path, log_path: Path) -> subprocess.CompletedProcess[str]:
    result = subprocess.run(command, cwd=str(cwd), check=False, capture_output=True, text=True)
    log_path.parent.mkdir(parents=True, exist_ok=True)
    with open(log_path, "a", encoding="utf-8") as handle:
        handle.write(f"$ {command_text(command)}\n")
        handle.write(f"exit_code={result.returncode}\n")
        if result.stdout:
            handle.write("\n[stdout]\n" + result.stdout)
        if result.stderr:
            handle.write("\n[stderr]\n" + result.stderr)
        handle.write("\n\n")
    return result


def quality_status_from_audit(run_dir: Path, returncode: int) -> str:
    audit_json = run_dir / "quality-audit.json"
    if not audit_json.exists():
        return "failed" if returncode != 0 else "unknown"
    try:
        status = str(read_json(audit_json).get("status") or "unknown")
    except Exception:
        return "failed"
    if status == "pass":
        return "passed"
    if status == "warning":
        return "warning"
    if status == "fail":
        return "failed"
    return status


def site_result_from_error(entry: SiteEntry, slug: str, run_dir: Path, error: str) -> SiteRunResult:
    ensure_dir(run_dir)
    write_json(
        run_dir / "site-error.json",
        {
            "generated_at": now_iso(),
            "site": {"name": entry.name, "url": entry.url, "owner": entry.owner},
            "error": error,
        },
    )
    return SiteRunResult(
        index=entry.index,
        slug=slug,
        name=entry.name,
        url=entry.url,
        owner=entry.owner,
        run_dir=run_dir,
        delivery_entry=None,
        status="unknown",
        score=0,
        findings_count=1,
        error=error,
        quality_status="not-run",
    )


def inspect_one_site(
    entry: SiteEntry,
    inventory: dict | None,
    *,
    batch_dir: Path,
    checks: str,
    guide_mode: str,
    skip_ai_guide_image: bool,
) -> SiteRunResult:
    slug = site_slug(entry.index, entry.name, entry.url)
    run_dir = batch_dir / "sites" / slug
    log_path = run_dir / "run.log"

    if entry.error or inventory is None:
        return site_result_from_error(entry, slug, run_dir, entry.error or "站点清单不可用")

    ensure_dir(run_dir)
    inventory_path = write_json(run_dir / "inventory.normalized.json", inventory)
    inspect_cmd = [
        sys.executable,
        str(SCRIPT_DIR / "inspect_sites.py"),
        "--inventory",
        str(inventory_path),
        "--out",
        str(run_dir),
        "--checks",
        checks_for_inspect(checks),
    ]
    if skip_ai_guide_image:
        inspect_cmd.append("--skip-ai-guide-image")
    else:
        inspect_cmd.extend(["--guide-mode", guide_mode])
    inspect_result = run_command(inspect_cmd, ROOT, log_path)
    if inspect_result.returncode != 0:
        return site_result_from_error(entry, slug, run_dir, f"巡检脚本退出码 {inspect_result.returncode}")

    final_report = run_dir / "final-report.json"
    if not final_report.exists():
        return site_result_from_error(entry, slug, run_dir, "未生成 final-report.json")

    audit_cmd = [
        sys.executable,
        str(SCRIPT_DIR / "audit_report_quality.py"),
        "--run-dir",
        str(run_dir),
        "--iteration",
        "1",
        "--out",
        str(run_dir / "quality-audit.md"),
    ]
    audit_result = run_command(audit_cmd, ROOT, log_path)
    quality_status = quality_status_from_audit(run_dir, audit_result.returncode)

    report = read_json(final_report)
    delivery_entry = run_dir / SITE_DELIVERY_HTML
    findings_count = len(report.get("priority_findings", []) or [])
    error = None
    if not delivery_entry.exists():
        error = "未生成 delivery/final-report.html"
        quality_status = "failed"
    return SiteRunResult(
        index=entry.index,
        slug=slug,
        name=entry.name,
        url=entry.url,
        owner=entry.owner,
        run_dir=run_dir,
        delivery_entry=delivery_entry if delivery_entry.exists() else None,
        status=str(report.get("overall_status") or "unknown"),
        score=as_int(report.get("score"), 0),
        findings_count=findings_count,
        error=error,
        report=report,
        quality_status=quality_status,
    )


def failure_finding(result: SiteRunResult, index: int) -> dict:
    return {
        "id": f"batch-failure-{index:03d}",
        "system_id": result.slug,
        "system_name": result.name,
        "severity": "medium",
        "priority": "P2",
        "title": "站点巡检未完成",
        "business_impact": "该站点本轮未形成完整自动化巡检结论，整体健康判断需要补充复测。",
        "technical_evidence": result.error or "unknown error",
        "target": result.url,
        "recommendation": "补齐 URL 或修复巡检运行环境后重新执行该站点巡检。",
        "evidence_files": [],
    }


def site_result_record(result: SiteRunResult) -> dict:
    return {
        "index": result.index,
        "slug": result.slug,
        "name": result.name,
        "url": result.url,
        "owner": result.owner,
        "status": result.status,
        "score": result.score,
        "findings": result.findings_count,
        "quality_status": result.quality_status,
        "error": result.error or "",
        "report": delivery_report_href(result),
    }


def reviewed_site_result(result: SiteRunResult) -> SiteRunResult:
    if result.error or not result.report:
        return result
    if "score" not in result.report and "overall_status" not in result.report:
        return result
    findings = result.report.get("priority_findings", []) or []
    has_reports = bool(result.report.get("module_reports") or findings)
    score = score_from_finding_clusters(findings, unknown_penalty=0 if has_reports else 20)
    status = health_status_from_score_and_findings(score, findings, has_reports=has_reports)
    reviewed = copy.copy(result)
    reviewed.score = score
    reviewed.status = status
    reviewed.report = copy.deepcopy(result.report)
    reviewed.report["score"] = score
    reviewed.report["overall_status"] = status
    return reviewed


def quality_failed_count(results: list[SiteRunResult]) -> int:
    return sum(1 for item in results if item.error or item.quality_status in {"failed", "not-run"})


def priority_site_records(results: list[SiteRunResult]) -> list[dict]:
    candidates = [item for item in results if item.status != "healthy" or item.error or item.quality_status in {"failed", "not-run"}]
    ordered = sorted(
        candidates,
        key=lambda item: (
            STATUS_ORDER.get(item.status, 2),
            1 if item.error else 0,
            -item.findings_count,
            -item.score,
        ),
        reverse=True,
    )
    records = []
    for item in ordered:
        if item.error:
            reason = item.error
        elif item.status != "healthy":
            reason = f"{STATUS_LABELS.get(item.status, item.status)}状态，{item.findings_count} 个优先问题"
        elif item.quality_status in {"failed", "not-run"}:
            reason = f"质量审计状态为 {item.quality_status}"
        else:
            reason = "建议例行复核"
        record = site_result_record(item)
        record["reason"] = reason
        records.append(record)
    return records


def common_issue_families(findings: list[dict]) -> list[dict]:
    families: dict[str, dict] = {}
    for finding in findings:
        title = str(finding.get("title") or "未命名问题")
        family = families.setdefault(
            title,
            {
                "title": title,
                "classification": "批次共性风险",
                "site_ids": set(),
                "sites": [],
                "owners": set(),
                "finding_count": 0,
                "severity": str(finding.get("severity") or "low"),
                "priority": str(finding.get("priority") or "P4"),
                "next_action": str(finding.get("recommendation") or "安排责任人复核并形成闭环。"),
            },
        )
        family["finding_count"] += 1
        sid = str(finding.get("system_id") or "")
        if sid:
            family["site_ids"].add(sid)
        site_name = str(finding.get("system_name") or sid or "未知站点")
        if site_name and site_name not in family["sites"]:
            family["sites"].append(site_name)
        owner = str(finding.get("owner") or "")
        if owner:
            family["owners"].add(owner)
        if str(finding.get("severity") or "") == "critical":
            family["severity"] = "critical"
        if str(finding.get("priority") or "") in {"P1", "P2"}:
            family["priority"] = str(finding.get("priority"))
        if title == "站点巡检未完成":
            family["classification"] = "批次覆盖缺口"
    rows = []
    for family in families.values():
        site_count = len(family["site_ids"]) or len(family["sites"])
        rows.append(
            {
                "title": family["title"],
                "classification": family["classification"],
                "site_count": site_count,
                "finding_count": family["finding_count"],
                "severity": family["severity"],
                "priority": family["priority"],
                "owners": sorted(family["owners"]),
                "sites": family["sites"][:8],
                "next_action": family["next_action"],
            }
        )
    return sorted(
        rows,
        key=lambda item: (
            1 if item["title"] == "站点巡检未完成" else 0,
            item["site_count"],
            item["finding_count"],
        ),
        reverse=True,
    )


def build_overall_report(results: list[SiteRunResult]) -> dict:
    results = [reviewed_site_result(item) for item in results]
    counts = {
        "total_systems": len(results),
        "healthy": sum(1 for item in results if item.status == "healthy"),
        "warning": sum(1 for item in results if item.status == "warning"),
        "critical": sum(1 for item in results if item.status == "critical"),
        "unknown": sum(1 for item in results if item.status == "unknown"),
    }
    if not results:
        overall_status = "unknown"
        average_score = 0
    else:
        overall_status = max((item.status for item in results), key=lambda value: STATUS_ORDER.get(value, 2))
        average_score = int(sum(item.score for item in results) / len(results))

    systems = []
    findings = []
    for result in results:
        systems.append(
            {
                "system_id": result.slug,
                "system_name": result.name,
                "status": result.status,
                "score": result.score,
                "findings": result.findings_count,
                "owner": result.owner,
                "url": result.url,
                "quality_status": result.quality_status,
            }
        )
        if result.error:
            findings.append(failure_finding(result, len(findings) + 1))

    for result in results:
        for finding in (result.report or {}).get("priority_findings", []) or []:
            copied = copy.deepcopy(finding)
            copied["id"] = f"{result.slug}-{copied.get('id') or len(findings) + 1}"
            copied["system_id"] = result.slug
            copied["system_name"] = result.name
            copied["evidence_files"] = []
            findings.append(copied)

    summary_text = (
        f"本次批量巡检覆盖 {len(results)} 个网站，"
        f"正常 {counts['healthy']} 个、告警 {counts['warning']} 个、"
        f"严重 {counts['critical']} 个、未知 {counts['unknown']} 个，"
        f"平均健康评分 {average_score}/100。"
    )
    if findings:
        summary_text += " 建议优先处理无法完成巡检、严重告警和影响访问的站点。"
    else:
        summary_text += " 未发现需要优先处理的问题。"

    batch_metrics = {
        "site_count": len(results),
        "average_score": average_score,
        "lowest_score": min((item.score for item in results), default=0),
        "failed": sum(1 for item in results if item.error),
        "quality_failed": quality_failed_count(results),
        "status_distribution": counts,
    }
    site_results = [site_result_record(item) for item in results]
    priority_sites = priority_site_records(results)
    issue_families = common_issue_families(findings)

    return {
        "generated_at": now_iso(),
        "report_kind": "multi-site-overall",
        "overall_status": overall_status,
        "score": average_score,
        "executive_summary": summary_text,
        "summary": counts,
        "batch_metrics": batch_metrics,
        "systems": systems,
        "site_results": site_results,
        "priority_sites": priority_sites,
        "common_issue_families": issue_families,
        "priority_findings": findings,
        "module_reports": [
            {
                "report_type": "multi-site-batch",
                "generated_at": now_iso(),
                "status": overall_status,
                "score": average_score,
                "summary": summary_text,
                "systems": systems,
                "findings": findings,
                "metrics": {
                    **counts,
                    "quality_failed": batch_metrics["quality_failed"],
                    "failed": sum(1 for item in results if item.error),
                },
            }
        ],
        "artifacts": {},
    }


def delivery_report_href(result: SiteRunResult) -> str:
    if not result.delivery_entry:
        return ""
    return f"sites/{result.slug}/delivery/final-report.html"


def build_batch_summary(results: list[SiteRunResult], batch_dir: Path, delivery_dir: Path) -> dict:
    results = [reviewed_site_result(item) for item in results]
    return {
        "generated_at": now_iso(),
        "batch_dir": str(batch_dir),
        "delivery_dir": str(delivery_dir),
        "counts": {
            "total": len(results),
            "healthy": sum(1 for item in results if item.status == "healthy"),
            "warning": sum(1 for item in results if item.status == "warning"),
            "critical": sum(1 for item in results if item.status == "critical"),
            "unknown": sum(1 for item in results if item.status == "unknown"),
            "failed": sum(1 for item in results if item.error),
            "quality_failed": sum(1 for item in results if item.quality_status == "failed"),
        },
        "sites": [
            {
                "index": item.index,
                "name": item.name,
                "url": item.url,
                "owner": item.owner,
                "status": item.status,
                "score": item.score,
                "findings": item.findings_count,
                "quality_status": item.quality_status,
                "error": item.error,
                "report": delivery_report_href(item),
            }
            for item in results
        ],
    }


def render_batch_index(results: list[SiteRunResult], *, overall_href: str | None) -> str:
    results = [reviewed_site_result(item) for item in results]
    counts = {
        "total": len(results),
        "healthy": sum(1 for item in results if item.status == "healthy"),
        "warning": sum(1 for item in results if item.status == "warning"),
        "critical": sum(1 for item in results if item.status == "critical"),
        "unknown": sum(1 for item in results if item.status == "unknown"),
    }
    rows = []
    for item in results:
        report_href = delivery_report_href(item)
        report_link = f'<a href="{html.escape(report_href)}">打开报告</a>' if report_href else '<span class="muted">未生成</span>'
        rows.append(
            "<tr>"
            f"<td>{item.index}</td>"
            f"<td><strong>{html.escape(item.name)}</strong><br><span>{html.escape(item.url or '未提供 URL')}</span></td>"
            f"<td>{html.escape(item.owner or '未指定')}</td>"
            f'<td><span class="badge {html.escape(item.status)}">{html.escape(STATUS_LABELS.get(item.status, item.status))}</span></td>'
            f"<td>{item.score}</td>"
            f"<td>{item.findings_count}</td>"
            f"<td>{html.escape(item.quality_status)}</td>"
            f"<td>{report_link}</td>"
            "</tr>"
        )
    overall_link = (
        f'<a class="primary-link" href="{html.escape(overall_href)}">打开整体总报告</a>'
        if overall_href
        else '<span class="muted">单站点巡检未生成独立整体总报告</span>'
    )
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>多站点网站健康巡检交付索引</title>
  <style>
    :root {{
      --paper: #f7f8fb;
      --ink: #172033;
      --muted: #667085;
      --line: #d8dee9;
      --panel: #ffffff;
      --accent: #1f6feb;
      --healthy: #0f766e;
      --warning: #a16207;
      --critical: #b42318;
      --unknown: #475467;
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      background: var(--paper);
      color: var(--ink);
      font-family: "Microsoft YaHei", "Segoe UI", Arial, sans-serif;
      line-height: 1.55;
    }}
    main {{ max-width: 1180px; margin: 0 auto; padding: 32px 20px 48px; }}
    header {{ display: flex; justify-content: space-between; gap: 24px; align-items: flex-end; margin-bottom: 24px; }}
    h1 {{ margin: 0 0 8px; font-size: 28px; letter-spacing: 0; }}
    p {{ margin: 0; color: var(--muted); }}
    .primary-link {{
      display: inline-flex;
      align-items: center;
      min-height: 40px;
      padding: 8px 14px;
      border-radius: 6px;
      background: var(--accent);
      color: #fff;
      text-decoration: none;
      font-weight: 700;
      white-space: nowrap;
    }}
    .stats {{ display: grid; grid-template-columns: repeat(5, minmax(0, 1fr)); gap: 12px; margin-bottom: 20px; }}
    .stat {{ background: var(--panel); border: 1px solid var(--line); border-radius: 8px; padding: 14px; }}
    .stat b {{ display: block; font-size: 24px; }}
    .stat span {{ color: var(--muted); font-size: 13px; }}
    .table-wrap {{ overflow-x: auto; background: var(--panel); border: 1px solid var(--line); border-radius: 8px; }}
    table {{ width: 100%; border-collapse: collapse; min-width: 920px; }}
    th, td {{ padding: 12px 14px; border-bottom: 1px solid var(--line); text-align: left; vertical-align: top; }}
    th {{ background: #eef2f7; font-size: 13px; color: #344054; }}
    td span {{ color: var(--muted); font-size: 12px; }}
    a {{ color: var(--accent); font-weight: 700; }}
    .badge {{ display: inline-block; min-width: 48px; padding: 3px 8px; border-radius: 999px; color: #fff; text-align: center; font-size: 12px; }}
    .healthy {{ background: var(--healthy); }}
    .warning {{ background: var(--warning); }}
    .critical {{ background: var(--critical); }}
    .unknown {{ background: var(--unknown); }}
    .muted {{ color: var(--muted); }}
    @media (max-width: 760px) {{
      header {{ display: block; }}
      header > div:last-child {{ margin-top: 14px; }}
      .stats {{ grid-template-columns: repeat(2, minmax(0, 1fr)); }}
    }}
  </style>
</head>
<body>
  <main>
    <header>
      <div>
        <h1>多站点网站健康巡检交付索引</h1>
        <p>生成时间：{html.escape(now_iso())}</p>
      </div>
      <div>{overall_link}</div>
    </header>
    <section class="stats" aria-label="巡检统计">
      <div class="stat"><b>{counts["total"]}</b><span>网站总数</span></div>
      <div class="stat"><b>{counts["healthy"]}</b><span>正常</span></div>
      <div class="stat"><b>{counts["warning"]}</b><span>告警</span></div>
      <div class="stat"><b>{counts["critical"]}</b><span>严重</span></div>
      <div class="stat"><b>{counts["unknown"]}</b><span>未知</span></div>
    </section>
    <section class="table-wrap">
      <table>
        <thead>
          <tr>
            <th>序号</th>
            <th>网站</th>
            <th>负责人</th>
            <th>状态</th>
            <th>评分</th>
            <th>问题数</th>
            <th>质量审计</th>
            <th>报告</th>
          </tr>
        </thead>
        <tbody>
          {"".join(rows)}
        </tbody>
      </table>
    </section>
  </main>
</body>
</html>
"""


def write_batch_run_log(results: list[SiteRunResult], target: Path) -> None:
    lines = [
        "# 多站点巡检批次日志",
        "",
        f"- 生成时间：`{now_iso()}`",
        f"- 网站数量：{len(results)}",
        f"- 失败数量：{sum(1 for item in results if item.error)}",
        "",
        "## 站点结果",
        "",
    ]
    for item in results:
        suffix = f"；错误：{item.error}" if item.error else ""
        lines.append(f"- {item.index}. {item.name}：{STATUS_LABELS.get(item.status, item.status)}，评分 {item.score}{suffix}")
    target.write_text("\n".join(lines) + "\n", encoding="utf-8-sig")


def render_overall_report(batch_dir: Path, results: list[SiteRunResult], guide_mode: str, skip_ai_guide_image: bool) -> Path | None:
    if len(results) <= 1:
        return None
    overall_dir = ensure_dir(batch_dir / "overall")
    report_path = write_json(overall_dir / "final-report.json", build_overall_report(results))
    render_cmd = [
        sys.executable,
        str(SCRIPT_DIR / "render_report.py"),
        "--report",
        str(report_path),
        "--out",
        str(overall_dir),
    ]
    if skip_ai_guide_image:
        render_cmd.append("--skip-ai-guide-image")
    else:
        render_cmd.extend(["--guide-mode", guide_mode])
    result = run_command(render_cmd, ROOT, overall_dir / "run.log")
    if result.returncode != 0:
        return None

    audit_cmd = [
        sys.executable,
        str(SCRIPT_DIR / "audit_report_quality.py"),
        "--run-dir",
        str(overall_dir),
        "--iteration",
        "1",
        "--out",
        str(overall_dir / "quality-audit.md"),
    ]
    run_command(audit_cmd, ROOT, overall_dir / "run.log")
    return overall_dir / SITE_DELIVERY_HTML


def prepare_delivery_batch(batch_dir: Path, results: list[SiteRunResult], overall_entry: Path | None) -> Path:
    delivery_dir = batch_dir / DELIVERY_BATCH_DIR
    remove_tree(delivery_dir)
    ensure_dir(delivery_dir)

    if overall_entry and overall_entry.exists():
        copy_tree(overall_entry.parent, delivery_dir / "overall")

    for result in results:
        if result.delivery_entry and result.delivery_entry.exists():
            copy_tree(result.delivery_entry.parent, delivery_dir / "sites" / result.slug / "delivery")

    overall_href = "overall/final-report.html" if overall_entry and overall_entry.exists() else None
    (delivery_dir / "index.html").write_text(render_batch_index(results, overall_href=overall_href), encoding="utf-8")
    summary = build_batch_summary(results, batch_dir, delivery_dir)
    write_json(delivery_dir / "batch-summary.json", summary)
    write_json(batch_dir / "batch-summary.json", summary)
    write_batch_run_log(results, delivery_dir / "batch-run-log.md")
    return delivery_dir


def run_all_sites(
    pairs: list[tuple[SiteEntry, dict | None]],
    *,
    batch_dir: Path,
    max_workers: int,
    checks: str,
    guide_mode: str,
    skip_ai_guide_image: bool,
    continue_on_error: bool,
) -> list[SiteRunResult]:
    results: list[SiteRunResult] = []
    with ThreadPoolExecutor(max_workers=max(1, max_workers)) as executor:
        futures = {
            executor.submit(
                inspect_one_site,
                entry,
                inventory,
                batch_dir=batch_dir,
                checks=checks,
                guide_mode=guide_mode,
                skip_ai_guide_image=skip_ai_guide_image,
            ): entry
            for entry, inventory in pairs
        }
        for future in as_completed(futures):
            entry = futures[future]
            try:
                result = future.result()
            except Exception as exc:  # noqa: BLE001
                slug = site_slug(entry.index, entry.name, entry.url)
                if not continue_on_error:
                    raise
                result = site_result_from_error(entry, slug, batch_dir / "sites" / slug, str(exc))
            if result.error and not continue_on_error:
                raise SystemExit(result.error)
            results.append(result)
    return sorted(results, key=lambda item: item.index)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run N+1 reports for multiple website health inspections.")
    parser.add_argument("--excel", required=True, help="简单 URL Excel 或标准资源清单 Excel。")
    parser.add_argument("--out", required=True, help="批次输出目录。")
    parser.add_argument("--max-workers", type=int, default=5, help="并发站点数，默认 5。")
    parser.add_argument("--checks", default="discovery,content,deep-browser", help="检查模块，默认公开页巡检。")
    parser.add_argument(
        "--guide-mode",
        choices=["auto", "ai-image", "html-guide", "reuse-existing"],
        default="auto",
        help="导读图生成方式；缺少 API Key 时 auto 会要求先确认用户选择。",
    )
    parser.add_argument("--skip-ai-guide-image", action=argparse.BooleanOptionalAction, default=False)
    parser.add_argument("--continue-on-error", action=argparse.BooleanOptionalAction, default=True)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    batch_dir = ensure_dir(resolved_path(args.out))
    pairs = load_site_inventory_pairs(resolved_path(args.excel), args.checks)
    if not pairs:
        raise SystemExit("Excel 中没有可巡检的网站记录。")

    results = run_all_sites(
        pairs,
        batch_dir=batch_dir,
        max_workers=args.max_workers,
        checks=args.checks,
        guide_mode=args.guide_mode,
        skip_ai_guide_image=args.skip_ai_guide_image,
        continue_on_error=args.continue_on_error,
    )
    overall_entry = render_overall_report(batch_dir, results, args.guide_mode, args.skip_ai_guide_image)
    delivery_dir = prepare_delivery_batch(batch_dir, results, overall_entry)
    print(delivery_dir / "index.html")
    if overall_entry:
        print(delivery_dir / "overall" / "final-report.html")


if __name__ == "__main__":
    main()
