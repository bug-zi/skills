from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation


SHEETS = {
    "系统清单": [
        ["system_id", "system_name", "priority", "owner", "environment", "enabled", "remark"],
        ["admin-platform", "统一管理平台", "critical", "运维组", "prod", "yes", "核心后台"],
    ],
    "网站页面": [
        ["system_id", "page_id", "page_name", "url", "expected_status", "timeout_ms", "enabled"],
        ["admin-platform", "login", "登录页", "https://example.com/login", 200, 5000, "yes"],
    ],
    "内容检查规则": [
        ["system_id", "page_id", "required_text", "forbidden_text", "min_length", "rule_enabled"],
        ["admin-platform", "login", "登录;统一管理平台", "404;500;系统维护;系统繁忙", 300, "yes"],
    ],
    "管理员功能流程": [
        ["flow_id", "system_id", "flow_name", "step_order", "action", "target", "value", "expect", "enabled"],
        ["admin-login", "admin-platform", "管理员登录检查", 1, "goto", "https://example.com/login", "", "", "yes"],
        ["admin-login", "admin-platform", "管理员登录检查", 2, "fill", "#username", "${ADMIN_USER}", "", "yes"],
        ["admin-login", "admin-platform", "管理员登录检查", 3, "fill", "#password", "${ADMIN_PASS}", "", "yes"],
        ["admin-login", "admin-platform", "管理员登录检查", 4, "click", "button[type=submit]", "", "", "yes"],
        ["admin-login", "admin-platform", "管理员登录检查", 5, "expect_url_contains", "", "", "/dashboard", "yes"],
    ],
    "主机清单": [
        ["host_id", "system_id", "host_name", "ip", "port", "auth_ref", "enabled"],
        ["app-01", "admin-platform", "应用服务器01", "10.0.0.11", 22, "prod-ssh-admin", "yes"],
    ],
    "微服务清单": [
        ["service_id", "system_id", "host_id", "service_name", "check_type", "check_target", "expected", "enabled"],
        ["admin-api", "admin-platform", "app-01", "后台 API", "http", "http://127.0.0.1:8080/health", "200", "yes"],
        ["nginx", "admin-platform", "app-01", "Nginx", "port", "443", "listening", "yes"],
    ],
    "凭证引用": [
        ["auth_ref", "auth_type", "env_keys", "remark"],
        ["prod-ssh-admin", "ssh", "SSH_USER,SSH_KEY_PATH", "生产 SSH"],
        ["admin-login", "browser_login", "ADMIN_USER,ADMIN_PASS", "管理员测试账号"],
    ],
    "全局配置": [
        ["key", "value"],
        ["default_timeout_ms", "5000"],
        ["max_parallel", "5"],
        ["report_language", "zh-CN"],
        ["allow_mutation", "false"],
        ["screenshot_on_failure", "true"],
        ["crawler_enabled", "false"],
        ["crawler_max_depth", "1"],
        ["crawler_max_pages", "20"],
    ],
}


def style_sheet(ws):
    fill = PatternFill("solid", fgColor="1B365D")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 4, 14), 42)


def create_template(output: str | Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in SHEETS.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
        style_sheet(ws)
    yes_no = DataValidation(type="list", formula1='"yes,no"', allow_blank=False)
    priority = DataValidation(type="list", formula1='"critical,high,medium,low"', allow_blank=False)
    actions = DataValidation(type="list", formula1='"goto,click,fill,wait_for,expect_text,expect_url_contains,expect_selector"', allow_blank=True)
    check_types = DataValidation(type="list", formula1='"systemd,port,http"', allow_blank=True)
    wb["系统清单"].add_data_validation(priority)
    priority.add("C2:C500")
    for sheet, col in [("系统清单", "F"), ("网站页面", "G"), ("内容检查规则", "F"), ("管理员功能流程", "I"), ("主机清单", "G"), ("微服务清单", "H")]:
        wb[sheet].add_data_validation(yes_no)
        yes_no.add(f"{col}2:{col}500")
    wb["管理员功能流程"].add_data_validation(actions)
    actions.add("E2:E500")
    wb["微服务清单"].add_data_validation(check_types)
    check_types.add("E2:E500")
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description="Create administrator Excel inventory template.")
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    print(create_template(args.out))


if __name__ == "__main__":
    main()

