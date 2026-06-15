from __future__ import annotations

import argparse
from collections import defaultdict
from pathlib import Path

from health_common import (
    as_bool,
    as_int,
    ensure_dir,
    has_dangerous_word,
    is_valid_url,
    now_iso,
    normalize_url,
    split_list,
    write_json,
)


SHEETS = {
    "系统清单": ["system_id", "system_name", "priority", "owner", "environment", "enabled", "remark"],
    "网站页面": ["system_id", "page_id", "page_name", "url", "expected_status", "timeout_ms", "enabled"],
    "内容检查规则": ["system_id", "page_id", "required_text", "forbidden_text", "min_length", "rule_enabled"],
    "管理员功能流程": ["flow_id", "system_id", "flow_name", "step_order", "action", "target", "value", "expect", "enabled"],
    "主机清单": ["host_id", "system_id", "host_name", "ip", "port", "auth_ref", "enabled"],
    "微服务清单": ["service_id", "system_id", "host_id", "service_name", "check_type", "check_target", "expected", "enabled"],
    "凭证引用": ["auth_ref", "auth_type", "env_keys", "remark"],
    "全局配置": ["key", "value"],
}


def require_openpyxl():
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as exc:
        raise SystemExit("Missing dependency: openpyxl. Install it with `python -m pip install openpyxl` to read Excel inventories.") from exc
    return load_workbook


def rows(ws):
    headers = [str(cell.value or "").strip() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        item = {header: value for header, value in zip(headers, row) if header}
        if any(value not in (None, "") for value in item.values()):
            yield row_number, item


def load_sheet_map(workbook) -> tuple[dict, list[dict]]:
    errors = []
    data = {}
    for sheet, required in SHEETS.items():
        if sheet not in workbook.sheetnames:
            errors.append({"level": "error", "sheet": sheet, "message": "缺少必要 Sheet"})
            data[sheet] = []
            continue
        ws = workbook[sheet]
        headers = [str(cell.value or "").strip() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        missing = [name for name in required if name not in headers]
        for name in missing:
            errors.append({"level": "error", "sheet": sheet, "message": f"缺少必要字段 {name}"})
        data[sheet] = list(rows(ws))
    return data, errors


def import_excel(path: str | Path) -> tuple[dict, dict]:
    load_workbook = require_openpyxl()
    workbook = load_workbook(path, data_only=True)
    data, messages = load_sheet_map(workbook)

    global_cfg = {}
    for _, row in data["全局配置"]:
        key = str(row.get("key") or "").strip()
        if key:
            global_cfg[key] = row.get("value")

    allow_mutation = as_bool(global_cfg.get("allow_mutation"), False)
    systems_by_id = {}
    for row_num, row in data["系统清单"]:
        if not as_bool(row.get("enabled"), True):
            continue
        sid = str(row.get("system_id") or "").strip()
        if not sid:
            messages.append({"level": "error", "sheet": "系统清单", "row": row_num, "message": "system_id 不能为空"})
            continue
        if sid in systems_by_id:
            messages.append({"level": "error", "sheet": "系统清单", "row": row_num, "message": f"system_id 重复: {sid}"})
            continue
        priority = str(row.get("priority") or "medium").strip().lower()
        if priority not in {"critical", "high", "medium", "low"}:
            messages.append({"level": "warning", "sheet": "系统清单", "row": row_num, "message": f"priority={priority} 不规范，已按 medium 处理"})
            priority = "medium"
        systems_by_id[sid] = {
            "id": sid,
            "name": str(row.get("system_name") or sid).strip(),
            "priority": priority,
            "owner": str(row.get("owner") or "").strip(),
            "environment": str(row.get("environment") or "").strip(),
            "remark": str(row.get("remark") or "").strip(),
            "websites": [],
            "admin_flows": [],
            "hosts": [],
            "services": [],
        }

    content_rules = {}
    for row_num, row in data["内容检查规则"]:
        if not as_bool(row.get("rule_enabled"), True):
            continue
        key = (str(row.get("system_id") or "").strip(), str(row.get("page_id") or "").strip())
        content_rules[key] = {
            "required_text": split_list(row.get("required_text")),
            "forbidden_text": split_list(row.get("forbidden_text")),
            "min_length": as_int(row.get("min_length"), 0),
        }

    page_keys = set()
    for row_num, row in data["网站页面"]:
        if not as_bool(row.get("enabled"), True):
            continue
        sid = str(row.get("system_id") or "").strip()
        pid = str(row.get("page_id") or "").strip()
        if sid not in systems_by_id:
            messages.append({"level": "error", "sheet": "网站页面", "row": row_num, "message": f"system_id 不存在: {sid}"})
            continue
        url = normalize_url(str(row.get("url") or "").strip())
        if not is_valid_url(url):
            messages.append({"level": "error", "sheet": "网站页面", "row": row_num, "message": f"URL 无效: {url}"})
            continue
        page_keys.add((sid, pid))
        systems_by_id[sid]["websites"].append(
            {
                "page_id": pid,
                "name": str(row.get("page_name") or pid or url).strip(),
                "url": url,
                "expected_status": as_int(row.get("expected_status"), 200),
                "timeout_ms": as_int(row.get("timeout_ms"), as_int(global_cfg.get("default_timeout_ms"), 5000)),
                "content_checks": content_rules.get((sid, pid), {"required_text": [], "forbidden_text": [], "min_length": 0}),
                "source": "excel",
            }
        )

    flows = defaultdict(list)
    for row_num, row in data["管理员功能流程"]:
        if not as_bool(row.get("enabled"), True):
            continue
        sid = str(row.get("system_id") or "").strip()
        if sid not in systems_by_id:
            messages.append({"level": "error", "sheet": "管理员功能流程", "row": row_num, "message": f"system_id 不存在: {sid}"})
            continue
        action = str(row.get("action") or "").strip()
        target = str(row.get("target") or "").strip()
        if not allow_mutation and (has_dangerous_word(action) or has_dangerous_word(target)):
            messages.append({"level": "error", "sheet": "管理员功能流程", "row": row_num, "message": "检测到危险操作，allow_mutation=false 时禁止"})
            continue
        flow_id = str(row.get("flow_id") or "").strip()
        flows[(sid, flow_id)].append(
            {
                "order": as_int(row.get("step_order"), 0),
                "action": action,
                "target": target,
                "value": "" if row.get("value") is None else str(row.get("value")),
                "expect": "" if row.get("expect") is None else str(row.get("expect")),
            }
        )
    for (sid, flow_id), steps in flows.items():
        first = steps[0] if steps else {}
        source_row = next((row for _, row in data["管理员功能流程"] if str(row.get("flow_id") or "").strip() == flow_id), {})
        systems_by_id[sid]["admin_flows"].append(
            {
                "flow_id": flow_id,
                "name": str(source_row.get("flow_name") or flow_id).strip(),
                "steps": sorted(steps, key=lambda item: item.get("order", 0)),
                "source": first.get("source", "excel"),
            }
        )

    host_ids = set()
    for row_num, row in data["主机清单"]:
        if not as_bool(row.get("enabled"), True):
            continue
        sid = str(row.get("system_id") or "").strip()
        hid = str(row.get("host_id") or "").strip()
        if sid not in systems_by_id:
            messages.append({"level": "error", "sheet": "主机清单", "row": row_num, "message": f"system_id 不存在: {sid}"})
            continue
        port = as_int(row.get("port"), 22)
        if port < 1 or port > 65535:
            messages.append({"level": "error", "sheet": "主机清单", "row": row_num, "message": f"端口超出范围: {port}"})
            continue
        host_ids.add(hid)
        systems_by_id[sid]["hosts"].append(
            {
                "host_id": hid,
                "name": str(row.get("host_name") or hid).strip(),
                "ip": str(row.get("ip") or "").strip(),
                "port": port,
                "auth_ref": str(row.get("auth_ref") or "").strip(),
            }
        )

    for row_num, row in data["微服务清单"]:
        if not as_bool(row.get("enabled"), True):
            continue
        sid = str(row.get("system_id") or "").strip()
        hid = str(row.get("host_id") or "").strip()
        if sid not in systems_by_id:
            messages.append({"level": "error", "sheet": "微服务清单", "row": row_num, "message": f"system_id 不存在: {sid}"})
            continue
        if hid and hid not in host_ids:
            messages.append({"level": "warning", "sheet": "微服务清单", "row": row_num, "message": f"host_id 未在主机清单中找到: {hid}"})
        check_type = str(row.get("check_type") or "").strip().lower()
        if check_type not in {"systemd", "port", "http"}:
            messages.append({"level": "warning", "sheet": "微服务清单", "row": row_num, "message": f"check_type={check_type} 暂不支持"})
        systems_by_id[sid]["services"].append(
            {
                "service_id": str(row.get("service_id") or "").strip(),
                "host_id": hid,
                "name": str(row.get("service_name") or row.get("service_id") or "").strip(),
                "check_type": check_type,
                "check_target": str(row.get("check_target") or "").strip(),
                "expected": str(row.get("expected") or "").strip(),
            }
        )

    global_normalized = {
        "allow_mutation": allow_mutation,
        "screenshot_on_failure": as_bool(global_cfg.get("screenshot_on_failure"), True),
        "max_parallel": as_int(global_cfg.get("max_parallel"), 5),
        "crawler": {
            "enabled": as_bool(global_cfg.get("crawler_enabled"), False),
            "max_depth": as_int(global_cfg.get("crawler_max_depth"), 1),
            "max_pages": as_int(global_cfg.get("crawler_max_pages"), 20),
            "same_origin_only": True,
            "request_interval_ms": 200,
            "obey_robots_txt": True,
        },
    }

    import_report = {
        "generated_at": now_iso(),
        "source": str(path),
        "status": "failed" if any(m["level"] == "error" for m in messages) else "passed",
        "messages": messages,
        "counts": {
            "systems": len(systems_by_id),
            "websites": sum(len(s["websites"]) for s in systems_by_id.values()),
            "admin_flows": sum(len(s["admin_flows"]) for s in systems_by_id.values()),
            "hosts": sum(len(s["hosts"]) for s in systems_by_id.values()),
            "services": sum(len(s["services"]) for s in systems_by_id.values()),
        },
    }
    inventory = {
        "generated_at": now_iso(),
        "source": {"type": "excel", "path": str(path)},
        "systems": list(systems_by_id.values()),
        "global": global_normalized,
    }
    return inventory, import_report


def main() -> None:
    parser = argparse.ArgumentParser(description="Import administrator Excel inventory.")
    parser.add_argument("--excel", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()

    out_dir = ensure_dir(args.out)
    inventory, report = import_excel(args.excel)
    write_json(Path(out_dir) / "inventory-import-report.json", report)
    if report["status"] == "passed":
        write_json(Path(out_dir) / "inventory.normalized.json", inventory)
        print(Path(out_dir) / "inventory.normalized.json")
    else:
        print(Path(out_dir) / "inventory-import-report.json")
        raise SystemExit(2)


if __name__ == "__main__":
    main()
